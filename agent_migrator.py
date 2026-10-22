"""
Agent Migration — the structure-aware agent migration workflow.

Agents are moved so that each one lands in the destination group of the
same name, not in one flat site. Three steps:

  1. READ the DESTINATION — walk account → sites → groups and record every
     scope's registration token into a JSON file (`build_scope_map`).
  2. MATCH the SOURCE against it — walk the source's scopes and pair each
     one with its same-named destination scope, so the operator can see
     what was found and what wasn't before anything moves
     (`build_match_plan`). Nothing is sent.
  3. MIGRATE — send `move-to-console` per group with that group's own
     token, then read each agent back to report what actually happened
     (`run_live_migration`, `verify_migration`).

Step 3 streams events so the UI can show every machine as it goes, with
the console's own reason when one doesn't make it. A scheduled run takes
the same path, so an unattended migration is just as auditable.

Around that sit a migration status report (counts per console-migration
status, optional passphrases, decommissioned agents) and a scheduler that
fires the live run at a chosen date and time.

Everything above the UI classes is a plain function, so the mapping and
migration logic is testable without a display.
"""
import base64
import csv
import customtkinter as ctk
import difflib
import hashlib
import io
import json
import os
import queue
import re
import threading
import time
import tkinter as tk
import tkinter.font as tkfont
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from tkinter import filedialog, messagebox

import config
from app import (run_async, cli_log, _ConsoleProxy, _ToolTip, UI_FONT,
                 MONO_FONT, CARD, CARD_ELEVATED, BORDER, ACCENT, ACCENT_HOVER,
                 BRAND, BRAND_HOVER, GREEN, GREEN_HOVER, WARN, NEUTRAL,
                 NEUTRAL_HOVER, GHOST, GHOST_HOVER, TEXT, TEXT_MUTED,
                 TEXT_FAINT)
from export_utils import export_agent_report
from s1_api import S1APIError
from theme import RADIUS_LG, RADIUS_MD, SIDEBAR_SEL, tk_track

# One move-to-console call carries at most this many agent ids.
MOVE_BATCH = 1000

# Agents read back per request when confirming a move. Smaller than
# MOVE_BATCH because these ids go in the query string, not a JSON body.
VERIFY_BATCH = 100

# Group listings are fetched one site at a time, so a 64-account tenant is
# hundreds of round trips. Fan them out — same worker count as
# S1API.get_many, which the rest of the app already runs against a console.
GROUP_WORKERS = 8

# Groups counted at the same time when matching. These are read-only
# GETs, so the pool can be wider than the mapping one.
COUNT_WORKERS = 12

# The four values S1 reports in `consoleMigrationStatus`.
MIGRATION_STATUSES = ["Migrated", "Pending", "N/A", "Failed"]

AGENT_REPORT_COLUMNS = [
    "Computer Name", "Migration Status", "Site", "Group", "OS",
    "Last User", "Decommissioned", "Passphrase", "UUID", "Agent ID",
]

_CONSOLE_RE = re.compile(r"(https://[^/\s]+\.sentinelone\.net)", re.I)


class MappingError(Exception):
    """A source scope has no same-named counterpart in the map file."""

    def __init__(self, kind: str, name: str, candidates=None):
        super().__init__(f"{kind} '{name}' is not in the map file")
        self.kind = kind
        self.name = name
        self.candidates = list(candidates or [])


def parse_console_url(value: str) -> str:
    """Reduce anything pasted from a browser to the console's base URL."""
    text = (value or "").strip()
    m = _CONSOLE_RE.search(text)
    return m.group(1) if m else text.rstrip("/")


def norm_name(value) -> str:
    """Scope names are matched case-insensitively, whitespace-collapsed."""
    return " ".join(str(value or "").casefold().split())


def loose_name(value) -> str:
    """A name with case, spacing and punctuation ignored: 'shira-telecom',
    'Shira_Telecom' and 'Shira Telecom' all read 'shira telecom'. Only for
    ranking and hints — matching stays exact."""
    return " ".join(re.sub(r"[\W_]+", " ", str(value or "").casefold()).split())


def batched(seq, size: int):
    """Yield `seq` in lists of at most `size` items."""
    items = list(seq)
    for i in range(0, len(items), max(1, size)):
        yield items[i:i + max(1, size)]


# ── step 1: map the destination ────────────────────────────────────────

def split_ids(value) -> list:
    """'100, 200 300\\n400' → ['100','200','300','400']."""
    return [p for p in re.split(r"[\s,;]+", str(value or "").strip()) if p]


def resolve_accounts(api, account_ids="") -> list:
    """The accounts to work on.

    Blank means EVERY account the token can see — the normal case on a
    multi-tenant console where you're migrating all 60-odd accounts, not
    one. Otherwise each id given is resolved, and any that can't be is
    named rather than silently dropped.
    """
    wanted = split_ids(account_ids)
    if not wanted:
        accounts = [dict(a) for a in api.get_accounts()]
        if not accounts:
            raise ValueError("This token cannot see any account.")
        return accounts
    found, missing = [], []
    for aid in wanted:
        acct = api.get_account_by_id(aid)
        if acct:
            found.append(dict(acct))
        else:
            missing.append(aid)
    if missing:
        raise ValueError(
            "Not found on this console (or not visible to the token): "
            + ", ".join(missing))
    return found


def _groups_for_sites(api, sites: list, workers: int = GROUP_WORKERS) -> list:
    """Each site as a map entry with its groups, input order preserved."""
    def one(site):
        srec = {"name": site.get("name", ""), "id": str(site.get("id", "")),
                "token": site.get("registrationToken") or "",
                "type": "site", "groups": []}
        for grp in api.get_groups(params={"siteIds": srec["id"]}):
            srec["groups"].append({
                "name": grp.get("name", ""), "id": str(grp.get("id", "")),
                "token": grp.get("registrationToken") or "", "type": "group"})
        return srec

    if len(sites) < 2 or workers < 2:
        return [one(s) for s in sites]
    with ThreadPoolExecutor(max_workers=min(workers, len(sites))) as ex:
        return list(ex.map(one, sites))


def build_scope_map(api, account_ids: str = "", site_id: str = "", log=None,
                    should_stop=None, workers: int = GROUP_WORKERS,
                    site_names=None, accounts=None) -> list:
    """Destination accounts → sites → groups, each with its registration token.

    `account_ids` blank maps EVERY account the token can see; pass one id,
    or several separated by commas/spaces, to narrow it.

    Shape matches the standalone tool's JSON so existing map files still
    load: a list of account dicts, each with `sites`, each with `groups`.
    """
    say = log or (lambda _m: None)
    stop = should_stop or (lambda: False)
    if accounts is None:
        accounts = resolve_accounts(api, account_ids)
    wanted = ({norm_name(n) for n in site_names}
              if site_names is not None else None)
    say(f"{len(accounts)} account(s) to map.")

    out = []
    for n, acct in enumerate(accounts, 1):
        if stop():
            say("Stopped — the map holds only the accounts walked so far.")
            break
        acct_id = str(acct.get("id", ""))
        entry = {"name": acct.get("name", ""), "id": acct_id,
                 "type": "account", "sites": []}
        params = {"accountId": acct_id}
        if split_ids(site_id):
            params["siteIds"] = ",".join(split_ids(site_id))
        sites = api.get_sites(params=params)
        if wanted is not None:
            sites = [s for s in sites if norm_name(s.get("name")) in wanted]
        entry["sites"] = _groups_for_sites(api, sites, workers)
        no_token = [s["name"] for s in entry["sites"] if not s["token"]]
        groups = sum(len(s["groups"]) for s in entry["sites"])
        say(f"[{n}/{len(accounts)}] {entry['name']}: {len(sites)} site(s), "
            f"{groups} group(s)"
            + (f"   ⚠ no registration token on: {', '.join(no_token)}"
               if no_token else ""))
        out.append(entry)
    return out


def map_counts(map_data) -> tuple:
    """(accounts, sites, groups) held in a map file."""
    accounts = len(map_data or [])
    sites = sum(len(a.get("sites") or []) for a in map_data or [])
    groups = sum(len(s.get("groups") or [])
                 for a in map_data or [] for s in a.get("sites") or [])
    return accounts, sites, groups


def find_by_name(items, name: str, kind: str) -> dict:
    """The entry in `items` whose name matches `name`, or MappingError."""
    needle = norm_name(name)
    for item in items or []:
        if norm_name(item.get("name")) == needle:
            return item
    raise MappingError(kind, name,
                       [str(i.get("name", "")) for i in items or []])


def token_for(group: dict, site: dict) -> str:
    """A group's own registration token, else its site's."""
    return (str(group.get("token") or "").strip()
            or str(site.get("token") or "").strip())


def rename_in_map(map_data, kind: str, old_name: str, new_name: str,
                  within=None) -> bool:
    """Rename the first `kind` entry called `old_name`, in place.

    Name mismatches between the two consoles are the single most common
    reason a mapped migration skips a scope; renaming the map entry to the
    SOURCE's spelling is what makes the next run match it.
    """
    old = norm_name(old_name)
    path = [norm_name(p) for p in within or []]

    def rename(entry):
        entry.setdefault("dest_name", entry.get("name", ""))
        entry["name"] = new_name
        return True

    for acct in map_data or []:
        if kind == "account":
            if norm_name(acct.get("name")) == old:
                return rename(acct)
            continue
        if path and norm_name(acct.get("name")) != path[0]:
            continue
        for site in acct.get("sites") or []:
            if kind == "site":
                if norm_name(site.get("name")) == old:
                    return rename(site)
                continue
            if len(path) > 1 and norm_name(site.get("name")) != path[1]:
                continue
            for grp in site.get("groups") or []:
                if norm_name(grp.get("name")) == old:
                    return rename(grp)
    return False


# ── reading the source's agents ────────────────────────────────────────

def _still_to_move(rows) -> list:
    return [a for a in rows
            if str(a.get("consoleMigrationStatus") or "") != "Migrated"]


def agents_for_group(api, group_id: str, skip_migrated: bool = True) -> list:
    """The agent RECORDS in one group that still need moving.

    Records, not ids, because the live view names every machine it is
    moving — an operator watching a migration needs to see "WIN-SRV-01",
    not "f4c1a09b8e2d4f7a".

    `consoleMigrationStatusesNin` is accepted by the move-to-console
    action's filter, but not every console accepts it as a GET /agents
    query param — one that rejects it would otherwise fail the whole
    group. A 400/422 falls back to filtering the rows we got back.
    """
    base = {"groupIds": str(group_id)}
    if not skip_migrated:
        return list(api.get_agents(params=base))
    try:
        rows = api.get_agents(
            params=dict(base, consoleMigrationStatusesNin="Migrated"))
    except S1APIError as e:
        if e.status_code not in (400, 422):
            raise
        rows = api.get_agents(params=base)
    return _still_to_move(rows)


def agent_ids_for_group(api, group_id: str, skip_migrated: bool = True) -> list:
    """Just the ids, for the counting pass."""
    return [str(a.get("id"))
            for a in agents_for_group(api, group_id, skip_migrated)
            if a.get("id")]


def _unmatched(err: MappingError, path: str, within=()) -> dict:
    return {"kind": err.kind, "name": err.name, "path": path,
            "within": list(within), "candidates": err.candidates}


# ── step 2: pair every source scope with its destination ───────────────
#
# Nothing moves here. The operator gets one row per source group saying
# exactly which destination scope it will land in and how many agents go
# with it — the old build only ever showed totals, so "did it find my
# scopes?" could not be answered until after the run.

PLAN_COLUMNS = ["Account", "Site", "Group", "Destination", "Agents", "Status"]

NO_MATCH_ACCOUNT = "account not on destination"
NO_MATCH_SITE = "site not on destination"
NO_MATCH_GROUP = "group not on destination"
NO_TOKEN = "no registration token"
ALL_DONE = "already migrated"
UNREADABLE = "cannot read agents"


def _plan_row(account, site, group, **kw) -> dict:
    row = {"Account": account, "Site": site, "Group": group,
           "Destination": "—", "Agents": 0, "Total": 0, "Status": "",
           "state": "blocked", "reason": "", "kind": "", "group_id": "",
           "token": "", "token_from": "", "counted": False,
           "dest_account": "", "dest_site": "", "dest_group": ""}
    row.update(kw)
    row["key"] = f"{row['Account']}\x1f{row['Site']}\x1f{row['Group']}"
    return row


def retally(plan: dict) -> dict:
    """Recount the headline numbers from the rows, in place."""
    plan["ready"] = plan["blocked"] = plan["done"] = plan["agents"] = 0
    for row in plan.get("rows") or []:
        if row["state"] == "ready":
            plan["ready"] += 1
            plan["agents"] += row["Agents"]
        elif row["state"] == "done":
            plan["done"] += 1
        else:
            plan["blocked"] += 1
    return plan


def _source_groups(api, sites: list, workers: int = GROUP_WORKERS) -> list:
    """Each site's source groups, in the order the sites were given."""
    def one(site):
        return api.get_groups(params={"siteIds": str(site.get("id", ""))})

    if len(sites) < 2 or workers < 2:
        return [one(s) for s in sites]
    with ThreadPoolExecutor(max_workers=min(workers, len(sites))) as ex:
        return list(ex.map(one, sites))


def _plan_one_account(api, acct, map_data, site_id, plan, say, stop, on_row,
                      workers=GROUP_WORKERS) -> None:
    """Pair one source account's scopes into `plan` (mutated in place)."""
    acct_name = acct.get("name", "")

    def add(row):
        plan["rows"].append(row)
        if row["state"] == "ready":
            plan["ready"] += 1
            plan["agents"] += row["Agents"]
        elif row["state"] == "done":
            plan["done"] += 1
        else:
            plan["blocked"] += 1
        if on_row:
            on_row(row)

    try:
        m_acct = find_by_name(map_data, acct_name, "account")
    except MappingError as e:
        plan["unmatched"].append(_unmatched(e, acct_name))
        say(f"  ✗ account '{acct_name}' is not on the destination")
        add(_plan_row(acct_name, "—", "—", Status=NO_MATCH_ACCOUNT,
                      kind="account",
                      reason=f"The destination has no account named "
                             f"'{acct_name}', so nothing under it can "
                             f"move."))
        return

    params = {"accountId": str(acct.get("id", ""))}
    if split_ids(site_id):
        params["siteIds"] = ",".join(split_ids(site_id))
    sites = api.get_sites(params=params)
    say(f"  {acct_name}: {len(sites)} source site(s)")
    groups_by_site = _source_groups(api, sites, workers)

    for site, site_groups in zip(sites, groups_by_site):
        if stop():
            plan["stopped"] = True
            return
        sname = site.get("name", "")
        try:
            m_site = find_by_name(m_acct.get("sites"), sname, "site")
        except MappingError as e:
            miss = _unmatched(e, acct_name, [acct_name])
            miss["groups"] = [{"name": str(g.get("name", "")),
                               "agents": int(g.get("totalAgents") or 0)}
                              for g in site_groups
                              if int(g.get("totalAgents") or 0)]
            plan["unmatched"].append(miss)
            say(f"    ✗ site '{sname}' is not on the destination")
            add(_plan_row(acct_name, sname, "—", Status=NO_MATCH_SITE,
                          kind="site",
                          reason=f"Account '{acct_name}' on the destination "
                                 f"has no site named '{sname}'."))
            continue

        for grp in site_groups:
            if stop():
                plan["stopped"] = True
                return
            gname = grp.get("name", "")
            total = int(grp.get("totalAgents") or 0)
            if total == 0:
                plan["empty_groups"] += 1
                continue
            try:
                m_grp = find_by_name(m_site.get("groups"), gname, "group")
            except MappingError as e:
                plan["unmatched"].append(
                    _unmatched(e, f"{acct_name} / {sname}",
                               [acct_name, sname]))
                say(f"    ✗ group '{gname}' is not on the destination")
                add(_plan_row(acct_name, sname, gname, Total=total,
                              Status=NO_MATCH_GROUP, kind="group",
                              reason=f"Site '{sname}' on the destination has "
                                     f"no group named '{gname}' — its "
                                     f"{total} agent(s) stay put."))
                continue

            token = token_for(m_grp, m_site)
            dest = " / ".join(str(e.get("dest_name") or e.get("name", ""))
                              for e in (m_acct, m_site, m_grp))
            common = dict(Destination=dest, Total=total,
                          group_id=str(grp.get("id", "")), token=token,
                          dest_account=m_acct.get("name", ""),
                          dest_site=m_site.get("name", ""),
                          dest_group=m_grp.get("name", ""))
            if not token:
                add(_plan_row(acct_name, sname, gname, Status=NO_TOKEN,
                              reason=f"Neither the destination group "
                                     f"'{gname}' nor its site '{sname}' has "
                                     f"a registration token. Create one on "
                                     f"the destination console.", **common))
                continue

            from_where = "group" if str(m_grp.get("token") or "").strip() \
                else "site"
            # The group listing's own total, so the pairing shows up at
            # once; refine_plan_counts corrects it to what still needs
            # moving. One GET per group is what used to make this slow.
            add(_plan_row(acct_name, sname, gname, state="ready",
                          Agents=total,
                          Status=f"ready · {from_where} token",
                          token_from=from_where, **common))


def refine_plan_counts(api, plan: dict, on_row=None, should_stop=None,
                       workers: int = COUNT_WORKERS) -> dict:
    """Correct every ready row's count to what still needs moving.

    The pairing pass uses the group listing's own total so the operator
    sees the whole plan at once. That total counts agents that already
    migrated, so each ready group is read once here to get the real
    number — in parallel, because one blocking GET per group is what
    made this step a coffee break on a 64-account console.
    """
    stop = should_stop or (lambda: False)
    rows = [r for r in plan.get("rows") or [] if r["state"] == "ready"]
    if not rows:
        plan["counted"] = True
        return plan

    def one(row):
        if stop():
            return row, None, None
        try:
            return row, len(agent_ids_for_group(api, row["group_id"])), None
        except Exception as e:
            return row, None, e

    with ThreadPoolExecutor(max_workers=max(1, min(workers, len(rows)))) as ex:
        for row, eligible, err in ex.map(one, rows):
            if err is not None:
                row.update(state="blocked", Status=UNREADABLE, reason=str(err))
            elif eligible is None:
                plan["stopped"] = True
                continue
            elif eligible == 0:
                row.update(state="done", Status=ALL_DONE, Agents=0,
                           reason=f"All {row['Total']} agent(s) have already "
                                  f"migrated.")
            else:
                row["Agents"] = eligible
            row["counted"] = True
            if on_row:
                on_row(row)

    retally(plan)
    plan["counted"] = True
    return plan


def build_match_plan(api, map_data, account_ids: str = "", site_id: str = "",
                     log=None, should_stop=None, on_row=None,
                     count_agents: bool = True,
                     workers: int = GROUP_WORKERS) -> dict:
    """Every source scope paired with the destination scope it maps to.

    Nothing is sent. Two passes, so the operator is never left staring at
    a spinner: the first pairs every scope and reports the group's own
    agent total, which is enough to answer "did it find my scopes and
    where is each one going?"; the second (`refine_plan_counts`) corrects
    those totals to what still needs moving.

    `on_row` fires for each row in both passes — the same row object, so
    a UI keyed on `row["key"]` updates in place rather than duplicating.
    """
    say = log or (lambda _m: None)
    stop = should_stop or (lambda: False)
    plan = {"rows": [], "unmatched": [], "accounts": 0, "ready": 0,
            "blocked": 0, "done": 0, "agents": 0, "empty_groups": 0,
            "stopped": False, "counted": False}

    accounts = resolve_accounts(api, account_ids)
    say(f"Reading {len(accounts)} source account(s).")
    for n, acct in enumerate(accounts, 1):
        if stop():
            plan["stopped"] = True
            break
        say(f"[{n}/{len(accounts)}] {acct.get('name', '')}")
        plan["accounts"] += 1
        _plan_one_account(api, acct, map_data, site_id, plan, say, stop,
                          on_row, workers)

    if count_agents and not plan["stopped"]:
        say(f"Counting agents in {plan['ready']} matched group(s)…")
        refine_plan_counts(api, plan, on_row, stop)
    return plan


_RENAME_ORDER = {"account": 0, "site": 1, "group": 2}


def apply_renames(map_data, renames) -> int:
    """Re-apply the operator's name fixes to a map, accounts first."""
    done = 0
    for r in sorted(renames or [],
                    key=lambda r: _RENAME_ORDER.get(r.get("kind"), 3)):
        if rename_in_map(map_data, r.get("kind", ""), r.get("old", ""),
                         r.get("new", ""), r.get("within") or None):
            done += 1
    return done


def read_destination(src_api, dst_api, account_ids: str = "",
                     site_ids: str = "", renames=(), log=None,
                     should_stop=None, workers: int = GROUP_WORKERS) -> tuple:
    """(destination map, every destination account name) for the source
    scope being migrated.

    Only the destination accounts — and, with a site filter, the sites —
    whose names match the source scope are read, so moving one account
    out of sixty walks one destination account. A name fix counts as a
    match, and is applied to the map before it is returned.
    """
    say = log or (lambda _m: None)
    renames = list(renames or [])
    sources = resolve_accounts(src_api, account_ids)
    want = {norm_name(a.get("name")) for a in sources}
    alias = {norm_name(r.get("old")): norm_name(r.get("new"))
             for r in renames if r.get("kind") == "account"}
    everyone = [dict(a) for a in dst_api.get_accounts()]
    if not everyone:
        raise ValueError("The destination token cannot see any account.")
    twins = [a for a in everyone
             if norm_name(a.get("name")) in want
             or alias.get(norm_name(a.get("name"))) in want]
    names = [str(a.get("name", "")) for a in everyone]

    site_names = None
    if split_ids(site_ids):
        chosen = src_api.get_sites(
            params={"siteIds": ",".join(split_ids(site_ids))})
        picked = {norm_name(s.get("name")) for s in chosen}
        site_names = picked | {norm_name(r.get("old")) for r in renames
                               if r.get("kind") == "site"
                               and norm_name(r.get("new")) in picked}

    say(f"Destination: {len(twins)} of {len(everyone)} account(s) match the "
        f"{len(sources)} source account(s) in scope.")
    if not twins:
        return [], names
    map_data = build_scope_map(dst_api, log=log, should_stop=should_stop,
                               workers=workers, site_names=site_names,
                               accounts=twins)
    apply_renames(map_data, renames)
    return map_data, names


# ── fixing names: the source beside the destination ────────────────────
#
# A scope with no same-named twin used to be one line in a flat list of
# dropdowns: no structure, no view of what the destination actually has
# in that place, and nothing to stop a pick that was already another
# scope's match. `NameComparison` is that picture. It is built from a
# checked plan and the destination read behind it, so it costs no call.

FIX_STATES = ("missing", "chosen", "waiting", "fixed", "unfixed", "same")
FIX_FILTERS = {"all": None, "fix": ("missing", "chosen", "waiting"),
               "done": ("chosen", "fixed", "unfixed")}


def _plural(count: int, word: str) -> str:
    return f"{count:,} {word}{'' if count == 1 else 's'}"


class NameComparison:
    """The source scope as a tree — account › site › group — each scope
    beside what the destination has in the same place.

    A scope's state is one of:
      same     the destination has a scope of this name in the same place
      fixed    paired by a name fix made earlier (`unfixed`: being undone)
      chosen   the operator picked its destination scope here
      missing  nothing of this name in the same place on the destination
      waiting  its parent isn't paired yet, so there is nothing to compare

    A pick is refused when the destination scope already pairs with
    another source scope: a name fix renames the destination entry, so a
    second scope pointed at it would take it away from the first.
    """

    def __init__(self, plan, map_data, dest_accounts=(), renames=(),
                 partial_sites=False):
        self.map = list(map_data or [])
        self.dest_accounts = [str(n) for n in dest_accounts or []
                              if str(n or "").strip()]
        self.applied = [dict(r) for r in renames or []]
        self.partial_sites = bool(partial_sites)
        self.nodes, self.roots, self.choices = {}, [], {}
        self._typed, self._unfixed, self._index = set(), set(), {}
        self._build(plan or {})

    @staticmethod
    def key(*names) -> tuple:
        return tuple(norm_name(n) for n in names)

    # ── the source tree ───────────────────────────────────────────────

    def _node(self, kind, names):
        key = self.key(*names)
        node = self.nodes.get(key)
        if node is None:
            parent = key[:-1] or None
            node = self.nodes[key] = {
                "key": key, "kind": kind, "name": str(names[-1]),
                "path": [str(n) for n in names], "parent": parent,
                "children": [], "agents": 0}
            (self.roots if parent is None
             else self.nodes[parent]["children"]).append(key)
        return node

    def _build(self, plan):
        blank = ("", "—")
        for row in plan.get("rows") or []:
            acct, site, group = (str(row.get(k) or "")
                                 for k in ("Account", "Site", "Group"))
            if acct in blank:
                continue
            self._node("account", [acct])
            if site in blank:
                continue
            self._node("site", [acct, site])
            if group not in blank:
                node = self._node("group", [acct, site, group])
                node["agents"] = max(node["agents"], int(
                    row.get("Total") or row.get("Agents") or 0))
        for miss in plan.get("unmatched") or []:
            within = list(miss.get("within") or [])
            if miss.get("kind") != "site" or len(within) != 1:
                continue
            names = [within[0], str(miss.get("name") or "")]
            self._node("account", names[:1])
            self._node("site", names)
            for grp in miss.get("groups") or []:
                node = self._node("group",
                                  names + [str(grp.get("name") or "")])
                node["agents"] = max(node["agents"],
                                     int(grp.get("agents") or 0))

    def agents(self, key) -> int:
        node = self.nodes[key]
        if node["kind"] == "group":
            return node["agents"]
        return sum(self.agents(child) for child in node["children"])

    # ── the destination beside it ─────────────────────────────────────

    def _lookup(self, entries, name):
        if not entries:
            return None
        index = self._index.get(id(entries))
        if index is None:
            index = self._index[id(entries)] = {}
            for entry in entries:
                index.setdefault(norm_name(entry.get("name")), entry)
        return index.get(norm_name(name))

    def _place(self, key):
        """(paired, entries): whether this scope's parent has a destination
        scope, and the destination scopes at this level inside it."""
        node = self.nodes[key]
        if node["parent"] is None:
            return True, self.map
        parent = self.dest(node["parent"])
        if parent is None:
            return False, None
        return True, parent.get("sites" if node["kind"] == "site"
                                else "groups")

    def dest(self, key):
        """The destination scope this source scope pairs with, or None."""
        paired, entries = self._place(key)
        if not paired:
            return None
        return self._lookup(entries,
                            self.choices.get(key, self.nodes[key]["name"]))

    def _within(self, key) -> list:
        return [norm_name(p) for p in self.nodes[key]["path"][:-1]]

    def _fixes_here(self, key) -> list:
        kind, within = self.nodes[key]["kind"], self._within(key)
        return [f for f in self.applied if f.get("kind") == kind
                and [norm_name(p) for p in f.get("within") or []] == within]

    def fix_for(self, key):
        """The name fix made earlier that pairs this scope, if any."""
        name = norm_name(self.nodes[key]["name"])
        return next((f for f in self._fixes_here(key)
                     if norm_name(f.get("new")) == name), None)

    def state(self, key) -> str:
        if key in self.choices:
            return "chosen"
        paired, entries = self._place(key)
        if not paired:
            return "waiting"
        if self._lookup(entries, self.nodes[key]["name"]) is None:
            return "missing"
        if self.fix_for(key) is not None:
            return "unfixed" if key in self._unfixed else "fixed"
        return "same"

    def dest_name(self, key) -> str:
        """The destination's own name for where this scope lands."""
        state = self.state(key)
        if state == "chosen":
            return self.choices[key]
        if state in ("fixed", "unfixed"):
            return str(self.fix_for(key).get("old") or "")
        entry = self.dest(key) if state == "same" else None
        return str(entry.get("name") or "") if entry else ""

    def where(self, key) -> str:
        """The destination scope this one is looked for in, in words."""
        parent = self.nodes[key]["parent"]
        if parent is None:
            return "the destination"
        return (f"the destination {self.nodes[parent]['kind']} "
                f"“{self.dest_name(parent)}”")

    def candidates(self, key) -> list:
        """The destination scopes in this scope's place, closest name
        first, each with whatever already pairs with it."""
        node = self.nodes[key]
        paired, entries = self._place(key)
        if not paired:
            return []
        fixes = self._fixes_here(key)
        by_new = {norm_name(f.get("new")): f for f in fixes}
        by_old = {norm_name(f.get("old")): f for f in fixes}
        if node["kind"] == "account" and self.dest_accounts:
            pool = [(n, n) for n in self.dest_accounts]
        else:
            pool = []
            for entry in entries or []:
                current = str(entry.get("name") or "")
                fix = by_new.get(norm_name(current))
                pool.append((current, str(fix.get("old")) if fix else current))
        siblings = (self.roots if node["parent"] is None
                    else self.nodes[node["parent"]]["children"])
        source = {norm_name(self.nodes[k]["name"]): self.nodes[k]["name"]
                  for k in siblings if k != key}
        picked = {norm_name(v): self.nodes[k]["name"]
                  for k, v in self.choices.items()
                  if k != key and self.nodes[k]["parent"] == node["parent"]}
        out, seen = [], set()
        for current, original in pool:
            n = norm_name(original)
            if not n or n in seen:
                continue
            seen.add(n)
            fix = by_old.get(n) or by_new.get(norm_name(current))
            out.append({"name": current, "shown": original,
                        "taken_by": (str(fix.get("new") or "") if fix
                                     else source.get(n, "")),
                        "chosen_for": picked.get(norm_name(current), ""),
                        "fixed": bool(fix)})
        target = node["name"].casefold()

        def closeness(c):
            other = c["shown"].casefold()
            score = difflib.SequenceMatcher(None, target, other).ratio()
            return score + (0.5 if other and (other in target
                                              or target in other) else 0)

        out.sort(key=lambda c: (bool(c["taken_by"] or c["chosen_for"]),
                                -closeness(c), c["shown"].casefold()))
        return out

    def check(self, key, name) -> tuple:
        """("free" | "taken" | "unknown", detail) for picking `name` here —
        the destination's name when free, who holds it when taken."""
        wanted = norm_name(name)
        for c in self.candidates(key):
            if wanted in (norm_name(c["shown"]), norm_name(c["name"])):
                holder = c["taken_by"] or c["chosen_for"]
                return ("taken", holder) if holder else ("free", c["name"])
        return "unknown", ""

    def elsewhere(self, key) -> list:
        """Where else the destination (as read) has a scope of this name."""
        node = self.nodes[key]
        if node["kind"] == "account":
            return []
        name, here = norm_name(node["name"]), self._place(key)[1]
        found = []
        for acct in self.map:
            sites = acct.get("sites") or []
            if node["kind"] == "site":
                if sites is not here and any(
                        norm_name(s.get("name")) == name for s in sites):
                    found.append(str(acct.get("name") or ""))
                continue
            for site in sites:
                groups = site.get("groups") or []
                if groups is not here and any(
                        norm_name(g.get("name")) == name for g in groups):
                    found.append(f"{acct.get('name', '')} › "
                                 f"{site.get('name', '')}")
        return found

    def reason(self, key) -> str:
        """Why this scope does or doesn't pair, in a sentence or two."""
        node = self.nodes[key]
        kind, name, state = node["kind"], node["name"], self.state(key)
        place = self.where(key)
        lead = place[:1].upper() + place[1:]
        art = "an" if kind == "account" else "a"
        if state == "same":
            return (f"{lead} has {art} {kind} named “{name}”, so it pairs by "
                    f"name — nothing to fix.")
        if state in ("fixed", "unfixed"):
            old = self.dest_name(key)
            if state == "unfixed":
                return (f"Fixed earlier to the destination {kind} “{old}”. "
                        f"That fix is removed when you apply, and the {kind} "
                        f"is compared by name again.")
            return (f"Fixed earlier: {place} has no {kind} named “{name}”, "
                    f"so its {kind} “{old}” is used instead.")
        if state == "chosen":
            pick = self.choices[key]
            if key in self._typed:
                return (f"You typed “{pick}”. No {kind} of that name was read "
                        f"from {place}, so it only pairs if one exists there "
                        f"with exactly this name.")
            return f"Will use the {kind} “{pick}” in {place}."
        if state == "waiting":
            parent = self.nodes[node["parent"]]
            return (f"Not compared yet: the {parent['kind']} "
                    f"“{parent['name']}” has no destination {parent['kind']} "
                    f"yet. Pair it first, and this {kind} is compared with "
                    f"what is inside it.")
        cands = self.candidates(key)
        free = sum(1 for c in cands if not (c["taken_by"] or c["chosen_for"]))
        text = f"{lead} has no {kind} named “{name}”."
        text += (f" It has {_plural(len(cands), kind)} — {free:,} not paired "
                 f"with any source {kind}." if cands
                 else f" No {kind}s were read from it.")
        if kind == "account" and not node["children"]:
            text += (" Its sites and groups are compared once it is paired "
                     "and checked again.")
        if kind == "site" and self.partial_sites:
            text += (" Only the sites named in step 1 were read, so its other "
                     "sites aren't listed — type the destination site's exact "
                     "name, or clear Sites in step 1 and check again.")
        return text

    def dest_text(self, key) -> str:
        """What the tree shows in its destination column."""
        state = self.state(key)
        if state == "same":
            return "✓ same name"
        if state == "fixed":
            return f"✎ “{self.dest_name(key)}” · fixed earlier"
        if state == "unfixed":
            return f"↺ “{self.dest_name(key)}” · removing fix"
        if state == "chosen":
            return (f"→ “{self.choices[key]}”"
                    + (" · typed" if key in self._typed else ""))
        if state == "waiting":
            parent = self.nodes[self.nodes[key]["parent"]]
            return f"· pair the {parent['kind']} first"
        return "✕ not found"

    # ── the operator's picks ──────────────────────────────────────────

    def choose(self, key, name, allow_unknown=False) -> bool:
        """Pair a scope that has no twin with the destination scope `name`.
        False when that scope already pairs with another — or, unless
        `allow_unknown`, when no scope of that name was read there."""
        if self.state(key) not in ("missing", "chosen"):
            return False
        verdict, value = self.check(key, name)
        if verdict == "taken":
            return False
        if verdict == "unknown":
            value = " ".join(str(name or "").split())
            if not (allow_unknown and value):
                return False
        if self.choices.get(key) != value:
            self._forget_below(key)
        self.choices[key] = value
        (self._typed.add if verdict == "unknown" else self._typed.discard)(key)
        return True

    def unchoose(self, key):
        self._forget_below(key)
        self.choices.pop(key, None)
        self._typed.discard(key)

    def _forget_below(self, key):
        """A parent's pick decides what its children are compared with."""
        for child in self.nodes[key]["children"]:
            self._forget_below(child)
            self.choices.pop(child, None)
            self._typed.discard(child)

    def typed(self, key) -> bool:
        return key in self._typed

    def unfix(self, key, keep=False) -> bool:
        """Mark a fix made earlier for removal — or `keep` it after all."""
        if self.fix_for(key) is None:
            return False
        (self._unfixed.discard if keep else self._unfixed.add)(key)
        return True

    def removed(self) -> list:
        """The fixes made earlier that are being removed."""
        return [self.fix_for(k) for k in sorted(self._unfixed)
                if self.fix_for(k) is not None]

    def chosen(self) -> list:
        """(scope, destination name) per pick, parents first."""
        picks = sorted(self.choices.items(),
                       key=lambda kv: (len(kv[0]), kv[0]))
        return [({"kind": self.nodes[k]["kind"], "name": self.nodes[k]["name"],
                  "within": self.nodes[k]["path"][:-1]}, dest)
                for k, dest in picks]

    def renames(self) -> list:
        """The picks as the name fixes `apply_renames` understands."""
        return [{"kind": item["kind"], "old": dest, "new": item["name"],
                 "within": list(item["within"])}
                for item, dest in self.chosen()]

    # ── what to show ──────────────────────────────────────────────────

    def visible(self, filt="all", query="", collapsed=()) -> list:
        """(key, depth) in display order: what the filter keeps, the scopes
        above it for context, and — while searching — everything under a
        match, folded or not."""
        wanted, q = FIX_FILTERS.get(filt), norm_name(query)

        def walk(key, inherited):
            node = self.nodes[key]
            hit = (inherited or not q or q in norm_name(node["name"])
                   or q in norm_name(self.dest_name(key)))
            below = [row for child in node["children"]
                     for row in walk(child, hit and bool(q))]
            keep = hit and (wanted is None or self.state(key) in wanted)
            if not keep and not below:
                return []
            row = [(key, len(key) - 1)]
            return row if key in collapsed and not q else row + below

        return [row for root in self.roots for row in walk(root, False)]

    def summary(self) -> dict:
        counts = dict.fromkeys(FIX_STATES, 0)
        for key in self.nodes:
            counts[self.state(key)] += 1
        return counts


# ── or: a site-token file instead of reading the destination ───────────
#
# Sometimes nobody has API access to the destination at all — a partner
# sends a sheet with each new site's name and its registration token
# (“Site Name / Site Token”). That is enough to migrate: every source
# site is matched by name to a row of the file, and its agents move with
# that row's token. Nothing is guessed. A name must match exactly
# (ignoring case and spacing), a name that would match in more than one
# place is refused, and everything that finds no partner is listed with
# the reason before a single agent moves.

SITE_TOKEN_FILETYPES = [
    ("Site tokens (CSV, TSV, Excel)", "*.csv *.tsv *.txt *.xlsx *.xlsm"),
    ("All files", "*.*"),
]

SITE_TOKEN_COLUMNS = ["Site", "Lands on", "Token", "Check"]

NOT_IN_FILE = "site not in the file"
NOT_ON_SOURCE = "no such site on the source"
FILE_PROBLEM = "fix this row of the file"
SAME_CONSOLE = "token is for the source console"
AMBIGUOUS = "name matches several sites"
NOTHING_TO_MOVE = "no agents to move"

# Words that say what a header names. A token column is anything with
# "token" in it; the rest go by whole words.
_HEAD_SOURCE = {"source", "from", "old", "origin", "originating", "current"}
_HEAD_DEST = {"destination", "dest", "target", "new", "to"}
_HEAD_NAME = {"site", "sites", "name", "customer", "client", "tenant"}

# "<site name> <whitespace> <long token>": a line pasted out of an e-mail
# that lost its tab on the way.
_NAME_THEN_TOKEN = re.compile(r"^\s*(.*?\S)\s+([A-Za-z0-9+/_=-]{40,})\s*$")


def decode_token(token) -> dict:
    """What a registration token says about itself, or {}.

    A SentinelOne registration token is base64 of a small JSON object —
    {"url": <console>, "site_key": <key>} — so it can be checked, and its
    console named, without asking any console.
    """
    text = "".join(str(token or "").split())
    if not text:
        return {}
    padded = text + "=" * (-len(text) % 4)
    for decode in (base64.b64decode, base64.urlsafe_b64decode):
        try:
            data = json.loads(decode(padded).decode("utf-8"))
        except (ValueError, TypeError):
            continue
        if isinstance(data, dict):
            return data
    return {}


def console_host(url) -> str:
    """'https://euce1-acme.sentinelone.net' → 'euce1-acme'."""
    text = parse_console_url(str(url or ""))
    return text.split("://", 1)[-1].split(".")[0] if text else ""


def token_hint(token) -> str:
    """Enough of a token to tell two apart on screen, never the secret."""
    key = str(decode_token(token).get("site_key") or "")
    return f"key …{key[-6:]}" if len(key) > 6 else mask_token(token)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _at(row, col) -> str:
    return row[col] if col is not None and col < len(row) else ""


def _decode_text(raw: bytes) -> str:
    """File bytes → text. Excel saves “CSV UTF-8” with a BOM, “Unicode
    Text” as UTF-16, and plain “CSV” in the Windows code page."""
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _text_rows(text: str) -> list:
    """Delimited text → rows of cells, whatever the separator.

    The separator is read off the first line — a header has no commas in
    it, a site name might. Blank lines are kept as empty rows so the row
    numbers quoted back match the file. Text pasted out of an e-mail can
    lose its tabs; then each line is split before the long token at its
    end.
    """
    text = str(text or "")
    first = next((ln for ln in text.splitlines() if ln.strip()), "")
    if not first:
        return []
    delim = max(("\t", ",", ";", "|"), key=first.count)
    if first.count(delim):
        return list(csv.reader(io.StringIO(text, newline=""),
                               delimiter=delim))
    rows = []
    for line in text.splitlines():
        m = _NAME_THEN_TOKEN.match(line)
        rows.append([m.group(1), m.group(2)] if m
                    else ([line.strip()] if line.strip() else []))
    return rows


def _header_kind(cell) -> str:
    """Which role a header cell names — token, account, group, source,
    dest or name — or "" for a column this ignores."""
    text = " ".join(re.sub(r"[^0-9a-z]+", " ", str(cell or "").casefold())
                    .split())
    if not text or len(text) > 40:
        return ""
    words = set(text.split())
    if "token" in text:
        return "token"
    if words & {"account", "accounts"}:
        return "account"
    if words & {"group", "groups"}:
        return "group"
    if words & _HEAD_SOURCE:
        return "source"
    if words & _HEAD_DEST:
        return "dest"
    if words & _HEAD_NAME:
        return "name"
    return ""


def _find_columns(rows) -> tuple:
    """(header row index, {role: column}) or (None, None).

    Looks a few rows down, so a title above the table doesn't matter.
    """
    for i, row in enumerate(rows[:10]):
        first = {}
        for j, cell in enumerate(row):
            first.setdefault(_header_kind(cell), j)
        if "token" not in first:
            continue
        src = next((first[k] for k in ("source", "name", "dest")
                    if k in first), None)
        if src is None:
            continue
        dst = first.get("dest")
        if dst is None and first.get("name", src) != src:
            dst = first["name"]
        return i, {"token": first["token"], "source": src,
                   "dest": src if dst is None else dst,
                   "account": first.get("account"),
                   "header": [c for c in row if c]}
    return None, None


def _guess_columns(rows):
    """Columns for a sheet with no header: the one full of tokens, and
    the first other column with anything in it."""
    width = max((len(r) for r in rows), default=0)
    best, best_hits = None, 0
    for j in range(width):
        cells = [r[j] for r in rows if j < len(r) and r[j]]
        hits = sum(1 for c in cells if decode_token(c).get("site_key"))
        if hits > best_hits and hits * 2 >= len(cells):
            best, best_hits = j, hits
    if best is None:
        return None
    name = next((j for j in range(width) if j != best
                 and any(j < len(r) and r[j] for r in rows)), None)
    if name is None:
        return None
    return {"token": best, "source": name, "dest": name, "account": None,
            "header": []}


def _check_token(entry) -> None:
    """Fill in the row's console, or say what is wrong with its token."""
    if not entry["token"]:
        entry["problem"] = "no token in the file"
        return
    info = decode_token(entry["token"])
    url = parse_console_url(str(info.get("url") or ""))
    if not url or not info.get("site_key"):
        entry["problem"] = ("the token does not decode — it is not a "
                            "SentinelOne registration token")
        return
    entry["console"] = url


def parse_site_tokens(rows, source: str = "") -> dict:
    """A site-token sheet → the rows step 2 matches the source against.

    Columns are found by their headers, wherever they sit: a token column
    (“Site Token”, “Token”, …) and a site column (“Site Name”, “Site”,
    …). Optional: “Source Site” when the source names the site
    differently, “Destination Site” for the name it lands under, and
    “Source Account” to say which account a site name belongs to. A sheet
    without a header works too, if one column is all tokens.

    Every row is checked before any console is asked: the token has to
    decode, a site may be listed only once, and a row with a site but no
    token is kept rather than dropped — so the plan can say what's wrong.
    """
    grid = [[_cell(c) for c in row] for row in rows or []]
    start, cols = _find_columns(grid)
    headless = start is None
    if headless:
        cols = _guess_columns(grid)
        if cols is None:
            raise ValueError(
                "No site and token columns found. The file needs a column "
                "of site names and one headed “Site Token” (or “Token”).")
        start = -1
    entries, skipped, seen = [], [], {}
    for n in range(start + 1, len(grid)):
        row = grid[n]
        token = "".join(_at(row, cols["token"]).split())
        name = _at(row, cols["source"])
        if not (name or token or _at(row, cols["dest"])):
            continue
        if headless and not decode_token(token).get("site_key"):
            skipped.append({"row": n + 1, "reason": "no token on this line"})
            continue
        if not name:
            skipped.append({"row": n + 1, "reason": "no site name"})
            continue
        entry = {"row": n + 1, "source": name,
                 "dest": _at(row, cols["dest"]) or name,
                 "account": _at(row, cols["account"]), "token": token,
                 "console": "", "problem": "", "paired": False}
        _check_token(entry)
        key = (norm_name(entry["account"]), norm_name(name))
        first = seen.get(key)
        if first is None:
            seen[key] = entry
            entries.append(entry)
            continue
        skipped.append({"row": n + 1, "reason": f"repeats row {first['row']}"})
        if token != first["token"]:
            first.setdefault("clash", [first["row"]]).append(n + 1)
            first["problem"] = ("listed more than once with different tokens "
                                f"(rows {', '.join(map(str, first['clash']))})")
    if not entries:
        raise ValueError("Found the columns, but no site rows under them.")
    return {"source": source, "header": cols["header"], "headless": headless,
            "entries": entries, "skipped": skipped,
            "consoles": sorted({e["console"] for e in entries
                                if e["console"]})}


def parse_site_tokens_text(text: str, source: str = "pasted text") -> dict:
    """Site tokens copied out of Excel or an e-mail."""
    return parse_site_tokens(_text_rows(text), source=source)


def read_site_tokens(path: str) -> dict:
    """Load a site-token file: CSV, TSV or text, or an Excel workbook."""
    ext = os.path.splitext(path)[1].casefold()
    if ext == ".xls":
        raise ValueError("Old-style .xls workbooks can't be read — open it "
                         "in Excel and save it as .xlsx or CSV.")
    if ext in (".xlsx", ".xlsm"):
        return _read_workbook(path)
    with open(path, "rb") as f:
        return parse_site_tokens(_text_rows(_decode_text(f.read())),
                                 source=path)


def _read_workbook(path: str) -> dict:
    """The first sheet of a workbook that holds a site-token table."""
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    error = None
    try:
        for sheet in book.worksheets:
            try:
                return parse_site_tokens(
                    sheet.iter_rows(min_row=1, values_only=True), source=path)
            except ValueError as e:
                error = error or e
    finally:
        book.close()
    raise error or ValueError("That workbook has no sheets.")


def file_destination(entry) -> str:
    """Where a file row's agents land: its site name and console."""
    host = console_host(entry.get("console"))
    return f"{entry['dest']}  ·  {host}" if host else entry["dest"]


def _entry_for(by_name, acct, site):
    """The file row naming this source site, or None. A row that says
    which account it belongs to beats one that doesn't."""
    rows = by_name.get(norm_name(site.get("name")), [])
    acct_name, acct_id = norm_name(acct.get("name")), str(acct.get("id", ""))
    for entry in rows:
        if entry["account"] and (norm_name(entry["account"]) == acct_name
                                 or entry["account"] == acct_id):
            return entry
    return next((e for e in rows if not e["account"]), None)


def build_file_plan(api, site_tokens, account_ids: str = "",
                    site_id: str = "", log=None, should_stop=None,
                    on_row=None, count_agents: bool = True,
                    workers: int = GROUP_WORKERS,
                    source_url: str = "") -> dict:
    """Every source site paired with the row of a site-token file that
    names it. Nothing is sent.

    Same plan shape as `build_match_plan`, so step 3, the counts, the
    reports and the scheduler work unchanged: one ready row per source
    group with agents, carrying its site's token from the file. On top,
    `unlisted` (source sites with agents the file doesn't mention) and
    `unused` (file rows no source site matched) feed the pairing dialog,
    and so does `spare`: with sites chosen, the file's other rows aren't
    checked against the source, but one of them may be the chosen site
    spelt differently.
    """
    say = log or (lambda _m: None)
    stop = should_stop or (lambda: False)
    entries = list((site_tokens or {}).get("entries") or [])
    plan = {"rows": [], "unmatched": [], "accounts": 0, "ready": 0,
            "blocked": 0, "done": 0, "agents": 0, "empty_groups": 0,
            "stopped": False, "counted": False, "mode": "file",
            "file": (site_tokens or {}).get("source", ""),
            "unlisted": [], "unused": [], "spare": [], "quiet_accounts": 0}
    by_name = {}
    for entry in entries:
        by_name.setdefault(norm_name(entry["source"]), []).append(entry)
    home = parse_console_url(source_url).casefold() if source_url else ""

    accounts = resolve_accounts(api, account_ids)
    say(f"Reading {len(accounts)} source account(s) against "
        f"{len(entries)} site(s) in the file.")
    walked = []
    for n, acct in enumerate(accounts, 1):
        if stop():
            plan["stopped"] = True
            break
        plan["accounts"] += 1
        params = {"accountId": str(acct.get("id", ""))}
        if split_ids(site_id):
            params["siteIds"] = ",".join(split_ids(site_id))
        sites = api.get_sites(params=params)
        say(f"[{n}/{len(accounts)}] {acct.get('name', '')}: "
            f"{len(sites)} source site(s)")
        for site, groups in zip(sites, _source_groups(api, sites, workers)):
            walked.append((acct, site, groups,
                           _entry_for(by_name, acct, site)))

    # One row reaching several source sites — "Default site" exists in
    # every account — is a guess, and a wrong guess hands one customer's
    # agents to another. Those are refused, never picked between.
    reach = {}
    for acct, site, _groups, entry in walked:
        if entry is not None:
            reach.setdefault(id(entry), []).append(
                f"{acct.get('name', '')} / {site.get('name', '')}")

    # With no account named, only the accounts the file talks about are in
    # play — otherwise every other customer's sites would be listed as
    # "not in the file" and bury the ones that matter.
    named = bool(split_ids(account_ids))
    in_play = {str(a.get("id", "")) for a, _s, _g, e in walked
               if e is not None}
    if not named:
        plan["quiet_accounts"] = plan["accounts"] - len(in_play)
    if plan["quiet_accounts"]:
        say(f"{plan['quiet_accounts']} account(s) have no site named in the "
            f"file and are left out of the plan.")

    used = {id(entry) for _a, _s, _g, entry in walked if entry is not None}
    leftovers = [e for e in entries if id(e) not in used]
    # The rows a site the file doesn't name could still take — unknown
    # after a stop, when the sites not yet read might use any of them.
    free = [] if plan["stopped"] else [e for e in leftovers
                                       if not e["problem"]]

    def add(row):
        plan["rows"].append(row)
        if on_row:
            on_row(row)

    for acct, site, groups, entry in walked:
        acct_name, sname = acct.get("name", ""), site.get("name", "")
        live = [g for g in groups if int(g.get("totalAgents") or 0) > 0]
        plan["empty_groups"] += len(groups) - len(live)
        total = sum(int(g.get("totalAgents") or 0) for g in live)
        if entry is None:
            if total and (named or str(acct.get("id", "")) in in_play):
                plan["unlisted"].append({
                    "account": acct_name, "account_id": str(acct.get("id", "")),
                    "site": sname, "site_id": str(site.get("id", "")),
                    "agents": total})
                add(_plan_row(acct_name, sname, "—", Total=total, kind="site",
                              Status=NOT_IN_FILE,
                              reason=f"The file has no row for '{sname}', so "
                                     f"its {total} agent(s) stay put."
                                     + _row_hint(sname, free)))
            continue
        dest = file_destination(entry)
        where = f"row {entry['row']} of the file"
        hits = reach.get(id(entry)) or []
        status = why = ""
        if len(hits) > 1:
            status = AMBIGUOUS
            why = (f"'{entry['source']}' on {where} matches {len(hits)} "
                   f"source sites ({'; '.join(hits[:4])}"
                   f"{' …' if len(hits) > 4 else ''}). Add a Source Account "
                   f"column to the file, or choose one source account.")
        elif entry["problem"]:
            status, why = FILE_PROBLEM, f"On {where}: {entry['problem']}."
        elif home and entry["console"].casefold() == home:
            status = SAME_CONSOLE
            why = (f"The token on {where} is for {console_host(home)} — the "
                   f"source console itself, so nothing would move.")
        if why:
            add(_plan_row(acct_name, sname, "—", Total=total, Status=status,
                          Destination=dest, reason=why))
            continue
        if not live:
            add(_plan_row(acct_name, sname, "—", state="done",
                          Status=NOTHING_TO_MOVE, Destination=dest,
                          reason=f"Matched {where}, but the source site has "
                                 f"no agents."))
            continue
        for grp in live:
            count = int(grp.get("totalAgents") or 0)
            add(_plan_row(acct_name, sname, grp.get("name", ""), state="ready",
                          Agents=count, Total=count,
                          Status="ready · token from the file",
                          token=entry["token"], token_from="file",
                          group_id=str(grp.get("id", "")), Destination=dest,
                          dest_site=entry["dest"],
                          reason=f"Site token from {where}."))

    if leftovers and split_ids(site_id) and not plan["stopped"]:
        plan["spare"] = leftovers
        say(f"{len(leftovers)} row(s) of the file are outside the chosen "
            f"sites and were not checked.")
    elif leftovers and not plan["stopped"]:
        for entry in leftovers:
            plan["unused"].append(entry)
            near, _same = _nearest(entry["source"], plan["unlisted"],
                                   lambda u: u["site"])
            hint = (f" The closest source site the file doesn't cover is "
                    f"'{near['site']}'." if near else "")
            add(_plan_row(entry["account"] or "—", entry["source"], "—",
                          kind="" if entry["problem"] else "file",
                          Status=NOT_ON_SOURCE,
                          Destination=file_destination(entry),
                          reason=f"Row {entry['row']}: no source site named "
                                 f"'{entry['source']}' in the "
                                 f"{plan['accounts']} account(s) read.{hint}"))

    retally(plan)
    if count_agents and not plan["stopped"]:
        say(f"Counting agents in {plan['ready']} matched group(s)…")
        refine_plan_counts(api, plan, on_row, stop)
    return plan


def pair_file_rows(pairs) -> int:
    """Point file rows at the source sites the operator picked for them.

    `pairs` is [(file row, unlisted source site)]. The row's source name
    becomes the site's and its account is pinned, so a pairing can never
    widen into a match in some other account. Where it lands is unchanged.
    """
    for entry, site in pairs:
        entry["source"] = site["site"]
        entry["account"] = site["account"]
        entry["paired"] = True
    return len(pairs)


def _closest_first(name, items, name_of) -> list:
    """`items` with the names closest to `name` first: the same name spelt
    differently ('shira-telecom' for 'Shira Telecom') on top."""
    key = loose_name(name)

    def rank(item):
        other = loose_name(name_of(item))
        return (other != key,
                -difflib.SequenceMatcher(None, key, other).ratio(), other)
    return sorted(items, key=rank)


def _nearest(name, items, name_of, cutoff: float = 0.6) -> tuple:
    """(item, same) — the item named closest to `name`, and whether it is
    the same name spelt differently — or (None, False) when none is close."""
    best = _closest_first(name, items, name_of)[:1]
    if not best:
        return None, False
    key, other = loose_name(name), loose_name(name_of(best[0]))
    if other == key:
        return best[0], True
    if difflib.SequenceMatcher(None, key, other).ratio() >= cutoff:
        return best[0], False
    return None, False


def rank_sites_by_name(name, sites) -> list:
    """`sites` with the names closest to `name` first."""
    return _closest_first(name, sites, lambda s: s.get("site"))


def rank_file_rows(name, rows) -> list:
    """The file's rows with the site names closest to `name` first."""
    return _closest_first(name, rows, lambda e: e.get("source"))


def _row_hint(site_name, rows) -> str:
    """What to tell the operator about a site the file doesn't name, given
    the rows of the file nothing uses."""
    entry, same = _nearest(site_name, rows, lambda e: e["source"])
    if same and norm_name(entry["source"]) == norm_name(site_name):
        return (f" Row {entry['row']} of the file names it for the account "
                f"'{entry['account']}' — choose it with Pair sites in step 2 "
                f"if it means this one.")
    if entry is not None and same:
        return (f" Row {entry['row']} of the file, '{entry['source']}', is "
                f"the same name spelt differently — choose it with Pair "
                f"sites in step 2.")
    if entry is not None:
        return (f" The closest row of the file nothing uses is "
                f"'{entry['source']}' (row {entry['row']}) — choose it with "
                f"Pair sites in step 2 if that's the one.")
    if rows:
        return (" If the file names it differently, choose its row with "
                "Pair sites in step 2.")
    return ""


# ── remembering the plan between sessions ──────────────────────────────
#
# Matching a 64-account console costs hundreds of round trips. The result
# only changes when the consoles or the map do, so it is cached against a
# fingerprint of exactly those inputs and restored on the next launch —
# step 3 is then usable straight away, with the plan's age on screen so a
# stale one is a visible choice rather than an accident.

PLAN_CACHE_NAME = "agent_match_plan.json"


def plan_cache_path() -> str:
    os.makedirs(config.CONFIG_DIR, exist_ok=True)
    return os.path.join(config.CONFIG_DIR, PLAN_CACHE_NAME)


def plan_fingerprint(source_url: str, account_ids: str, site_id: str,
                     map_path: str, map_data) -> dict:
    """What the plan depended on. Any change invalidates the cache."""
    blob = json.dumps(map_data or [], sort_keys=True, default=str)
    return {
        "source": str(source_url or ""),
        "accounts": " ".join(split_ids(account_ids)),
        "site": " ".join(split_ids(site_id)),
        "map_path": str(map_path or ""),
        "map_hash": hashlib.sha256(blob.encode("utf-8")).hexdigest()[:16],
    }


def save_plan(plan: dict, fingerprint: dict, path: str = "") -> str:
    """Cache a plan. Registration tokens are deliberately NOT written —
    they live in the map file, and are restored from it on load."""
    rows = []
    for row in plan.get("rows") or []:
        clean = {k: v for k, v in row.items() if k != "token"}
        rows.append(clean)
    payload = {"saved_at": datetime.now().isoformat(timespec="seconds"),
               "fingerprint": fingerprint,
               "plan": dict(plan, rows=rows)}
    target = path or plan_cache_path()
    with open(target, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2)
    return target


def load_plan(fingerprint: dict, map_data, path: str = "") -> tuple:
    """(plan, saved_at) for a cached plan that still matches, else (None, "").

    The tokens stripped by `save_plan` are re-read from the CURRENT map,
    so a rotated token is picked up and a scope that has since vanished
    from the map turns back into a blocked row instead of a bad move.
    """
    target = path or plan_cache_path()
    try:
        with open(target, encoding="utf-8") as f:
            payload = json.load(f)
    except (OSError, ValueError):
        return None, ""
    if payload.get("fingerprint") != fingerprint:
        return None, ""
    plan = payload.get("plan") or {}
    if not isinstance(plan.get("rows"), list):
        return None, ""
    for row in plan["rows"]:
        row.setdefault("token", "")
    rehydrate_tokens(plan, map_data)
    return plan, str(payload.get("saved_at") or "")


def rehydrate_tokens(plan: dict, map_data) -> dict:
    """Re-resolve each ready row's destination token from the map."""
    for row in plan.get("rows") or []:
        if row.get("state") != "ready":
            continue
        try:
            acct = find_by_name(map_data, row["dest_account"], "account")
            site = find_by_name(acct.get("sites"), row["dest_site"], "site")
            grp = find_by_name(site.get("groups"), row["dest_group"], "group")
        except MappingError:
            row.update(state="blocked", Status=NO_MATCH_GROUP, token="",
                       reason=f"'{row['dest_group']}' is no longer in the "
                              f"map file. Read the destination again.")
            continue
        token = token_for(grp, site)
        row["token"] = token
        if not token:
            row.update(state="blocked", Status=NO_TOKEN,
                       reason=f"The destination scope no longer has a "
                              f"registration token. Create one, then read "
                              f"the destination again.")
    return retally(plan)


def forget_plan(path: str = "") -> None:
    try:
        os.remove(path or plan_cache_path())
    except OSError:
        pass


def plan_age(saved_at: str) -> str:
    """'4 minutes ago' for a cached plan's timestamp."""
    try:
        when = datetime.fromisoformat(saved_at)
    except (TypeError, ValueError):
        return "earlier"
    secs = max(0, int((datetime.now() - when).total_seconds()))
    if secs < 90:
        return "just now"
    mins = secs // 60
    if mins < 60:
        return f"{mins} minute{'s' if mins != 1 else ''} ago"
    hours = mins // 60
    if hours < 24:
        return f"{hours} hour{'s' if hours != 1 else ''} ago"
    days = hours // 24
    return f"{days} day{'s' if days != 1 else ''} ago"


DEST_MAP_NAME = "agent_destination_map.json"


def dest_map_path() -> str:
    os.makedirs(config.CONFIG_DIR, exist_ok=True)
    return os.path.join(config.CONFIG_DIR, DEST_MAP_NAME)


def save_destination(map_data, destination: str, renames=(),
                     path: str = "") -> str:
    """Keep the last destination read, with the console it came from and
    the name fixes applied to it. Owner-only: it holds registration
    tokens."""
    target = path or dest_map_path()
    payload = {"destination": str(destination or ""),
               "saved_at": datetime.now().isoformat(timespec="seconds"),
               "renames": list(renames or []), "map": map_data or []}
    with open(target, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2)
    try:
        os.chmod(target, 0o600)
    except OSError:
        pass
    return target


def load_destination(path: str = "") -> dict:
    """What `save_destination` wrote, or {} if there is nothing usable."""
    try:
        with open(path or dest_map_path(), encoding="utf-8") as f:
            payload = json.load(f)
    except (OSError, ValueError):
        return {}
    if not isinstance(payload, dict) or not isinstance(payload.get("map"),
                                                       list):
        return {}
    return payload


def cached_plan_scope(path: str = "") -> dict:
    """The fingerprint the cached plan was saved under, or {}."""
    try:
        with open(path or plan_cache_path(), encoding="utf-8") as f:
            payload = json.load(f)
    except (OSError, ValueError):
        return {}
    fingerprint = payload.get("fingerprint") if isinstance(payload, dict) \
        else None
    return fingerprint if isinstance(fingerprint, dict) else {}


def plan_blockers(plan: dict) -> dict:
    """Blocked rows grouped by what is wrong, for the summary line."""
    out = {}
    for row in plan.get("rows") or []:
        if row["state"] == "blocked":
            out.setdefault(row["Status"], []).append(row)
    return out


# ── step 3: move them, and watch each one ──────────────────────────────

LIVE_COLUMNS = ["Computer", "OS", "Group", "Destination", "Result", "Detail"]

# consoleMigrationStatus → (what to show, why)
_SETTLED = {
    "Migrated": ("moved", "Confirmed on the destination console."),
    "Pending": ("pending", "Accepted — the agent moves when it next "
                           "checks in."),
    "Failed": ("failed", "The console reports the migration as failed. "
                         "Check the agent is online and the destination "
                         "token is still valid."),
    "N/A": ("pending", "The console has not picked the request up yet."),
}
_ACCEPTED = "Accepted by the console — not confirmed on the agent yet."


def mask_token(token: str) -> str:
    """Registration tokens are secrets and the live log is screenshotted,
    so only ever show enough to tell two tokens apart."""
    text = str(token or "")
    return f"…{text[-6:]}" if len(text) > 6 else ("…" if text else "(none)")


def verify_migration(api, agent_ids) -> dict:
    """agent id → the console's current consoleMigrationStatus.

    The bulk endpoint answers with a count, not per-agent results, so this
    read-back is the only way to tell "the console accepted 400 agents"
    from "400 agents actually moved".
    """
    out = {}
    for chunk in batched([str(a) for a in agent_ids], VERIFY_BATCH):
        try:
            rows = api.get_agents(params={"ids": ",".join(chunk)})
        except Exception:
            continue                      # a failed read-back is not a
        for a in rows:                    # failed migration
            out[str(a.get("id"))] = str(a.get("consoleMigrationStatus") or "")
    return out


def where_of(row: dict) -> str:
    return f"{row['Account']} / {row['Site']} / {row['Group']}"


def _live_record(agent: dict, row: dict) -> dict:
    aid = str(agent.get("id", ""))
    return {"id": aid,
            "Computer": agent.get("computerName", "") or aid,
            "OS": agent.get("osName", ""),
            "Account": row["Account"], "Site": row["Site"],
            "Group": row["Group"], "Destination": row["Destination"],
            "Result": "queued", "Detail": "", "Agent ID": aid,
            "offline": agent.get("isActive") is False}


def _settle_batch(api, records, batch, run, emit, verify) -> None:
    """Turn "the console accepted N" into a per-agent answer."""
    statuses = verify_migration(api, batch) if verify else {}
    for aid in batch:
        rec = records[aid]
        state, detail = _SETTLED.get(statuses.get(aid, ""),
                                     ("pending", _ACCEPTED))
        if state == "pending" and rec.get("offline"):
            detail = ("Accepted — this agent is offline, so it moves the "
                      "next time it connects.")
        rec["Result"], rec["Detail"] = state, detail
        run[state] += 1
        emit({"type": "agent", "rec": dict(rec)})


def run_live_migration(api, rows, on_event=None, should_stop=None,
                       skip_migrated: bool = True,
                       verify: bool = True) -> dict:
    """Move every ready row of a match plan, reporting each agent as it goes.

    `on_event` receives dicts describing the run as it happens — the group
    being worked, the request going out, and each agent's outcome. That is
    what the step-3 panel draws; pass None for an unattended run.
    """
    emit = on_event or (lambda _e: None)
    stop = should_stop or (lambda: False)
    ready = [r for r in rows if r.get("state") == "ready"]
    run = {"agents": [], "errors": [], "by_group": [], "groups": 0,
           "sent": 0, "affected": 0, "moved": 0, "pending": 0, "failed": 0,
           "stopped": False, "verified": bool(verify)}
    emit({"type": "start", "groups": len(ready),
          "agents": sum(int(r.get("Agents") or 0) for r in ready)})

    for index, row in enumerate(ready, 1):
        if stop():
            run["stopped"] = True
            break
        emit({"type": "group", "index": index, "total": len(ready),
              "row": row})
        try:
            agents = agents_for_group(api, row["group_id"], skip_migrated)
        except Exception as e:
            run["errors"].append(f"{where_of(row)}: {e}")
            emit({"type": "error", "where": where_of(row), "error": str(e)})
            continue

        run["groups"] += 1
        records = {}
        for agent in agents:
            rec = _live_record(agent, row)
            if rec["id"]:
                records[rec["id"]] = rec
                run["agents"].append(rec)
        # Copies: the consumer is on another thread and must see each
        # record as it was when the event fired, not as it ends up.
        emit({"type": "roster", "row": row,
              "recs": [dict(r) for r in records.values()]})

        tally = {"row": row, "candidates": len(records), "moved": 0,
                 "errors": []}
        chunks = list(batched(list(records), MOVE_BATCH))
        for n, batch in enumerate(chunks, 1):
            if stop():
                run["stopped"] = True
                break
            extra = {"groupIds": [str(row["group_id"])]}
            if skip_migrated:
                extra["consoleMigrationStatusesNin"] = ["Migrated"]
            for aid in batch:
                records[aid]["Result"] = "sending"
                emit({"type": "agent", "rec": dict(records[aid])})
            emit({"type": "request", "batch": n, "of": len(chunks),
                  "count": len(batch), "row": row,
                  "endpoint": "/agents/actions/move-to-console",
                  "token": mask_token(row["token"]), "filter": extra})
            try:
                res = api.move_agents_to_console(batch, row["token"], extra)
                affected = int(((res or {}).get("data") or {})
                               .get("affected") or 0)
            except Exception as e:
                detail = str(e)
                run["errors"].append(f"{where_of(row)}: {detail}")
                tally["errors"].append(detail)
                for aid in batch:
                    rec = records[aid]
                    rec["Result"], rec["Detail"] = "failed", detail
                    run["failed"] += 1
                    emit({"type": "agent", "rec": dict(rec)})
                emit({"type": "response", "ok": False, "batch": n,
                      "error": detail})
                continue
            run["sent"] += len(batch)
            run["affected"] += affected
            tally["moved"] += affected
            emit({"type": "response", "ok": True, "batch": n,
                  "affected": affected, "of": len(batch)})
            _settle_batch(api, records, batch, run, emit, verify)

        run["by_group"].append(tally)
        emit({"type": "group_done", "tally": tally})

    emit({"type": "done", "run": run})
    return run


# ── migration status report ────────────────────────────────────────────

def _agent_row(api, agent: dict, status: str,
               fetch_passphrases: bool = False) -> dict:
    row = {
        "Computer Name": agent.get("computerName", ""),
        "Migration Status": status,
        "Site": agent.get("siteName", ""),
        "Group": agent.get("groupName", ""),
        "OS": agent.get("osName", ""),
        "Last User": agent.get("lastLoggedInUserName", ""),
        "Decommissioned": "Yes" if agent.get("isDecommissioned") else "No",
        "Passphrase": "not fetched",
        "UUID": agent.get("uuid", ""),
        "Agent ID": str(agent.get("id", "")),
    }
    if fetch_passphrases and row["Agent ID"]:
        try:
            data = api.get_agent_passphrase(row["Agent ID"]) or []
            row["Passphrase"] = (data[0].get("passphrase", "")
                                 if data else "unavailable")
        except Exception as e:
            row["Passphrase"] = f"error: {e}"
    return row


def count_agents_by_status(api, account_ids: str = "", site_id: str = "",
                           fetch_passphrases: bool = False, log=None,
                           should_stop=None) -> dict:
    """Agents per console-migration status, plus decommissioned agents.

    `account_ids` may name several accounts (comma/space separated); blank
    with no site means every agent the token can see.
    """
    say = log or (lambda _m: None)
    stop = should_stop or (lambda: False)
    sites = split_ids(site_id)
    if sites:
        scope = {"siteIds": ",".join(sites)}
    else:
        ids = split_ids(account_ids)
        scope = {"accountIds": ",".join(ids)} if ids else {}

    counts = {s: 0 for s in MIGRATION_STATUSES}
    rows, decommissioned = [], []
    for status in MIGRATION_STATUSES:
        if stop():
            break
        params = dict(scope)
        params["consoleMigrationStatuses"] = status
        agents = api.get_agents(params=params)
        counts[status] = len(agents)
        say(f"{status}: {len(agents)} agent(s)")
        for agent in agents:
            rows.append(_agent_row(api, agent, status, fetch_passphrases))

    if not stop():
        params = dict(scope)
        params["isDecommissioned"] = "true"
        try:
            for agent in api.get_agents(params=params):
                decommissioned.append(
                    _agent_row(api, agent, "Decommissioned",
                               fetch_passphrases))
            say(f"Decommissioned: {len(decommissioned)} agent(s)")
        except Exception as e:
            say(f"Could not list decommissioned agents: {e}")

    return {"counts": counts, "agents": rows,
            "decommissioned": decommissioned}


# ── scheduling ─────────────────────────────────────────────────────────

def parse_schedule(date_str: str, time_str: str, now=None) -> datetime:
    """'YYYY-MM-DD' + 'HH:MM' → a datetime that is in the future."""
    text = f"{(date_str or '').strip()} {(time_str or '').strip()}"
    try:
        when = datetime.strptime(text, "%Y-%m-%d %H:%M")
    except ValueError:
        raise ValueError("Use date YYYY-MM-DD and time HH:MM (24-hour).")
    if when <= (now or datetime.now()):
        raise ValueError("That moment has already passed — pick a future "
                         "date and time.")
    return when


def format_countdown(delta: timedelta) -> str:
    total = int(max(0, delta.total_seconds()))
    hours, rem = divmod(total, 3600)
    minutes, seconds = divmod(rem, 60)
    days, hours = divmod(hours, 24)
    head = f"{days}d " if days else ""
    return f"{head}{hours:02d}:{minutes:02d}:{seconds:02d}"


def sites_without_token(map_data) -> list:
    """Sites nothing can be migrated into: no site token, and at least one
    group under them with no token of its own either.

    Worth surfacing before a run rather than as a per-group error during
    one — it's a destination-console fix, not something to retry.
    """
    stranded = []
    for acct in map_data or []:
        for site in acct.get("sites") or []:
            if site.get("token"):
                continue
            groups = site.get("groups") or []
            if not groups or any(not g.get("token") for g in groups):
                stranded.append((acct.get("name", ""), site.get("name", "")))
    return stranded


def plan_summary_text(plan: dict) -> str:
    """Plain-English outcome of one match pass."""
    if not plan:
        return "No match yet."
    lines = ["SOURCE → DESTINATION MATCH — nothing has moved.", "",
             f"Accounts read:      {plan.get('accounts', 0)}",
             f"Groups matched:     {plan.get('ready', 0)}",
             f"Agents to migrate:  {plan.get('agents', 0)}",
             f"Cannot migrate:     {plan.get('blocked', 0)}",
             f"Already migrated:   {plan.get('done', 0)}",
             f"Empty groups:       {plan.get('empty_groups', 0)}"]
    if plan.get("mode") == "file":
        lines[2:2] = [f"Tokens from:        {plan.get('file') or 'a file'}"]
        lines += [f"Sites not in file:  {len(plan.get('unlisted') or [])}",
                  f"File rows unused:   {len(plan.get('unused') or [])}"]
    if plan.get("stopped"):
        lines.append("Stopped early by the operator.")
    blocked = [r for r in plan.get("rows") or [] if r["state"] == "blocked"]
    if blocked:
        lines += ["", f"Will not migrate ({len(blocked)}):"]
        for row in blocked[:40]:
            lines.append(f"  • {where_of(row)} — {row['Status']}: "
                         f"{row['reason']}")
        if len(blocked) > 40:
            lines.append(f"  … and {len(blocked) - 40} more")
    return "\n".join(lines)


def live_summary_text(run: dict) -> str:
    """Plain-English outcome of one live migration, per agent."""
    if not run:
        return "No migration yet."
    lines = ["LIVE MIGRATION", "",
             f"Groups sent:        {run.get('groups', 0)}",
             f"Agents sent:        {run.get('sent', 0)}",
             f"Accepted by console:{run.get('affected', 0):>4}",
             f"Confirmed moved:    {run.get('moved', 0)}",
             f"Pending check-in:   {run.get('pending', 0)}",
             f"Failed:             {run.get('failed', 0)}"]
    if not run.get("verified"):
        lines.append("Per-agent confirmation was off, so 'moved' is what "
                     "the console accepted, not what arrived.")
    if run.get("stopped"):
        lines.append("Stopped early by the operator.")

    failures = [a for a in run.get("agents") or []
                if a.get("Result") == "failed"]
    if failures:
        lines += ["", f"Did not migrate ({len(failures)}):"]
        for agent in failures[:40]:
            lines.append(f"  • {agent['Computer']} "
                         f"({agent['Account']} / {agent['Site']} / "
                         f"{agent['Group']}): {agent['Detail']}")
        if len(failures) > 40:
            lines.append(f"  … and {len(failures) - 40} more")
    errors = run.get("errors") or []
    if errors:
        lines += ["", f"Errors ({len(errors)}):"]
        for err in errors[:20]:
            lines.append(f"  • {err}")
    return "\n".join(lines)


def status_summary_text(result: dict) -> str:
    """Plain-English outcome of one migration status report."""
    if not result:
        return "No report yet."
    counts = result.get("counts") or {}
    width = max([len(s) for s in MIGRATION_STATUSES] + [14])
    lines = ["Agents by console-migration status", ""]
    for status in MIGRATION_STATUSES:
        lines.append(f"  {status.ljust(width)}  {counts.get(status, 0)}")
    lines.append(f"  {'Decommissioned'.ljust(width)}  "
                 f"{len(result.get('decommissioned') or [])}")
    lines += ["", f"Total listed: {sum(counts.values())} agent(s)"]
    if any(r.get("Passphrase") not in ("", "not fetched")
           for r in result.get("agents") or []):
        lines += ["", "Passphrases are included in the export — treat that "
                      "file as sensitive."]
    return "\n".join(lines)


# ── rich report builders ───────────────────────────────────────────────
#
# Each export turns its result into the structured dict that
# export_utils.export_agent_report renders as a polished HTML document (or
# Excel / CSV / JSON). These are plain functions on plain dicts, so the
# report can be built and asserted without a console or a window.

def _stat(label, value, cls: str = "") -> dict:
    return {"label": label, "value": value, "class": cls}


def build_live_report(run: dict, source: str = "", map_path: str = "",
                      map_label: str = "Destination map file") -> dict:
    """Structured report for one live migration run."""
    run = run or {}
    agents = run.get("agents") or []
    moved, pending = run.get("moved", 0), run.get("pending", 0)
    failed = run.get("failed", 0)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    failures = [a for a in agents if a.get("Result") == "failed"]
    err_rows = [{"Error": e} for e in (run.get("errors") or [])]
    group_rows = [{
        "Group": where_of(t["row"]),
        "Destination": t["row"].get("Destination", ""),
        "Candidates": t.get("candidates", 0),
        "Accepted by console": t.get("moved", 0),
        "Errors": len(t.get("errors") or []),
    } for t in (run.get("by_group") or [])]

    sections = []
    if failures:
        sections.append({
            "title": "Did not migrate — manual action required",
            "icon": "✕", "color": "#ff6b81",
            "columns": ["Computer", "OS", "Group", "Destination", "Detail"],
            "rows": failures})
    if err_rows:
        sections.append({
            "title": "Errors", "icon": "⚠", "color": "#fdcb6e",
            "columns": ["Error"], "rows": err_rows})
    if group_rows:
        sections.append({
            "title": "Per group", "icon": "▦", "color": "#74b9ff",
            "columns": ["Group", "Destination", "Candidates",
                        "Accepted by console", "Errors"], "rows": group_rows})
    sections.append({
        "title": "Every agent", "icon": "🖥", "color": "#ffffff",
        "columns": LIVE_COLUMNS + ["Agent ID"], "rows": agents,
        "badge_cols": ["Result"], "empty": "No agents were processed."})

    note = None
    if run.get("stopped"):
        note = {"kind": "warn",
                "text": "The run was stopped early, so this report is partial."}
    elif not run.get("verified") and (moved or pending):
        note = {"kind": "info",
                "text": "Per-agent confirmation was off, so “moved” is what "
                        "the console accepted, not what was verified on each "
                        "agent. Run the Status report to confirm they arrived."}

    return {
        "title": "Agent Migration — Live Run", "icon": "🚚",
        "subtitle": "S1 Command Center — agents moved to the destination "
                    "console",
        "meta_line": f"Generated {now} • {len(agents)} agent(s) • "
                     f"{moved} moved • {pending} pending • "
                     f"{failed} failed",
        "info": [
            ("Source console", source or "—"),
            (map_label, map_path or "—"),
            ("Groups migrated", run.get("groups", 0)),
            ("Agents sent", run.get("sent", 0)),
            ("Accepted by the console", run.get("affected", 0)),
            ("Per-agent confirmation", "on" if run.get("verified") else "off"),
        ],
        "stats": [
            _stat("Moved", moved),
            _stat("Pending check-in", pending, "warn"),
            _stat("Failed", failed, "accent" if failed else "muted"),
            _stat("Groups", run.get("groups", 0), "blue"),
            _stat("Accepted", run.get("affected", 0), "blue"),
        ],
        "note": note, "sections": sections,
        "flat": {"columns": LIVE_COLUMNS + ["Agent ID"], "rows": agents},
    }


def build_status_report(result: dict, console: str = "", scope: str = "",
                        passphrases: bool = False) -> dict:
    """Structured report for one migration status run."""
    result = result or {}
    counts = result.get("counts") or {}
    agents = list(result.get("agents") or [])
    decommissioned = list(result.get("decommissioned") or [])
    all_rows = agents + decommissioned
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    has_pass = any(r.get("Passphrase") not in ("", "not fetched", None)
                   for r in all_rows)
    note = None
    if has_pass:
        note = {"kind": "danger",
                "text": "This report contains agent passphrases — treat the "
                        "file as sensitive, store it securely and delete it "
                        "when finished. Each passphrase was recorded in the "
                        "source console's activity log when it was fetched."}

    return {
        "title": "Agent Migration — Status Report", "icon": "📋",
        "subtitle": "S1 Command Center — agents by console-migration status",
        "meta_line": f"Generated {now} • {len(all_rows)} agent(s)",
        "info": [
            ("Console", console or "—"),
            ("Scope", scope or "whole console"),
            ("Passphrases included", "yes" if has_pass else "no"),
        ],
        "stats": [
            _stat("Migrated", counts.get("Migrated", 0)),
            _stat("Pending", counts.get("Pending", 0), "warn"),
            _stat("Failed", counts.get("Failed", 0),
                  "accent" if counts.get("Failed") else "muted"),
            _stat("N/A", counts.get("N/A", 0), "muted"),
            _stat("Decommissioned", len(decommissioned), "muted"),
        ],
        "note": note,
        "sections": [{
            "title": "Agents", "icon": "🖥", "color": "#ffffff",
            "columns": AGENT_REPORT_COLUMNS, "rows": all_rows,
            "badge_cols": ["Migration Status"],
            "empty": "No agents in that scope."}],
        "flat": {"columns": AGENT_REPORT_COLUMNS, "rows": all_rows},
    }


def build_match_report(plan: dict, map_path: str = "",
                       map_label: str = "") -> dict:
    """Structured report for one source→destination match pass."""
    plan = plan or {}
    rows = plan.get("rows") or []
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    def _plain(rs):
        return [{k: r.get(k, "") for k in PLAN_COLUMNS} for r in rs]

    def _with_reason(rs):
        return [{**{k: r.get(k, "") for k in PLAN_COLUMNS},
                 "Reason": r.get("reason", "")} for r in rs]

    ready = [r for r in rows if r.get("state") == "ready"]
    blocked = [r for r in rows if r.get("state") == "blocked"]
    done = [r for r in rows if r.get("state") == "done"]

    sections = [
        {"title": "Will migrate", "icon": "✓", "color": "#00e0a4",
         "columns": PLAN_COLUMNS, "rows": _plain(ready),
         "empty": "No group is ready to migrate."},
        {"title": "Will not migrate", "icon": "✕", "color": "#ff6b81",
         "columns": PLAN_COLUMNS + ["Reason"], "rows": _with_reason(blocked),
         "empty": "Nothing is blocked."},
    ]
    if done:
        sections.append(
            {"title": "Already migrated", "icon": "≡", "color": "#74b9ff",
             "columns": PLAN_COLUMNS, "rows": _plain(done)})

    file_mode = plan.get("mode") == "file"
    info = [(map_label or ("Site-token file" if file_mode
                           else "Destination map file"), map_path or "—"),
            ("Accounts read", plan.get("accounts", 0)),
            ("Agents to migrate", plan.get("agents", 0))]
    if file_mode:
        info += [("Source sites not in the file",
                  len(plan.get("unlisted") or [])),
                 ("File rows with no source site",
                  len(plan.get("unused") or []))]

    return {
        "title": "Agent Migration — Match Plan", "icon": "🧭",
        "subtitle": "S1 Command Center — source scopes matched to the "
                    "destination (nothing has moved)",
        "meta_line": f"Generated {now} • {len(rows)} scope(s)",
        "info": info,
        "stats": [
            _stat("Ready", plan.get("ready", 0)),
            _stat("Agents to migrate", plan.get("agents", 0), "blue"),
            _stat("Blocked", plan.get("blocked", 0),
                  "accent" if plan.get("blocked") else "muted"),
            _stat("Already migrated", plan.get("done", 0), "muted"),
            _stat("Empty groups", plan.get("empty_groups", 0), "muted"),
        ],
        "sections": sections,
        "flat": {"columns": PLAN_COLUMNS + ["State", "Reason"],
                 "rows": [{**{k: r.get(k, "") for k in PLAN_COLUMNS},
                           "State": r.get("state", ""),
                           "Reason": r.get("reason", "")} for r in rows]},
    }


# ═══════════════════════════════════════════════════════════════════════
#  UI building blocks
# ═══════════════════════════════════════════════════════════════════════

def _hex(colour) -> str:
    """One hex string from a CustomTkinter (light, dark) colour pair.

    Plain Tk widgets cannot take the pair, and the live feed rows have to
    be plain Tk: a CTkLabel queues canvas redraw work on every
    `configure`, and a few hundred of those is what froze the window.
    """
    if isinstance(colour, (tuple, list)):
        dark = str(ctk.get_appearance_mode()).lower() == "dark"
        return colour[1] if dark else colour[0]
    return str(colour)


class _ChoiceDialog(ctk.CTkToplevel):
    """Tick what you want from a list the console searches for us.

    Scope ids are long, opaque and easy to mistype, and nobody knows
    them by heart. Asking the console and showing names next to the ids
    turns "paste the right number" into a choice.

    Big consoles rule out the obvious implementation twice over, so this
    breaks both habits:

    * **We never ask for the whole console.** `fetch(query, limit)` runs
      on a worker thread and lets the API do the matching, so opening
      costs one small page instead of a full cursor walk. Consoles with
      thousands of sites used to mean waiting for every one of them.
    * **We never build a widget per result.** Rows come from a fixed
      pool of plain Tk widgets, rebound on each render. Growing a
      `CTkScrollableFrame` is superlinear — measured here at 2.3s for 50
      rows, 112s for 200 and 14 minutes for 400, because every insert
      relayouts the whole tree. A pool of `PAGE` is built once and each
      render after that is a handful of `configure` calls.

    Ticks survive searching: something ticked and then filtered away
    stays ticked, because `_chosen` is keyed by id and outlives the rows
    on screen. The footer says how many are ticked but out of view, so
    the count never looks wrong.

    Call `ask()`; it returns the chosen ids, or None if cancelled.
    """

    PAGE = 50            # rows in the pool, and the page we ask for
    DEBOUNCE_MS = 300    # quiet time before a keystroke hits the API

    def __init__(self, master, title, fetch, chosen=(), note="",
                 noun="item", source="the console"):
        super().__init__(master)
        self.title(title)
        self.geometry("560x560")
        self.minsize(440, 380)
        self._fetch = fetch
        self._noun = noun
        self._source = source
        # dict, not set: insertion order is the order we hand back
        self._chosen = dict.fromkeys(
            str(c) for c in chosen if str(c).strip())
        self._shown: list = []       # [(id, name, sub)] currently rendered
        self._more = False           # console has more than we asked for
        self._gen = 0                # newest search wins, stale ones drop
        self._timer = None
        self._result = None

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(3, weight=1)

        ctk.CTkLabel(self, text=title, anchor="w",
                     font=(UI_FONT, 15, "bold")).grid(
            row=0, column=0, sticky="ew", padx=16, pady=(14, 0))
        ctk.CTkLabel(self, text=note or "Search, then tick what you want.",
                     anchor="w", justify="left", wraplength=500,
                     font=(UI_FONT, 11), text_color=TEXT_MUTED).grid(
            row=1, column=0, sticky="ew", padx=16, pady=(2, 8))

        bar = ctk.CTkFrame(self, fg_color="transparent")
        bar.grid(row=2, column=0, sticky="ew", padx=16)
        bar.grid_columnconfigure(0, weight=1)
        self._search = ctk.CTkEntry(bar, height=32,
                                    placeholder_text="Search by name…")
        self._search.grid(row=0, column=0, sticky="ew")
        self._search.bind("<KeyRelease>", self._queue_search)
        ctk.CTkButton(bar, text="All", width=54, height=30, font=(UI_FONT, 11),
                      fg_color=NEUTRAL, hover_color=NEUTRAL_HOVER,
                      command=lambda: self._set_all(True)).grid(
            row=0, column=1, padx=(8, 0))
        ctk.CTkButton(bar, text="None", width=54, height=30,
                      font=(UI_FONT, 11), fg_color=NEUTRAL,
                      hover_color=NEUTRAL_HOVER,
                      command=lambda: self._set_all(False)).grid(
            row=0, column=2, padx=(6, 0))

        self._build_list(row=3)

        self._status = ctk.CTkLabel(self, text=f"Searching {source}…",
                                    anchor="w", font=(UI_FONT, 11),
                                    text_color=TEXT_FAINT)
        self._status.grid(row=4, column=0, sticky="ew", padx=18, pady=(0, 6))

        foot = ctk.CTkFrame(self, fg_color="transparent")
        foot.grid(row=5, column=0, sticky="ew", padx=16, pady=(0, 14))
        foot.grid_columnconfigure(0, weight=1)
        self._count = ctk.CTkLabel(foot, text="", anchor="w",
                                   font=(UI_FONT, 11),
                                   text_color=TEXT_MUTED)
        self._count.grid(row=0, column=0, sticky="w")
        ctk.CTkButton(foot, text="Cancel", width=90, height=34,
                      fg_color=NEUTRAL, hover_color=NEUTRAL_HOVER,
                      command=self._cancel).grid(row=0, column=1, padx=(0, 8))
        self._ok = ctk.CTkButton(foot, text="Use selection", width=140,
                                 height=34, font=(UI_FONT, 12, "bold"),
                                 fg_color=BRAND, hover_color=BRAND_HOVER,
                                 command=self._accept)
        self._ok.grid(row=0, column=2)

        self._tally()
        self.protocol("WM_DELETE_WINDOW", self._cancel)
        self.bind("<Escape>", lambda _e: self._cancel())
        self.bind("<Return>", lambda _e: self._accept())
        # Paint an empty window first, then go and ask. The operator sees
        # the dialog immediately even when the console is slow.
        self.after(0, self._search_now)

    # ── the list: one canvas, one pool, built once ────────────────────

    def _build_list(self, row):
        wrap = ctk.CTkFrame(self, fg_color=CARD_ELEVATED,
                            corner_radius=RADIUS_MD)
        wrap.grid(row=row, column=0, sticky="nsew", padx=16, pady=(10, 4))
        wrap.grid_columnconfigure(0, weight=1)
        wrap.grid_rowconfigure(0, weight=1)
        bg = _hex(CARD_ELEVATED)
        self._canvas = tk.Canvas(wrap, highlightthickness=0, bd=0, bg=bg)
        self._canvas.grid(row=0, column=0, sticky="nsew", padx=(6, 0), pady=6)
        scroll = ctk.CTkScrollbar(wrap, command=self._canvas.yview)
        scroll.grid(row=0, column=1, sticky="ns", padx=(2, 4), pady=6)
        self._canvas.configure(yscrollcommand=_queued_set(self, scroll))
        self._inner = tk.Frame(self._canvas, bg=bg)
        self._inner.grid_columnconfigure(0, weight=1)
        self._window = self._canvas.create_window(
            (0, 0), window=self._inner, anchor="nw")
        self._inner.bind("<Configure>", lambda _e: self._canvas.configure(
            scrollregion=self._canvas.bbox("all")))
        self._canvas.bind("<Configure>", lambda e: self._canvas.itemconfigure(
            self._window, width=e.width))
        self._rows = [self._make_row(i, bg) for i in range(self.PAGE)]
        self._blank = tk.Label(self._inner, text="", bg=bg,
                               fg=_hex(TEXT_FAINT), anchor="w",
                               justify="left", font=(UI_FONT, 12))
        self._wheel(self._canvas)
        self._wheel(self._inner)

    def _make_row(self, slot, bg):
        var = tk.BooleanVar(value=False)
        frame = tk.Frame(self._inner, bg=bg)
        frame.grid_columnconfigure(0, weight=1)
        chk = tk.Checkbutton(
            frame, variable=var, text="", anchor="w", justify="left",
            bg=bg, fg=_hex(TEXT), activebackground=bg,
            activeforeground=_hex(TEXT), selectcolor=_hex(CARD),
            highlightthickness=0, bd=0, padx=4, font=(UI_FONT, 12),
            command=lambda s=slot: self._toggle(s))
        chk.grid(row=0, column=0, sticky="ew", padx=(6, 6))
        sub = tk.Label(frame, text="", bg=bg, fg=_hex(TEXT_FAINT),
                       anchor="w", justify="left", font=(MONO_FONT, 10))
        sub.grid(row=1, column=0, sticky="ew", padx=(30, 6))
        for w in (frame, chk, sub):
            self._wheel(w)
        return {"frame": frame, "var": var, "chk": chk, "sub": sub,
                "id": None, "sig": None, "placed": False}

    def _wheel(self, widget):
        def on_wheel(event):
            num = getattr(event, "num", None)
            if num == 4:
                step = -1
            elif num == 5:
                step = 1
            else:
                delta = getattr(event, "delta", 0)
                step = -int(delta / 120) if abs(delta) >= 120 else -int(delta)
                step = step or (-1 if delta > 0 else 1)
            self._canvas.yview_scroll(step, "units")
            return "break"
        for seq in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            widget.bind(seq, on_wheel, add="+")

    def _render(self):
        """Rebind the pool. No widget is created or destroyed here."""
        rows = self._shown
        for i, row in enumerate(self._rows):
            if i >= len(rows):
                if row["placed"]:
                    row["frame"].grid_remove()
                    row["placed"] = False
                row["id"] = None
                continue
            ident, name, sub = rows[i]
            sig = (name, sub)
            if row["sig"] != sig:
                row["chk"].configure(text=name or ident)
                row["sub"].configure(text=sub or "")
                row["sig"] = sig
            row["id"] = ident
            # set() does not fire the Checkbutton command, so this cannot
            # loop back into _toggle and corrupt the selection
            row["var"].set(ident in self._chosen)
            if not row["placed"]:
                row["frame"].grid(row=i, column=0, sticky="ew", pady=1)
                row["placed"] = True
        if rows:
            self._blank.grid_remove()
        else:
            self._blank.configure(
                text=f"No {self._noun} matches that search."
                if self._search.get().strip()
                else f"No {self._noun} came back from {self._source}.")
            self._blank.grid(row=self.PAGE, column=0, sticky="ew",
                             padx=12, pady=12)
        self._canvas.yview_moveto(0)
        self._tally()

    # ── searching: debounced, off-thread, newest wins ─────────────────

    def _queue_search(self, _event=None):
        if self._timer is not None:
            try:
                self.after_cancel(self._timer)
            except Exception:
                pass
        self._timer = self.after(self.DEBOUNCE_MS, self._search_now)

    def _search_now(self):
        self._timer = None
        self._gen += 1
        gen = self._gen
        query = self._search.get().strip()
        self._status.configure(text=f"Searching {self._source}…")

        # Ask for one more than we can show: that is how we know whether
        # to tell the operator there is more behind the search.
        def work():
            items = self._fetch(query, self.PAGE + 1) or []
            return list(items), bool(getattr(items, "more", False))

        def done(payload):
            items, more = payload
            if gen != self._gen:
                return           # a later keystroke already answered
            self._more = more or len(items) > self.PAGE
            self._shown = items[:self.PAGE]
            self._render()
            self._say()

        def failed(exc):
            if gen != self._gen:
                return
            self._shown = []
            self._render()
            self._status.configure(text=f"Could not search: {exc}")

        self._spawn(work, done, failed)

    def _spawn(self, work, done, failed):
        """Run `work` off-thread and answer on the UI thread.

        Deliberately not `run_async`: this dialog can be dismissed while
        a search is still in flight, and scheduling onto a destroyed
        widget throws from inside the worker. Tests override this to run
        inline, which is also why the search logic never touches threads
        directly.
        """
        def run():
            try:
                payload, err = work(), None
            except Exception as exc:          # noqa: BLE001 - reported in UI
                payload, err = None, exc
            try:
                self.after(0, lambda: failed(err) if err else done(payload))
            except Exception:
                pass          # dialog closed while we were searching
        threading.Thread(target=run, daemon=True).start()

    def _say(self):
        n = len(self._shown)
        if not n:
            self._status.configure(text="")
        elif self._more:
            self._status.configure(
                text=f"Showing the first {n}. Type a name to narrow it "
                     f"down — the search covers all of {self._source}.")
        else:
            self._status.configure(
                text=f"Showing all {n} {self._noun}(s) that match.")

    def _set_all(self, value):
        """Applies to what is on screen, not to the whole console — with
        a search box in play, "All" must never tick what you filtered
        away and cannot see."""
        for ident, _n, _s in self._shown:
            if value:
                self._chosen[ident] = True
            else:
                self._chosen.pop(ident, None)
        self._render()

    def _toggle(self, slot):
        ident = self._rows[slot]["id"]
        if ident is None:
            return
        if self._rows[slot]["var"].get():
            self._chosen[ident] = True
        else:
            self._chosen.pop(ident, None)
        self._tally()

    def _tally(self):
        n = len(self._chosen)
        on_screen = sum(1 for i, _n, _s in self._shown if i in self._chosen)
        hidden = n - on_screen
        if not n:
            text = "Nothing ticked — that means everything"
        elif hidden:
            text = f"{n} ticked ({hidden} not shown by this search)"
        else:
            text = f"{n} ticked"
        self._count.configure(text=text)
        self._ok.configure(text="Use selection" if n else "Use all")

    # ── result ────────────────────────────────────────────────────────

    def _accept(self):
        self._result = list(self._chosen)
        self.destroy()

    def _cancel(self):
        self._result = None
        self.destroy()

    def ask(self):
        self._centre()
        self.after(60, self._grab)
        self.wait_window()
        return self._result

    def _centre(self):
        """Over the window that opened it, not the corner of the screen."""
        try:
            parent = self.master.winfo_toplevel()
            parent.update_idletasks()
            w, h = 560, 560
            x = parent.winfo_rootx() + (parent.winfo_width() - w) // 2
            y = parent.winfo_rooty() + (parent.winfo_height() - h) // 3
            self.geometry(f"{w}x{h}+{max(x, 0)}+{max(y, 0)}")
        except Exception:
            pass

    def _grab(self):
        try:
            self.lift()
            self.focus_force()
            self.grab_set()
            self._search.focus_set()
        except Exception:
            pass


class _Stats(ctk.CTkFrame):
    """A row of headline numbers. Figures are monospaced so columns line
    up and don't jitter as they update.

    Tiles are built once and then re-labelled. Destroying and rebuilding
    them on every update looks harmless, but creating CustomTkinter
    widgets from inside a redraw makes Tk re-enter `update_idletasks`,
    and during a live migration that cascade freezes the window.
    """

    def __init__(self, master, **kw):
        kw.setdefault("fg_color", "transparent")
        # Until the first tile exists this frame has no children, and an
        # empty CTkFrame asks for 200px — a band of nothing between the
        # buttons and the results. It grows to fit the tiles once shown.
        kw.setdefault("height", 1)
        super().__init__(master, **kw)
        self._tiles = []

    def _tile(self, index):
        while len(self._tiles) <= index:
            i = len(self._tiles)
            frame = ctk.CTkFrame(self, fg_color=CARD_ELEVATED,
                                 corner_radius=RADIUS_MD)
            value = ctk.CTkLabel(frame, text="",
                                 font=(MONO_FONT, 20, "bold"))
            value.pack(anchor="w", padx=14, pady=(10, 0))
            label = ctk.CTkLabel(frame, text="", font=(UI_FONT, 11),
                                 text_color=TEXT_MUTED)
            label.pack(anchor="w", padx=14, pady=(0, 10))
            frame.grid(row=0, column=i, padx=(0, 8), pady=2, sticky="ew")
            frame.grid_remove()
            self.grid_columnconfigure(i, weight=1)
            self._tiles.append({"frame": frame, "value": value,
                                "label": label, "sig": None})
        return self._tiles[index]

    def show(self, items):
        """items: [(value, label, colour)]"""
        for i, (value, label, colour) in enumerate(items):
            tile = self._tile(i)
            sig = (str(value), label, colour or TEXT)
            if tile["sig"] != sig:
                tile["sig"] = sig
                tile["value"].configure(text=sig[0], text_color=sig[2])
                tile["label"].configure(text=label)
            tile["frame"].grid()
        for tile in self._tiles[len(items):]:
            tile["frame"].grid_remove()

    def clear(self):
        for tile in self._tiles:
            tile["frame"].grid_remove()


class _EventPump:
    """Carry events from a worker thread to the widgets, in batches.

    Tk may only be touched from the main thread, and a live migration can
    emit thousands of events — redrawing on each one makes the window
    unusable. Everything lands in a queue and is drained on a timer.
    """

    def __init__(self, widget, handler, interval: int = 90,
                 per_tick: int = 2000, on_flush=None):
        self._widget = widget
        self._handler = handler
        self._interval = interval
        self._per_tick = per_tick
        self._on_flush = on_flush
        self._q = queue.Queue()
        self._job = None

    def push(self, event):
        """Called from the worker thread."""
        self._q.put(event)

    def start(self):
        if self._job is None:
            self._tick()

    def stop(self):
        if self._job is not None:
            try:
                self._widget.after_cancel(self._job)
            except Exception:
                pass
            self._job = None

    def drain(self):
        """Handle everything queued right now, then repaint once.

        The handler only updates state; drawing happens in `on_flush`, so
        a burst of ten thousand events costs one repaint instead of ten
        thousand. That is the difference between a live view and a frozen
        window.
        """
        seen = 0
        while seen < self._per_tick:
            try:
                event = self._q.get_nowait()
            except queue.Empty:
                break
            seen += 1
            try:
                self._handler(event)
            except Exception as exc:
                cli_log(f"live view: {exc}", "error")
        if seen and self._on_flush:
            try:
                self._on_flush()
            except Exception as exc:
                cli_log(f"live view: {exc}", "error")
        return seen

    def _tick(self):
        self.drain()
        self._job = self._widget.after(self._interval, self._tick)


class _CommandLog(ctk.CTkFrame):
    """The requests actually going out, as they go out.

    Operators asked to see the command, not a paraphrase of it — this is
    the endpoint, the filter and how many ids rode along. Tokens are
    masked because this pane ends up in screenshots.
    """

    MAX_LINES = 200

    def __init__(self, master, **kw):
        kw.setdefault("fg_color", CARD_ELEVATED)
        kw.setdefault("corner_radius", RADIUS_MD)
        super().__init__(master, **kw)
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)
        ctk.CTkLabel(self, text="API CALLS", anchor="w",
                     font=(UI_FONT, 10, "bold"), text_color=TEXT_FAINT).grid(
            row=0, column=0, sticky="ew", padx=12, pady=(8, 0))
        self._box = ctk.CTkTextbox(self, font=(MONO_FONT, 11), height=110,
                                   fg_color="transparent", wrap="none",
                                   activate_scrollbars=True)
        self._box.grid(row=1, column=0, sticky="nsew", padx=6, pady=(0, 6))
        self._box.configure(state="disabled")
        self._lines = 0
        self._pending = []

    def add(self, text):
        """Queue a line. Written on the next `flush`, because each write
        to the textbox costs a redraw of far more than the textbox."""
        self._pending.append(text)

    def flush(self):
        if not self._pending:
            return
        block, self._pending = "\n".join(self._pending) + "\n", []
        self._box.configure(state="normal")
        self._box.insert("end", block)
        self._lines += block.count("\n")
        if self._lines > self.MAX_LINES:
            cut = self._lines - self.MAX_LINES
            self._box.delete("1.0", f"{cut + 1}.0")
            self._lines = self.MAX_LINES
        self._box.see("end")
        self._box.configure(state="disabled")

    def clear(self):
        self._pending = []
        self._box.configure(state="normal")
        self._box.delete("1.0", "end")
        self._box.configure(state="disabled")
        self._lines = 0


class _Form(ctk.CTkFrame):
    """A labelled form card.

    Every field gets a visible label AND persistent helper text underneath
    — placeholder-only labels and hover-only "?" tooltips both hide the
    thing the operator most needs to know (what blank means).
    """

    BUTTON_W = 88        # so fields with and without one line up

    def __init__(self, master, **kw):
        kw.setdefault("fg_color", CARD)
        kw.setdefault("corner_radius", RADIUS_MD)
        super().__init__(master, **kw)
        # The input takes a quarter of any spare width and the trailing
        # column soaks up the rest. Letting it take all of it strands the
        # label at one edge of a wide window and the button at the other.
        self.grid_columnconfigure(1, weight=1, minsize=240)
        self.grid_columnconfigure(3, weight=3)
        ctk.CTkFrame(self, fg_color="transparent", width=1, height=1).grid(
            row=0, column=3)
        self._row = 0

    def field(self, label, placeholder="", helper="", button=None, show=None):
        r = self._row
        ctk.CTkLabel(self, text=label, anchor="w",
                     font=(UI_FONT, 12, "bold")).grid(
            row=r, column=0, sticky="w", padx=(14, 10), pady=(12, 0))
        entry = ctk.CTkEntry(self, placeholder_text=placeholder, height=34)
        if show:
            entry.configure(show=show)
        entry.grid(row=r, column=1, sticky="ew", pady=(12, 0))
        if button:
            text, cmd = button
            ctk.CTkButton(self, text=text, width=self.BUTTON_W, height=32,
                          fg_color=NEUTRAL, hover_color=NEUTRAL_HOVER,
                          font=(UI_FONT, 12), command=cmd).grid(
                row=r, column=2, padx=(8, 14), pady=(12, 0))
        else:
            # Height 1, because an empty CTkFrame defaults to 200px tall
            # and would stretch this row — leaving the helper text
            # stranded half a panel below the box it describes.
            ctk.CTkFrame(self, fg_color="transparent",
                         width=self.BUTTON_W, height=1).grid(
                row=r, column=2, padx=(8, 14))
        if helper:
            ctk.CTkLabel(self, text=helper, anchor="w", justify="left",
                         font=(UI_FONT, 11), text_color=TEXT_FAINT).grid(
                row=r + 1, column=1, columnspan=2, sticky="w", pady=(3, 0))
        self._row += 2
        return entry

    def block(self, widget, pad=(10, 12)):
        widget.grid(row=self._row, column=0, columnspan=3, sticky="ew",
                    padx=14, pady=pad)
        self._row += 1
        return widget

    def pad(self, height=10):
        ctk.CTkFrame(self, fg_color="transparent",
                     width=1, height=height).grid(row=self._row, column=0)
        self._row += 1


# ═══════════════════════════════════════════════════════════════════════
#  The page's building blocks
# ═══════════════════════════════════════════════════════════════════════

_UNSET = object()


def _retext(widget, **opts):
    """configure() only the options that changed: every configure() on a
    visible CustomTkinter widget redraws it."""
    cache = getattr(widget, "_retext_cache", None) or {}
    todo = {}
    for key, value in opts.items():
        try:
            current = widget.cget(key)
        except Exception:
            current = cache.get(key, _UNSET)
        if current != value:
            todo[key] = value
    if todo:
        widget.configure(**todo)
        widget._retext_cache = {**cache, **todo}


def _grid_shown(widget, shown):
    """grid() or grid_remove() a widget that has been gridded once."""
    if shown:
        if not widget.winfo_manager():
            widget.grid()
    elif widget.winfo_manager():
        widget.grid_remove()


def _ghost(parent, text, command, width=110, height=32, **kw):
    """A secondary action — a quiet outline, so the one filled button on
    the page is always the next thing to do."""
    opts = dict(text=text, command=command, width=width, height=height,
                font=(UI_FONT, 12), fg_color=GHOST, hover_color=GHOST_HOVER,
                text_color=TEXT, text_color_disabled=TEXT_FAINT,
                border_width=1, border_color=BORDER, corner_radius=RADIUS_MD)
    opts.update(kw)
    return ctk.CTkButton(parent, **opts)


_PRIMARY = dict(fg_color=BRAND, hover_color=BRAND_HOVER, text_color="#FFFFFF",
                border_width=0)
_GHOSTED = dict(fg_color=GHOST, hover_color=GHOST_HOVER, text_color=TEXT,
                border_width=1)
_GO = dict(fg_color=GREEN, hover_color=GREEN_HOVER, text_color="#FFFFFF",
           border_width=0)


def _wrap(frame, label, pad=8):
    """Keep `label` wrapped to the width of `frame` as the window resizes."""
    def resize(event):
        try:
            scale = ctk.ScalingTracker.get_widget_scaling(label)
        except Exception:
            scale = 1.0
        width = max(240, int(event.width / scale) - pad)
        last = getattr(label, "_wrap_px", None)
        if last is None or width < last or width - last > 2:
            label._wrap_px = width
            label.configure(wraplength=width)
    frame.bind("<Configure>", resize, add="+")


def _wheel_rows(event) -> int:
    """Rows to scroll for one mouse-wheel event, on any platform."""
    num = getattr(event, "num", None)
    if num in (4, 5):
        return -3 if num == 4 else 3
    delta = int(getattr(event, "delta", 0) or 0)
    if abs(delta) >= 120:
        return -3 * int(delta / 120)
    return -delta


def _queued_set(widget, bar):
    state = {"pos": None, "job": None}

    def apply():
        state["job"] = None
        try:
            bar.set(*state["pos"])
        except tk.TclError:
            pass

    def queue(first, last):
        state["pos"] = (first, last)
        if state["job"] is None:
            state["job"] = widget.after(15, apply)
    return queue


def _short(text, limit: int) -> str:
    """`text`, cut to `limit` characters with an ellipsis."""
    text = str(text or "")
    return text if len(text) <= limit else text[:max(1, limit - 1)] + "…"


def _col(title, key, weight=1, anchor="w", mono=False, bold=False,
         minsize=0) -> dict:
    return {"title": title, "key": key, "weight": weight, "anchor": anchor,
            "mono": mono, "bold": bold, "minsize": minsize}


def scope_path(row) -> str:
    """'Acme  ›  HQ  ›  Servers', leaving out the parts a row doesn't have."""
    return "  ›  ".join(str(p) for p in (row.get("Account"), row.get("Site"),
                                          row.get("Group"))
                        if p and p != "—")


def plan_note(row) -> str:
    """The short 'why' — or 'with which token' — beside a plan row."""
    if row.get("state") != "ready" or row.get("token_from") == "file":
        return str(row.get("reason") or "")
    return f"with the destination {row.get('token_from') or 'group'}'s token"


class _StepCard(ctk.CTkFrame):
    """One numbered step: a badge saying where it stands, its title and
    caption down the left, its controls in `body`."""

    # state → (badge fill, glyph or None for the number, glyph colour,
    #          card edge, title colour)
    LOOK = {
        "todo": (CARD_ELEVATED, None, TEXT_MUTED, BORDER, TEXT),
        "active": (BRAND, None, "#FFFFFF", BRAND, TEXT),
        "done": (GREEN, "✓", "#FFFFFF", BORDER, TEXT),
        "warn": (WARN, "!", "#FFFFFF", WARN, TEXT),
        "error": (ACCENT, "!", "#FFFFFF", ACCENT, TEXT),
        "locked": (CARD_ELEVATED, None, TEXT_FAINT, BORDER, TEXT_FAINT),
    }

    def __init__(self, master, number, title, caption, **kw):
        kw.setdefault("fg_color", CARD)
        kw.setdefault("corner_radius", RADIUS_LG)
        kw.setdefault("border_width", 1)
        kw.setdefault("border_color", BORDER)
        super().__init__(master, **kw)
        self.number = number
        self.step_state = None
        self._painted = None
        self.grid_columnconfigure(0, minsize=176)
        self.grid_columnconfigure(1, weight=1)
        rail = ctk.CTkFrame(self, fg_color="transparent")
        rail.grid(row=0, column=0, sticky="nw", padx=(16, 8), pady=12)
        self._badge = ctk.CTkFrame(rail, width=28, height=28, corner_radius=14,
                                   fg_color=CARD_ELEVATED)
        self._badge.grid(row=0, column=0, rowspan=2, sticky="n", padx=(0, 10))
        self._glyph = ctk.CTkLabel(self._badge, text=str(number), width=20,
                                   height=20, fg_color="transparent",
                                   text_color=TEXT_MUTED,
                                   font=(UI_FONT, 13, "bold"))
        self._glyph.place(relx=0.5, rely=0.5, anchor="center")
        self._title = ctk.CTkLabel(rail, text=title, anchor="w", height=20,
                                   font=(UI_FONT, 14, "bold"))
        self._title.grid(row=0, column=1, sticky="w")
        self._caption = ctk.CTkLabel(rail, text=caption, anchor="w", height=16,
                                     font=(UI_FONT, 11), text_color=TEXT_FAINT)
        self._caption.grid(row=1, column=1, sticky="w")
        self.body = ctk.CTkFrame(self, fg_color="transparent")
        self.body.grid(row=0, column=1, sticky="ew", padx=(0, 16), pady=10)
        self.body.grid_columnconfigure(0, weight=1)
        self.set("todo")

    def set(self, state, caption=None):
        if caption is not None:
            _retext(self._caption, text=caption)
        self.step_state = state
        if state == self._painted:
            return
        self._painted = state
        fill, glyph, ink, edge, title = self.LOOK[state]
        self._badge.configure(fg_color=fill)
        self._glyph.configure(text=glyph or str(self.number), text_color=ink)
        self.configure(border_color=edge)
        self._title.configure(text_color=title)


class _VirtualTable(ctk.CTkFrame):
    """Any number of rows, drawn by a fixed pool of plain Tk rows.

    A CustomTkinter widget per cell is what froze the old lists: each one
    paints its own canvas and a CTkScrollableFrame relayouts the whole tree
    on every insert — 32 s to show 100 plan rows. Here the pool is built
    once, placed by hand and re-labelled as the view scrolls, so 50,000
    rows cost what 50 do. Records are plain dicts; `put` only stores them
    and `render` draws.
    """

    VISIBLE = 60

    def __init__(self, master, columns, empty="", on_select=None,
                 colour=None, **kw):
        kw.setdefault("fg_color", CARD_ELEVATED)
        kw.setdefault("corner_radius", RADIUS_MD)
        super().__init__(master, **kw)
        self.columns = [c if isinstance(c, dict) else _col(*c)
                        for c in columns]
        self._colour = colour
        self._on_select = on_select
        self._empty = empty
        self._records = {}
        self._order = []
        self._view = []
        self._predicate = None
        self._dirty = False
        self._offset = 0
        self._fit = 1
        self._follow = self._tail = False
        self._selected = None
        self._bar_sig = None
        self._hexes = {}
        self._c = {}
        self._resolve()
        self._apply_scale()
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)

        self._head = tk.Frame(self, bg=self._c["bg"], bd=0,
                              highlightthickness=0)
        self._head.grid(row=0, column=0, sticky="ew", padx=(6, 0),
                        pady=(8, 2))
        self._heads = []
        for j, col in enumerate(self.columns):
            self._layout(self._head, j, col)
            lab = tk.Label(self._head, text=col["title"].upper(),
                           anchor=col["anchor"], bg=self._c["bg"],
                           fg=self._c["faint"], font=self._head_font,
                           width=1, bd=0, highlightthickness=0)
            lab.grid(row=0, column=j, sticky="ew", padx=4)
            self._heads.append(lab)
        self._body = tk.Frame(self, bg=self._c["bg"], bd=0,
                              highlightthickness=0)
        self._body.grid(row=1, column=0, sticky="nsew", padx=(6, 0),
                        pady=(0, 6))
        self._bar = ctk.CTkScrollbar(self, command=self._yview)
        self._bar.grid(row=1, column=1, sticky="ns", padx=(2, 4), pady=(0, 6))
        self._set_bar = _queued_set(self, self._bar)
        self._pool = [self._make_row(i) for i in range(self.VISIBLE)]
        self._note = tk.Label(self._body, text=empty, anchor="nw",
                              justify="left", bg=self._c["bg"],
                              fg=self._c["faint"],
                              font=self._note_font, bd=0)
        self._body.bind("<Configure>", self._on_resize)
        self._wheel(self._body)
        self._wheel(self._note)
        self.render()
        tk_track(self, _VirtualTable._retheme)

    # ── the pool ──────────────────────────────────────────────────────

    def _apply_scale(self, scale=None):
        """Fonts and row height in pixels, scaled the way CustomTkinter
        scales its own widgets, so the list reads like the rest of the page."""
        if scale is None:
            try:
                scale = ctk.ScalingTracker.get_widget_scaling(self)
            except Exception:
                scale = 1.0
        self._scale = scale

        def px(size):
            return -max(8, round(size * scale))
        self._fonts = {(m, b): ((MONO_FONT, px(11)) if m
                                else (UI_FONT, px(12)))
                       + (("bold",) if b else ())
                       for m in (False, True) for b in (False, True)}
        self._head_font = (UI_FONT, px(10), "bold")
        self._note_font = (UI_FONT, px(12))
        self._row_h = tkfont.Font(root=self, family=UI_FONT,
                                  size=px(12)).metrics("linespace") \
            + round(10 * scale)

    def _set_scaling(self, *args, **kwargs):
        super()._set_scaling(*args, **kwargs)
        self._apply_scale(args[0] if args else None)
        for j, (lab, col) in enumerate(zip(self._heads, self.columns)):
            self._layout(self._head, j, col)
            lab.configure(font=self._head_font)
        self._note.configure(font=self._note_font)
        for row in self._pool:
            for j, (lab, col) in enumerate(zip(row["cells"], self.columns)):
                self._layout(row["frame"], j, col)
                lab.configure(font=self._fonts[(col["mono"], col["bold"])])
            if row["shown"]:
                row["frame"].place_forget()
                row["shown"] = False
            row["sig"] = None
        height = self._body.winfo_height()
        if height > 1:
            self._fit = max(1, min(self.VISIBLE, height // self._row_h))
        self.render()

    def _layout(self, frame, j, col):
        minsize = round(col["minsize"] * self._scale)
        if col["weight"]:
            frame.grid_columnconfigure(j, weight=col["weight"], uniform="cols",
                                       minsize=minsize)
        else:
            frame.grid_columnconfigure(j, weight=0, minsize=minsize)

    def _make_row(self, slot):
        frame = tk.Frame(self._body, bg=self._c["bg"], bd=0,
                         highlightthickness=0)
        frame.grid_rowconfigure(0, weight=1)
        cells = []
        for j, col in enumerate(self.columns):
            self._layout(frame, j, col)
            lab = tk.Label(frame, text="", anchor=col["anchor"],
                           bg=self._c["bg"], fg=self._c["text"],
                           font=self._fonts[(col["mono"], col["bold"])],
                           width=1, bd=0, highlightthickness=0, padx=0)
            lab.grid(row=0, column=j, sticky="nsew", padx=4)
            cells.append(lab)
        for widget in [frame] + cells:
            widget.bind("<Button-1>", lambda _e, s=slot: self._click(s))
            self._wheel(widget)
        return {"frame": frame, "cells": cells, "key": None, "sig": None,
                "shown": False}

    def _resolve(self):
        self._hexes = {}
        self._c = {"bg": _hex(CARD_ELEVATED), "stripe": _hex(CARD),
                   "pick": _hex(SIDEBAR_SEL), "text": _hex(TEXT),
                   "faint": _hex(TEXT_FAINT)}

    def _hexof(self, token):
        key = token if isinstance(token, str) else tuple(token)
        hit = self._hexes.get(key)
        if hit is None:
            hit = self._hexes[key] = _hex(token)
        return hit

    def _retheme(self):
        self._resolve()
        for widget in (self._head, self._body):
            widget.configure(bg=self._c["bg"])
        self._note.configure(bg=self._c["bg"], fg=self._c["faint"])
        for lab in self._heads:
            lab.configure(bg=self._c["bg"], fg=self._c["faint"])
        for row in self._pool:
            row["sig"] = None
        self.render()

    # ── data in ───────────────────────────────────────────────────────

    def clear(self, message=None):
        if message is not None:
            self._empty = message
        self._records, self._order, self._view = {}, [], []
        self._dirty = False
        self._offset = 0
        self._selected = None
        self.render()

    def put(self, rec, key=None):
        """Record one row, new or changed. Draws nothing — see render()."""
        k = key if key is not None else (rec.get("key") or rec.get("id"))
        if k is None:
            return
        if k not in self._records:
            self._order.append(k)
            if self._predicate is None and not self._dirty:
                self._view.append(k)
            else:
                self._dirty = True
        elif self._predicate is not None:
            self._dirty = True
        self._records[k] = rec

    def set_rows(self, rows):
        self.clear()
        for i, row in enumerate(rows):
            k = row.get("key") or row.get("id")
            self.put(row, key=i if k is None else k)
        self.render()

    def filter_by(self, predicate):
        self._predicate = predicate
        self._dirty = True
        self._offset = 0
        self.render()

    def set_follow(self, follow: bool):
        """Stick to the newest row while a run is live."""
        self._tail = self._follow = bool(follow)

    def counts(self) -> tuple:
        """(rows the filter lets through, all rows)"""
        if self._dirty:
            self._rebuild()
        return len(self._view), len(self._order)

    def records(self) -> list:
        return [self._records[k] for k in self._order]

    def _rebuild(self):
        test = self._predicate
        self._view = [k for k in self._order
                      if test is None or test(self._records[k])]
        self._dirty = False

    # ── drawing ───────────────────────────────────────────────────────

    def cell(self, rec, key):
        """(text, colour) for one cell — subclasses decide."""
        value = rec.get(key, "")
        colour = self._colour(rec) if self._colour else None
        return ("" if value is None else str(value)), colour or TEXT

    def render(self):
        if self._dirty:
            self._rebuild()
        total, fit = len(self._view), self._fit
        top = max(0, total - fit)
        if self._follow:
            self._offset = top
        self._offset = min(max(0, self._offset), top)
        shown = self._view[self._offset:self._offset + fit]
        for slot, row in enumerate(self._pool):
            if slot >= len(shown):
                if row["shown"]:
                    row["frame"].place_forget()
                    row["shown"] = False
                row["key"] = None
                continue
            key = shown[slot]
            rec = self._records[key]
            index = self._offset + slot
            bg = (self._c["pick"] if key == self._selected
                  else self._c["stripe"] if index % 2 else self._c["bg"])
            cells = tuple((str(text), self._hexof(colour)) for text, colour in
                          (self.cell(rec, col["key"]) for col in self.columns))
            sig = (bg, cells)
            row["key"] = key
            if row["sig"] != sig:
                row["sig"] = sig
                row["frame"].configure(bg=bg)
                for lab, (text, fg) in zip(row["cells"], cells):
                    lab.configure(text=text, fg=fg, bg=bg)
            if not row["shown"]:
                row["frame"].place(x=0, y=slot * self._row_h, relwidth=1,
                                   height=self._row_h)
                row["shown"] = True
        if total:
            if self._note.winfo_manager():
                self._note.place_forget()
        else:
            text = self._empty if not self._order \
                else "Nothing matches this filter."
            if self._note.cget("text") != text:
                self._note.configure(text=text)
            if not self._note.winfo_manager():
                self._note.place(x=8, y=10, relwidth=1, width=-16)
        span = ((self._offset / total, (self._offset + len(shown)) / total)
                if total > fit else (0.0, 1.0))
        if self._bar_sig != span:
            self._bar_sig = span
            self._set_bar(*span)

    # ── scrolling and picking ─────────────────────────────────────────

    def _on_resize(self, event):
        self._note.configure(wraplength=max(200, int(event.width) - 24))
        fit = max(1, min(self.VISIBLE, int(event.height) // self._row_h))
        if fit != self._fit:
            self._fit = fit
            self.render()

    def scroll(self, rows):
        if not rows:
            return
        total = len(self._view)
        self._offset = min(max(0, self._offset + rows),
                           max(0, total - self._fit))
        self._follow = self._tail and self._offset >= total - self._fit
        self.render()

    def _yview(self, *args):
        total = len(self._view)
        if not args or not total:
            return
        if args[0] == "moveto":
            self.scroll(int(round(float(args[1]) * total)) - self._offset)
        elif args[0] == "scroll":
            step = int(float(args[1]))
            pages = len(args) > 2 and args[2] == "pages"
            self.scroll(step * self._fit if pages else step)

    def _wheel(self, widget):
        for seq in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            widget.bind(seq, self._on_wheel, add="+")

    def _on_wheel(self, event):
        self.scroll(_wheel_rows(event))
        return "break"

    def _click(self, slot):
        key = self._pool[slot]["key"]
        if key is None:
            return
        self._selected = key
        self.render()
        if self._on_select:
            self._on_select(self._records[key])


class _PlanList(_VirtualTable):
    """Every source scope beside where it lands — any number of them."""

    MARK = {"ready": ("✓", GREEN), "done": ("≡", TEXT_FAINT),
            "blocked": ("✕", ACCENT)}

    def __init__(self, master, **kw):
        super().__init__(master, [
            _col("", "_mark", 0, bold=True, minsize=26),
            _col("Source", "_source", 4),
            _col("Lands in", "_dest", 4),
            _col("Agents", "_count", 1, anchor="e", mono=True),
            _col("Details", "_note", 5)], **kw)

    def set_filter(self, value):
        """all | ready | blocked | done"""
        self.filter_by(None if value == "all"
                       else (lambda r, v=value: r.get("state") == v))

    def cell(self, row, key):
        state = row.get("state")
        if key == "_mark":
            return self.MARK.get(state, ("•", TEXT_MUTED))
        if key == "_source":
            return scope_path(row), TEXT if state != "done" else TEXT_MUTED
        if key == "_dest":
            if state == "blocked":
                return row.get("Status", ""), ACCENT
            return (row.get("Destination", ""),
                    TEXT if state == "ready" else TEXT_MUTED)
        if key == "_count":
            if state == "ready":
                n = f"{int(row.get('Agents') or 0):,}"
                return (n, TEXT) if row.get("counted") else (f"≈{n}",
                                                             TEXT_FAINT)
            total = int(row.get("Total") or 0)
            return (f"{total:,}" if total else ""), TEXT_FAINT
        return plan_note(row), TEXT_MUTED if state == "blocked" else TEXT_FAINT


class _LiveFeed(_VirtualTable):
    """Every agent in the run, with what the console said about it."""

    PILL = {
        "queued": ("queued", TEXT_FAINT),
        "sending": ("sending…", BRAND),
        "moved": ("moved", GREEN),
        "pending": ("pending", WARN),
        "failed": ("FAILED", ACCENT),
    }

    def __init__(self, master, **kw):
        super().__init__(master, [
            _col("Result", "Result", 0, bold=True, minsize=92),
            _col("Computer", "Computer", 3, mono=True),
            _col("OS", "OS", 2),
            _col("Site  ›  Group", "_where", 3),
            _col("Details", "Detail", 5)], **kw)

    def set_filter(self, value):
        """all | moved | pending | failed"""
        self.filter_by(None if value == "all"
                       else (lambda r, v=value: r.get("Result") == v))

    def cell(self, rec, key):
        state = rec.get("Result", "queued")
        if key == "Result":
            return self.PILL.get(state, (state, TEXT_MUTED))
        if key == "_where":
            return (f"{rec.get('Site', '')}  ›  {rec.get('Group', '')}",
                    TEXT_MUTED)
        if key == "Detail":
            return (rec.get("Detail", ""),
                    ACCENT if state == "failed" else TEXT_FAINT)
        if key == "OS":
            return rec.get("OS", ""), TEXT_MUTED
        return rec.get(key, ""), TEXT


class _TableDialog(ctk.CTkToplevel):
    """A read-only list in its own window — the rows of a loaded file."""

    def __init__(self, master, title, note, columns, rows, colour=None):
        super().__init__(master)
        self.title(title)
        self.geometry("880x520")
        self.minsize(560, 320)
        ctk.CTkLabel(self, text=note, anchor="w", justify="left",
                     wraplength=820, font=(UI_FONT, 12),
                     text_color=TEXT_MUTED).pack(fill="x", padx=16,
                                                 pady=(14, 8))
        table = _VirtualTable(self, columns, colour=colour,
                              empty="Nothing to show.")
        table.pack(fill="both", expand=True, padx=16, pady=(0, 10))
        table.set_rows(rows)
        _ghost(self, "Close", self.destroy, width=100).pack(
            anchor="e", padx=16, pady=(0, 14))
        self.bind("<Escape>", lambda _e: self.destroy())
        self.transient(master.winfo_toplevel())
        self.after(120, self._raise)

    def _raise(self):
        try:
            self.lift()
            self.focus_force()
        except Exception:
            pass


class _ScheduleDialog(ctk.CTkToplevel):
    """When the unattended run should fire."""

    def __init__(self, page):
        super().__init__(page)
        self._page = page
        self.title("Schedule the migration")
        self.geometry("540x350")
        self.resizable(False, False)
        ctk.CTkLabel(self, text="Schedule the migration", anchor="w",
                     font=(UI_FONT, 15, "bold")).pack(fill="x", padx=18,
                                                      pady=(16, 2))
        ctk.CTkLabel(
            self, anchor="w", justify="left", wraplength=500,
            font=(UI_FONT, 12), text_color=TEXT_MUTED,
            text="Runs step 3 unattended at the time below, by this "
                 "computer's clock, and moves exactly what the current plan "
                 "lists. Keep the app open, the machine awake and on VPN "
                 "until it fires.").pack(fill="x", padx=18, pady=(0, 10))
        form = _Form(self)
        form.pack(fill="x", padx=16)
        soon = datetime.now() + timedelta(hours=1)
        self.date = form.field("Date", "YYYY-MM-DD",
                               "Your computer's local date.")
        self.date.insert(0, soon.strftime("%Y-%m-%d"))
        self.time = form.field("Time", "HH:MM (24-hour)",
                               "Your computer's local time.")
        self.time.insert(0, soon.strftime("%H:%M"))
        form.pad(8)
        buttons = ctk.CTkFrame(self, fg_color="transparent")
        buttons.pack(fill="x", padx=16, pady=14)
        ctk.CTkButton(buttons, text="⏱  Arm", width=120, height=34,
                      font=(UI_FONT, 13, "bold"), fg_color=ACCENT,
                      hover_color=ACCENT_HOVER,
                      command=self._arm).pack(side="right")
        _ghost(buttons, "Cancel", self.destroy, width=100, height=34).pack(
            side="right", padx=8)
        self.bind("<Escape>", lambda _e: self.destroy())
        self.transient(page.winfo_toplevel())
        self.after(120, self._grab)

    def _grab(self):
        try:
            self.lift()
            self.focus_force()
            self.grab_set()
        except Exception:
            pass

    def _arm(self):
        if self._page._arm(self.date.get(), self.time.get(), parent=self):
            self.destroy()


class AgentMigratorPage(ctk.CTkFrame):
    """Agent migration on one screen, in three steps.

    1 · What to move — the source accounts and sites, and where they go:
        the destination console, or a site-token file when there is no API
        access to it.
    2 · Check — reads what it needs from both sides and pairs every source
        group with where it lands. Nothing moves.
    3 · Migrate — moves exactly what the check listed, and names every
        machine as it goes.

    The results sit underneath in tabs: the plan, the agents, the API
    calls going out, and the status report. There is no dry-run switch:
    step 2 cannot move anything, and step 3 only unlocks on a plan that
    still matches what step 1 says.
    """

    MODE_CONSOLE = "Destination console"
    MODE_FILE = "Site-token file"

    # Seconds between counter repaints while something is running.
    COUNTS_EVERY = 0.3

    TABS = (("plan", "Plan", 104), ("agents", "Agents", 116),
            ("calls", "API calls", 100), ("status", "Status report", 126))
    HINTS = {
        "plan": "Every source group beside where it lands — click a row for "
                "the whole story.",
        "agents": "Every agent in the run, with what the console said "
                  "about it.",
        "calls": "The requests going out, as they go out — tokens masked.",
        "status": "The agents in the step 1 scope, by console migration "
                  "status.",
    }
    PLAN_FILTERS = {"Everything": "all", "Will migrate": "ready",
                    "Can't migrate": "blocked", "Already done": "done"}
    LIVE_FILTERS = {"All agents": "all", "Moved": "moved",
                    "Pending": "pending", "Failed": "failed"}
    FEED_EMPTY = ("Every agent appears here as it is sent, with what the "
                  "console said about it.")
    CHECK_IDLE = ("Reads both sides and pairs every source group with where "
                  "it lands. Nothing is moved.")
    FIX_TIP = ("Point each source scope that found no same-named twin at "
               "the destination scope it really is. Remembered, and "
               "re-applied every time the destination is read again.")
    PAIR_TIP = ("Choose the row of the file each source site's agents go "
                "to, when the file spells the site differently — "
                "“shira-telecom” for “Shira Telecom”. Nothing is sent.")

    def __init__(self, master, app, **kw):
        super().__init__(master, fg_color="transparent", **kw)
        # Before anything is built: CustomTkinter flushes geometry while
        # widgets are created, and a page that has already asked for more
        # height than the window has pushes the app's status bar off-screen.
        self.grid_propagate(False)
        self.app = app
        self.log = _ConsoleProxy(app)
        self._dest_mode = "console"
        self._site_tokens = None    # a loaded site-token file
        self._file_map = None       # a destination map loaded from a file
        self._file_path = ""
        self._map_data = None       # the map the current plan came from
        self._map_path = ""
        self._dest_accounts = []    # every destination account name
        self._renames = []          # name fixes, re-applied on every read
        self._dest_version = 0      # bumped whenever "where to" changes
        self._plan = None
        self._plan_inputs = None
        self._plan_saved_at = ""
        self._report = None         # the run made from the current plan
        self._last_run = None       # the last run, kept for its export
        self._status = None         # the last status report
        self._busy = None           # "check" | "live" | "status" | None
        self._error = {}
        self._stop = threading.Event()
        self._live_counts = {"moved": 0, "pending": 0, "failed": 0,
                             "sent": 0, "total": 0}
        self._live_where = ""
        self._last_counts_paint = 0.0
        self._check_phase = ""
        self._last_check_paint = 0.0
        self._check_job = None
        self._gate_job = None
        self._run_started = 0.0
        self._timer_job = None
        self._sched_job = None
        self._sched_when = None
        self._tab = None

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(4, weight=1)
        self._build_header()
        self.cards = [self._build_what_card(), self._build_check_card(),
                      self._build_migrate_card()]
        self._build_results()
        self._show_tab("plan")
        self._refresh_gates()

    # ── layout ────────────────────────────────────────────────────────

    def _build_header(self):
        hdr = ctk.CTkFrame(self, fg_color="transparent")
        hdr.grid(row=0, column=0, sticky="ew", padx=20, pady=(16, 10))
        hdr.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(hdr, text="Agent Migration", anchor="w",
                     font=(UI_FONT, 22, "bold")).grid(row=0, column=0,
                                                      sticky="w")
        ctk.CTkLabel(hdr, anchor="w", font=(UI_FONT, 12),
                     text_color=TEXT_MUTED,
                     text="Move agents from the source console to the "
                          "destination — each one into the site and group of "
                          "the same name.").grid(row=1, column=0, sticky="w")
        self._sched_badge = ctk.CTkLabel(hdr, text="",
                                         font=(UI_FONT, 12, "bold"),
                                         text_color=ACCENT)
        self._sched_badge.grid(row=0, column=1, rowspan=2, sticky="e")

    def _build_what_card(self):
        card = _StepCard(self, 1, "What to move", "Blank = everything")
        card.grid(row=1, column=0, sticky="ew", padx=20, pady=(0, 8))
        body = card.body

        scope = ctk.CTkFrame(body, fg_color="transparent")
        scope.grid(row=0, column=0, sticky="ew")
        scope.grid_columnconfigure((1, 4), weight=1)
        ctk.CTkLabel(scope, text="Accounts", font=(UI_FONT, 12),
                     text_color=TEXT_MUTED).grid(row=0, column=0, sticky="w",
                                                 padx=(0, 8))
        self.src_accounts = ctk.CTkEntry(scope, height=32,
                                         placeholder_text="all accounts")
        self.src_accounts.grid(row=0, column=1, sticky="ew")
        self._acct_btn = _ghost(
            scope, "Choose…", lambda: self._choose_accounts(self.src_accounts),
            width=84)
        self._acct_btn.grid(row=0, column=2, padx=(6, 18))
        ctk.CTkLabel(scope, text="Sites", font=(UI_FONT, 12),
                     text_color=TEXT_MUTED).grid(row=0, column=3, sticky="w",
                                                 padx=(0, 8))
        self.src_sites = ctk.CTkEntry(scope, height=32,
                                      placeholder_text="all sites")
        self.src_sites.grid(row=0, column=4, sticky="ew")
        self._site_btn = _ghost(
            scope, "Choose…",
            lambda: self._choose_sites(self.src_accounts, self.src_sites),
            width=84)
        self._site_btn.grid(row=0, column=5, padx=(6, 0))
        for entry in (self.src_accounts, self.src_sites):
            for seq in ("<KeyRelease>", "<FocusOut>"):
                entry.bind(seq, self._queue_gates, add="+")

        into = ctk.CTkFrame(body, fg_color="transparent")
        into.grid(row=1, column=0, sticky="ew", pady=(8, 0))
        ctk.CTkLabel(into, text="Into", font=(UI_FONT, 12),
                     text_color=TEXT_MUTED).grid(row=0, column=0, sticky="w",
                                                 padx=(0, 26))
        self.dest_switch = ctk.CTkSegmentedButton(
            into, values=[self.MODE_CONSOLE, self.MODE_FILE],
            font=(UI_FONT, 12), height=30, command=self._on_dest_mode)
        self.dest_switch.set(self.MODE_CONSOLE)
        self.dest_switch.grid(row=0, column=1, sticky="w")
        self._file_btns = ctk.CTkFrame(into, fg_color="transparent")
        self._file_btns.grid(row=0, column=2, sticky="w", padx=(12, 0))
        self._load_btn = _ghost(self._file_btns, "Load file…",
                                self._load_dest_file, width=104)
        self._load_btn.grid(row=0, column=0)
        self._paste_btn = _ghost(self._file_btns, "Paste",
                                 self._paste_site_tokens, width=70)
        self._paste_btn.grid(row=0, column=1, padx=(6, 0))
        _ToolTip(self._paste_btn,
                 "Copy the Site Name and Site Token columns — headers "
                 "included — out of Excel or the e-mail, then click Paste.")
        self._rows_btn = _ghost(self._file_btns, "View rows",
                                self._view_file_rows, width=96)
        self._rows_btn.grid(row=0, column=2, padx=(6, 0))
        self._connect_btn = _ghost(into, "Connections…",
                                   self._open_connections, width=118)
        self._connect_btn.grid(row=0, column=3, sticky="w", padx=(12, 0))

        self._what_line = ctk.CTkLabel(body, text="", anchor="w",
                                       justify="left",
                                       font=(UI_FONT, 12, "bold"),
                                       text_color=TEXT_MUTED)
        self._what_line.grid(row=2, column=0, sticky="ew", pady=(8, 0))
        _wrap(body, self._what_line)
        return card

    def _build_check_card(self):
        card = _StepCard(self, 2, "Check", "Nothing moves")
        card.grid(row=2, column=0, sticky="ew", padx=20, pady=(0, 8))
        body = card.body
        acts = ctk.CTkFrame(body, fg_color="transparent")
        acts.grid(row=0, column=0, sticky="w")
        big = dict(height=38, corner_radius=RADIUS_MD,
                   font=(UI_FONT, 13, "bold"))
        self.check_btn = ctk.CTkButton(
            acts, text="🔍  Check what will move", width=216,
            border_color=BORDER, text_color_disabled=TEXT_FAINT,
            command=self._run_match, **big)
        self.check_btn.grid(row=0, column=0, padx=(0, 8))
        self.check_stop_btn = ctk.CTkButton(
            acts, text="■  Stop", width=100, fg_color=ACCENT,
            hover_color=ACCENT_HOVER, command=self._stop_run, **big)
        self.check_stop_btn.grid(row=0, column=1, padx=(0, 8))
        self.fix_btn = _ghost(acts, "✎  Fix names", self._fix_names,
                              width=126, height=38, text_color=WARN)
        self.fix_btn.grid(row=0, column=2, padx=(0, 8))
        self._fix_tip = _ToolTip(self.fix_btn, self.FIX_TIP)
        self.plan_export_btn = _ghost(acts, "Export plan", self._export_match,
                                      width=112, height=38)
        self.plan_export_btn.grid(row=0, column=3)
        self._check_line = ctk.CTkLabel(body, text="", anchor="w",
                                        justify="left", font=(UI_FONT, 12),
                                        text_color=TEXT_MUTED)
        self._check_line.grid(row=1, column=0, sticky="ew", pady=(6, 0))
        _wrap(body, self._check_line)
        return card

    def _build_migrate_card(self):
        card = _StepCard(self, 3, "Migrate", "Moves them for real")
        card.grid(row=3, column=0, sticky="ew", padx=20, pady=(0, 8))
        body = card.body
        acts = ctk.CTkFrame(body, fg_color="transparent")
        acts.grid(row=0, column=0, sticky="w")
        big = dict(height=38, corner_radius=RADIUS_MD,
                   font=(UI_FONT, 13, "bold"))
        self.live_btn = ctk.CTkButton(
            acts, text="▶  Migrate", width=232, border_color=BORDER,
            text_color_disabled=TEXT_FAINT, command=self._run_live, **big)
        self.live_btn.grid(row=0, column=0, padx=(0, 8))
        self.live_stop_btn = ctk.CTkButton(
            acts, text="■  Stop", width=100, fg_color=ACCENT,
            hover_color=ACCENT_HOVER, command=self._stop_run, **big)
        self.live_stop_btn.grid(row=0, column=1, padx=(0, 8))
        self.sched_btn = _ghost(acts, "⏱  Schedule…", self._open_schedule,
                                width=126, height=38)
        self.sched_btn.grid(row=0, column=2, padx=(0, 8))
        _ToolTip(self.sched_btn,
                 "Run step 3 unattended at a date and time you choose — it "
                 "moves exactly what the current plan lists.")
        self.disarm_btn = _ghost(acts, "Disarm", self._disarm, width=96,
                                 height=38, text_color=ACCENT)
        self.disarm_btn.grid(row=0, column=3, padx=(0, 8))
        self.live_export_btn = _ghost(acts, "Export results",
                                      self._export_live, width=126, height=38)
        self.live_export_btn.grid(row=0, column=4)

        meta = ctk.CTkFrame(body, fg_color="transparent")
        meta.grid(row=1, column=0, sticky="ew", pady=(6, 0))
        meta.grid_columnconfigure(1, weight=1)
        self.verify_var = ctk.BooleanVar(value=True)
        self.verify_switch = ctk.CTkSwitch(
            meta, text="Confirm each agent arrived", variable=self.verify_var,
            font=(UI_FONT, 12), progress_color=GREEN, switch_width=34,
            switch_height=18, command=self._refresh_gates)
        self.verify_switch.grid(row=0, column=0, sticky="w")
        _ToolTip(self.verify_switch,
                 "The move call only answers with a count. Reading each agent "
                 "back is what turns that into moved, pending or failed per "
                 "machine — one extra call per 100 agents.")
        self._timer_lbl = ctk.CTkLabel(meta, text="", font=(MONO_FONT, 12),
                                       text_color=TEXT_MUTED)
        self._timer_lbl.grid(row=0, column=2, sticky="e")

        self.live_prog = ctk.CTkProgressBar(body, height=8, corner_radius=4,
                                            progress_color=GREEN)
        self.live_prog.grid(row=2, column=0, sticky="ew", pady=(8, 0))
        self.live_prog.set(0)
        self._live_line = ctk.CTkLabel(body, text="", anchor="w",
                                       justify="left", font=(UI_FONT, 12),
                                       text_color=TEXT_MUTED)
        self._live_line.grid(row=3, column=0, sticky="ew", pady=(6, 0))
        _wrap(body, self._live_line)
        return card

    def _build_results(self):
        wrap = ctk.CTkFrame(self, fg_color="transparent")
        wrap.grid(row=4, column=0, sticky="nsew", padx=20, pady=(2, 14))
        wrap.grid_columnconfigure(0, weight=1)
        wrap.grid_rowconfigure(1, weight=1)
        bar = ctk.CTkFrame(wrap, fg_color="transparent")
        bar.grid(row=0, column=0, sticky="ew", pady=(0, 6))
        self._tab_btns = {}
        for key, text, width in self.TABS:
            btn = ctk.CTkButton(bar, text=text, width=width, height=30,
                                corner_radius=RADIUS_MD, border_width=1,
                                font=(UI_FONT, 12, "bold"),
                                command=lambda k=key: self._show_tab(k))
            btn.pack(side="left", padx=(0, 6))
            self._tab_btns[key] = btn
        self._tab_hint = ctk.CTkLabel(bar, text="", font=(UI_FONT, 11),
                                      text_color=TEXT_FAINT, anchor="w")
        self._tab_hint.pack(side="left", padx=(8, 0), fill="x", expand=True)
        host = ctk.CTkFrame(wrap, fg_color="transparent")
        host.grid(row=1, column=0, sticky="nsew")
        host.grid_columnconfigure(0, weight=1)
        host.grid_rowconfigure(0, weight=1)
        self._tabs = {"plan": self._build_plan_tab(host),
                      "agents": self._build_agents_tab(host),
                      "calls": self._build_calls_tab(host),
                      "status": self._build_panel_status(host)}
        for frame in self._tabs.values():
            frame.grid(row=0, column=0, sticky="nsew")
            frame.grid_remove()

    def _list_tab(self, host, values, command):
        frame = ctk.CTkFrame(host, fg_color="transparent")
        frame.grid_columnconfigure(0, weight=1)
        frame.grid_rowconfigure(1, weight=1)
        top = ctk.CTkFrame(frame, fg_color="transparent")
        top.grid(row=0, column=0, sticky="ew", pady=(0, 6))
        chooser = ctk.CTkSegmentedButton(top, values=values,
                                         font=(UI_FONT, 11), command=command)
        chooser.set(values[0])
        chooser.pack(side="left")
        count = ctk.CTkLabel(top, text="", font=(UI_FONT, 11),
                             text_color=TEXT_FAINT)
        count.pack(side="left", padx=12)
        detail = ctk.CTkLabel(frame, text="", anchor="w", justify="left",
                              font=(UI_FONT, 11), text_color=TEXT_MUTED)
        detail.grid(row=2, column=0, sticky="ew", pady=(6, 0))
        _wrap(frame, detail)
        return frame, chooser, count, detail

    def _build_plan_tab(self, host):
        frame, self.plan_filter, self._plan_count, self._plan_detail = \
            self._list_tab(host, list(self.PLAN_FILTERS), self._on_plan_filter)
        self.plan_list = _PlanList(frame, on_select=self._on_plan_pick,
                                   empty=self._plan_placeholder())
        self.plan_list.grid(row=1, column=0, sticky="nsew")
        return frame

    def _build_agents_tab(self, host):
        frame, self.live_filter, self._feed_count, self._agent_detail = \
            self._list_tab(host, list(self.LIVE_FILTERS), self._on_live_filter)
        self.live_feed = _LiveFeed(frame, on_select=self._on_agent_pick,
                                   empty=self.FEED_EMPTY)
        self.live_feed.grid(row=1, column=0, sticky="nsew")
        return frame

    def _build_calls_tab(self, host):
        self.cmd_log = _CommandLog(host)
        return self.cmd_log

    def _build_panel_status(self, host):
        frame = ctk.CTkFrame(host, fg_color="transparent")
        frame.grid_columnconfigure(0, weight=1)
        frame.grid_rowconfigure(3, weight=1)
        top = ctk.CTkFrame(frame, fg_color="transparent")
        top.grid(row=0, column=0, sticky="ew")
        self.status_btn = ctk.CTkButton(
            top, text="📈  Run status report", width=190, height=34,
            corner_radius=RADIUS_MD, font=(UI_FONT, 12, "bold"),
            text_color_disabled=TEXT_FAINT, command=self._run_status,
            **_PRIMARY)
        self.status_btn.grid(row=0, column=0, padx=(0, 8))
        self.status_export_btn = _ghost(top, "Export", self._export_status,
                                        width=90, height=34)
        self.status_export_btn.grid(row=0, column=1, padx=(0, 16))
        self.pass_var = ctk.BooleanVar(value=False)
        ctk.CTkSwitch(top, text="Include passphrases", variable=self.pass_var,
                      font=(UI_FONT, 12), progress_color=WARN,
                      switch_width=34, switch_height=18,
                      command=self._on_passphrase_toggle).grid(
            row=0, column=2, sticky="w")
        note = ctk.CTkLabel(
            frame, anchor="w", justify="left", font=(UI_FONT, 11),
            text_color=TEXT_FAINT,
            text="Counts every agent in the step 1 scope by its console "
                 "migration status — the report you hand back after a "
                 "migration window. Leave passphrases off unless the customer "
                 "needs them: one API call per agent, every fetch is recorded "
                 "in the source console's activity log, and the export "
                 "becomes sensitive.")
        note.grid(row=1, column=0, sticky="ew", pady=(6, 6))
        _wrap(frame, note)
        self.status_stats = _Stats(frame)
        self.status_stats.grid(row=2, column=0, sticky="ew", pady=(0, 6))
        self.status_table = _VirtualTable(
            frame, [_col("Computer", "Computer Name", 3, mono=True),
                    _col("Migration status", "Migration Status", 2),
                    _col("Site", "Site", 2), _col("Group", "Group", 2),
                    _col("Decommissioned", "Decommissioned", 1)],
            empty="No report yet — it counts the agents in the step 1 scope.",
            colour=self._status_colour)
        self.status_table.grid(row=3, column=0, sticky="nsew")
        return frame

    # ── state: the one place that decides what can be done now ────────

    def _inputs(self):
        """Everything a plan depends on. A plan built from other inputs is
        stale, and step 3 will not move it."""
        src = self.app.source_api
        if self._dest_mode == "file":
            dest = ("file", self._dest_version)
        else:
            dst = getattr(self.app, "dest_api", None)
            dest = ("console", str(getattr(dst, "base_url", "") or ""),
                    self._dest_version)
        return (str(getattr(src, "base_url", "") or ""),
                " ".join(split_ids(self.src_accounts.get())),
                " ".join(split_ids(self.src_sites.get())), dest)

    def _plan_stale(self):
        return self._plan is not None and self._plan_inputs != self._inputs()

    def _dest_ready(self):
        """Step 1 knows where the agents go, whichever way it is set."""
        if self._dest_mode == "file":
            return (bool((self._site_tokens or {}).get("entries"))
                    or bool(self._file_map))
        return getattr(self.app, "dest_api", None) is not None

    def _queue_gates(self, _e=None):
        self._error.pop("check", None)
        if self._gate_job is not None:
            try:
                self.after_cancel(self._gate_job)
            except Exception:
                pass
        self._gate_job = self.after(150, self._refresh_gates)

    def _refresh_gates(self):
        """Show, enable and label every control for where the operator is."""
        self._gate_job = None
        busy = self._busy
        idle = busy is None
        armed = self._sched_when is not None
        file_mode = self._dest_mode == "file"
        src = self.app.source_api
        dst = getattr(self.app, "dest_api", None)

        s1, text, colour = self._what_state()
        self.cards[0].set(s1)
        _retext(self._what_line, text=text, text_color=colour)
        editable = "normal" if idle and not armed else "disabled"
        for widget in (self.src_accounts, self.src_sites, self._acct_btn,
                       self._site_btn, self._load_btn, self._paste_btn,
                       self._rows_btn, self.dest_switch):
            _retext(widget, state=editable)
        _grid_shown(self._file_btns, file_mode)
        _grid_shown(self._rows_btn, bool(self._site_tokens))
        _grid_shown(self._connect_btn,
                    src is None or (not file_mode and dst is None))
        _retext(self._load_btn,
                **(_GHOSTED if self._dest_ready() else _PRIMARY))

        ready1 = s1 in ("done", "warn") and self._dest_ready()
        s2, text, colour = self._check_state(ready1)
        self.cards[1].set(s2)
        _retext(self._check_line, text=text, text_color=colour)
        fresh = self._plan is not None and not self._plan_stale()
        spent = self._report is not None
        can_check = ready1 and idle and not armed
        _grid_shown(self.check_btn, busy != "check")
        _grid_shown(self.check_stop_btn, busy == "check")
        _retext(self.check_btn, state="normal" if can_check else "disabled",
                text="↻  Check again" if self._plan is not None
                else "🔍  Check what will move",
                **(_PRIMARY if can_check and (spent or not fresh)
                   else _GHOSTED))
        _grid_shown(self.fix_btn, fresh and idle and not armed
                    and self._fixable())
        _retext(self.fix_btn, text="✎  Pair sites" if file_mode
                else "✎  Fix names")
        self._fix_tip.text = self.PAIR_TIP if file_mode else self.FIX_TIP
        _grid_shown(self.plan_export_btn,
                    bool((self._plan or {}).get("rows")) and busy != "check")

        s3, text, colour = self._live_state()
        self.cards[2].set(s3)
        _retext(self._live_line, text=text, text_color=colour)
        agents = (self._plan or {}).get("agents", 0)
        can_go = (fresh and not spent and agents > 0 and idle
                  and src is not None)
        _grid_shown(self.live_btn, busy != "live")
        _grid_shown(self.live_stop_btn, busy == "live")
        _retext(self.live_btn,
                state="normal" if can_go and not armed else "disabled",
                text=(f"▶  Migrate {agents:,} agent{'' if agents == 1 else 's'}"
                      if fresh and not spent and agents else "▶  Migrate"),
                **(_GO if can_go and not armed else _GHOSTED))
        _grid_shown(self.sched_btn, not armed)
        _retext(self.sched_btn, state="normal" if can_go else "disabled")
        _grid_shown(self.disarm_btn, armed)
        _grid_shown(self.live_export_btn,
                    bool(((self._last_run or {}).get("run") or {})
                         .get("agents")) and busy != "live")
        _retext(self.verify_switch,
                state="disabled" if busy == "live" else "normal")
        ran = busy == "live" or self._report is not None
        _grid_shown(self.live_prog, ran)
        _grid_shown(self._timer_lbl, ran)

        _retext(self.status_btn,
                state="normal" if idle and src is not None else "disabled",
                text="Counting…" if busy == "status"
                else "📈  Run status report")
        status = self._status or {}
        _grid_shown(self.status_export_btn,
                    bool(status.get("agents") or status.get("decommissioned"))
                    and busy != "status")
        _retext(self._sched_badge, text=(
            f"⏱  {format_countdown(self._sched_when - datetime.now())}"
            if armed else ""))
        self._paint_list_counts()

    def _scope_words(self):
        accounts = split_ids(self.src_accounts.get())
        sites = split_ids(self.src_sites.get())

        def n(count, noun):
            return f"{count} {noun}{'' if count == 1 else 's'}"
        if sites:
            return n(len(sites), "site") + (
                f" in {n(len(accounts), 'account')}" if accounts else "")
        return n(len(accounts), "account") if accounts else "Every account"

    def _what_state(self):
        """(card state, line, colour) for step 1."""
        src = self.app.source_api
        if src is None:
            return ("error", "✕  The SOURCE console isn't connected — connect "
                    "it on the Connections page.", ACCENT)
        source = self._console_name(src)
        scope = self._scope_words()
        if self._dest_mode == "console":
            dst = getattr(self.app, "dest_api", None)
            if dst is None:
                return ("error", "✕  The DESTINATION console isn't connected "
                        "— connect it on the Connections page, or choose "
                        "“Site-token file” if there is no API access to it.",
                        ACCENT)
            return ("done", f"➜  {scope} on {source}  →  same-named "
                    f"accounts, sites and groups on "
                    f"{self._console_name(dst)}", TEXT)
        if self._site_tokens:
            data = self._site_tokens
            entries = data.get("entries") or []
            bad = sum(1 for e in entries if e["problem"])
            hosts = [console_host(c) for c in data.get("consoles") or []]
            name = os.path.basename(data.get("source") or "") or "pasted text"
            head = (f"{name}: {len(entries)} site"
                    f"{'' if len(entries) == 1 else 's'}, tokens for "
                    f"{', '.join(hosts) or 'no console'}")
            if bad == len(entries):
                return ("error", f"✕  {head} — no row can be used. View rows "
                        f"says why; fix the file and load it again.", ACCENT)
            if bad:
                return ("warn", f"⚠  {head} — {bad} row"
                        f"{'' if bad == 1 else 's'} can't be used (View "
                        f"rows). {scope} on {source} moves by site name.",
                        WARN)
            if len(hosts) > 1:
                return ("warn", f"⚠  {head} — tokens for {len(hosts)} "
                        f"different consoles; check this is the right file.",
                        WARN)
            return ("done", f"➜  {scope} on {source}  →  {head}. Each source "
                    f"site moves with its row's token.", TEXT)
        if self._file_map:
            accounts, sites, groups = map_counts(self._file_map)
            name = os.path.basename(self._file_path) or "a destination map"
            return ("done", f"➜  {scope} on {source}  →  {name} (a saved "
                    f"destination map: {accounts} accounts, {sites} sites, "
                    f"{groups} groups).", TEXT)
        return ("active", "Load the partner's sheet of site names and "
                "registration tokens (CSV or Excel), or paste its two "
                "columns. A destination map saved earlier (.json) works "
                "too.", TEXT_MUTED)

    def _check_state(self, ready1):
        """(card state, line, colour) for step 2."""
        plan = self._plan
        if self._busy == "check":
            return "active", self._check_progress_text(), TEXT_MUTED
        if not ready1:
            return "locked", self.CHECK_IDLE, TEXT_FAINT
        if self._error.get("check"):
            return ("error", f"✕  Could not check: {self._error['check']}",
                    ACCENT)
        if plan is None:
            return "active", self.CHECK_IDLE, TEXT_MUTED
        if self._plan_stale():
            return ("warn", "⚠  Step 1 changed since this plan was made — "
                    "check again before migrating.", WARN)
        ready, agents = plan.get("ready", 0), plan.get("agents", 0)
        bits = [f"✓  {ready:,} group{'' if ready == 1 else 's'} ready · "
                f"{agents:,} agent{'' if agents == 1 else 's'} to move"]
        if plan.get("blocked"):
            bits.append(f"✕  {plan['blocked']:,} can't move")
        if plan.get("done"):
            bits.append(f"{plan['done']:,} already done")
        if plan.get("mode") == "file":
            if plan.get("unlisted"):
                bits.append(f"{len(plan['unlisted'])} source site(s) not in "
                            f"the file")
            if plan.get("unused"):
                bits.append(f"{len(plan['unused'])} file row(s) matched "
                            f"nothing")
        text = "   ·   ".join(bits)
        if plan.get("stopped"):
            return ("warn", "■  Stopped early, so this plan is partial — "
                    "check again for the full picture.   " + text, WARN)
        if self._plan_saved_at:
            text += (f"   ·   plan from {plan_age(self._plan_saved_at)} — "
                     f"check again if anything changed")
        return ("done" if ready else "warn"), text, TEXT

    def _live_state(self):
        """(card state, line, colour) for step 3."""
        plan, run = self._plan, self._report
        if self._busy == "live":
            return "active", self._live_progress_text(), TEXT
        if self._error.get("live"):
            return ("error", f"✕  The migration failed: {self._error['live']}",
                    ACCENT)
        if run is not None:
            return self._outcome(run)
        if not plan:
            return ("locked", "Unlocks once step 2 has a plan with agents to "
                    "move.", TEXT_FAINT)
        if self._plan_stale():
            return ("locked", "Check again first — the plan no longer matches "
                    "step 1.", TEXT_FAINT)
        agents = plan.get("agents", 0)
        if not agents:
            return ("locked", "Nothing to move — every matched group is empty "
                    "or already migrated.", TEXT_FAINT)
        if self._sched_when is not None:
            left = format_countdown(self._sched_when - datetime.now())
            return ("active", f"⏱  Armed — moving {agents:,} agent(s) in "
                    f"{left} (at {self._sched_when:%Y-%m-%d %H:%M}). Keep "
                    f"the app open and the machine awake.", ACCENT)
        confirm = ("each one confirmed after the move"
                   if self.verify_var.get() else
                   "not confirmed per agent (“moved” = accepted)")
        groups = plan.get("ready", 0)
        return ("todo", f"{agents:,} agent{'' if agents == 1 else 's'} in "
                f"{groups:,} group{'' if groups == 1 else 's'}, "
                f"{self._console_name(self.app.source_api)} → "
                f"{self._dest_name()} — {confirm}.", TEXT_MUTED)

    @staticmethod
    def _outcome(run):
        moved, pending = run.get("moved", 0), run.get("pending", 0)
        failed = run.get("failed", 0)
        again = " Check again to see what's left."
        if not run.get("agents"):
            if run.get("errors"):
                return ("error", f"✕  Nothing was sent — "
                        f"{len(run['errors'])} group(s) could not be read. "
                        f"The API calls tab has the errors.", ACCENT)
            return ("done", "✓  Nothing needed moving — those agents had "
                    "already migrated.", GREEN)
        if failed:
            return (("warn" if moved or pending else "error"),
                    f"✕  {moved:,} moved · {pending:,} pending · {failed:,} "
                    f"failed — the Agents tab lists every failure with the "
                    f"console's reason.{again}", ACCENT)
        if run.get("stopped"):
            return ("warn", f"■  Stopped early — {moved:,} moved and "
                    f"{pending:,} pending before the stop.{again}", WARN)
        if not run.get("verified"):
            return ("warn", f"■  {run.get('sent', 0):,} sent, not confirmed "
                    f"per agent (confirmation was off) — run the Status "
                    f"report to check they arrived.", WARN)
        if pending:
            return ("warn", f"■  {moved:,} moved · {pending:,} accepted, "
                    f"moving on their next check-in — run the Status report "
                    f"later to confirm.", WARN)
        if run.get("errors"):
            return ("warn", f"⚠  {moved:,} moved, but {len(run['errors'])} "
                    f"group(s) could not be read — the API calls tab has the "
                    f"errors.{again}", WARN)
        return ("done", f"✓  All {moved:,} agent(s) moved and confirmed on "
                f"the destination.", GREEN)

    def _fixable(self):
        plan = self._plan or {}
        if plan.get("mode") == "file":
            return bool(plan.get("unlisted") and self._pair_rows(plan))
        return bool(plan.get("unmatched") or self._renames)

    def _check_progress_text(self):
        total = self.plan_list.counts()[1]
        text = f"⏳  {self._check_phase or 'Starting…'}"
        if total:
            text += (f"   ·   {total:,} scope{'' if total == 1 else 's'} "
                     f"listed so far")
        return text

    def _live_progress_text(self):
        c = self._live_counts
        done = c["moved"] + c["pending"] + c["failed"]
        text = (f"{done:,} of {c['total']:,} processed   ·   "
                f"{c['moved']:,} moved   ·   {c['pending']:,} pending   ·   "
                f"{c['failed']:,} failed")
        return text + (f"\n{self._live_where}" if self._live_where else "")

    def _show_tab(self, key):
        if key == self._tab:
            return
        self._tab = key
        for name, frame in self._tabs.items():
            _grid_shown(frame, name == key)
        for name, btn in self._tab_btns.items():
            on = name == key
            _retext(btn, fg_color=BRAND if on else GHOST,
                    hover_color=BRAND_HOVER if on else GHOST_HOVER,
                    text_color="#FFFFFF" if on else TEXT_MUTED,
                    border_color=BRAND if on else BORDER)
        _retext(self._tab_hint, text=self.HINTS[key])

    def _paint_list_counts(self):
        titles = {key: title for key, title, _w in self.TABS}
        for key, table, label, noun in (
                ("plan", self.plan_list, self._plan_count, "scope"),
                ("agents", self.live_feed, self._feed_count, "agent")):
            shown, total = table.counts()
            if shown != total:
                words = f"{shown:,} of {total:,} shown"
            else:
                words = (f"{total:,} {noun}{'' if total == 1 else 's'}"
                         if total else "")
            _retext(label, text=words)
            _retext(self._tab_btns[key],
                    text=f"{titles[key]} · {total:,}" if total
                    else titles[key])

    # ── shared bits ───────────────────────────────────────────────────

    def _say(self, msg):
        self.after(0, lambda m=msg: self.log.log(m))

    def _source_api(self):
        api = self.app.source_api
        if api is None:
            raise ValueError("The SOURCE console isn't connected — connect "
                             "it on the Connections page.")
        return api

    def _dest_api(self):
        api = getattr(self.app, "dest_api", None)
        if api is None:
            raise ValueError("The DESTINATION console isn't connected — "
                             "connect it on the Connections page, or switch "
                             "step 1 to a site-token file.")
        return api

    @staticmethod
    def _console_name(api):
        if api is None:
            return "—"
        return str(getattr(api, "base_url", "")).replace(
            "https://", "").split(".")[0] or "connected"

    def _dest_name(self):
        """The console the agents are going to, in a word or two."""
        if self._dest_mode == "file":
            if self._site_tokens:
                hosts = [console_host(c) for c in
                         self._site_tokens.get("consoles") or []]
                return ", ".join(hosts) or "—"
            return os.path.basename(self._file_path) or "—"
        return self._console_name(getattr(self.app, "dest_api", None))

    def _dest_source(self):
        """(label, value) naming where the plan's tokens came from."""
        plan = self._plan or {}
        if plan.get("mode") == "file":
            return "Site-token file", plan.get("file") or "a site-token file"
        if self._dest_mode == "file":
            return "Destination map file", self._file_path or "—"
        dst = getattr(self.app, "dest_api", None)
        return "Destination console", (self._map_path or str(
            getattr(dst, "base_url", "") or "—"))

    @staticmethod
    def _fill(entry, ids):
        entry.delete(0, "end")
        entry.insert(0, ", ".join(ids))

    def _open_connections(self):
        show = getattr(self.app, "_show", None)
        if callable(show) and "Connections" in (getattr(self.app, "pages",
                                                        None) or {}):
            show("Connections")

    def _choose_accounts(self, entry):
        """Tick source accounts the console finds for us.

        The dialog opens straight away and does its own searching on a
        worker thread, so a console with thousands of accounts costs one
        small page rather than a full walk before anything appears.
        """
        try:
            api = self._source_api()
        except ValueError as e:
            messagebox.showwarning("Console", str(e))
            return

        def fetch(query, limit):
            return [(str(a["id"]), a.get("name") or "(unnamed)",
                     f"id {a['id']}")
                    for a in (api.search_accounts(query, limit) or [])
                    if a.get("id")]

        chosen = _ChoiceDialog(
            self.winfo_toplevel(), "Choose accounts", fetch,
            chosen=split_ids(entry.get()), noun="account",
            source="the source console",
            note="Tick the source accounts to migrate. Leave everything "
                 "unticked to cover them all.").ask()
        if chosen is None:
            return
        self._fill(entry, chosen)
        self.log.log(f"Accounts: {', '.join(chosen) or 'all of them'}")
        self._queue_gates()

    def _choose_sites(self, account_entry, site_entry):
        """Sites for whichever accounts are in the field beside this one."""
        try:
            api = self._source_api()
        except ValueError as e:
            messagebox.showwarning("Console", str(e))
            return
        account_ids = split_ids(account_entry.get())
        scope = (f"in the {len(account_ids)} account(s) beside this field"
                 if account_ids else "across every account this token sees")

        def fetch(query, limit):
            out = []
            for s in (api.search_sites(
                    query, ",".join(account_ids), limit) or []):
                ident = str(s.get("id", ""))
                if not ident:
                    continue
                # The site payload already carries its account name, so
                # there is no need to walk /accounts just to label a row.
                acct = s.get("accountName") or ""
                out.append((ident, s.get("name") or "(unnamed)",
                            f"{acct} · id {ident}" if acct else f"id {ident}"))
            return out

        chosen = _ChoiceDialog(
            self.winfo_toplevel(), "Choose sites", fetch,
            chosen=split_ids(site_entry.get()), noun="site",
            source="the source console",
            note=f"Searching {scope}. Tick the sites to include — leave "
                 f"everything unticked for all of them.").ask()
        if chosen is None:
            return
        self._fill(site_entry, chosen)
        self.log.log(f"Sites: {', '.join(chosen) or 'all of them'}")
        self._queue_gates()

    # ── step 1 · where the agents go ──────────────────────────────────

    def _on_dest_mode(self, value):
        """Switch where the agents go.

        A plan never survives the switch: it was built from the other
        tokens, and step 3 must only ever move what the operator reviewed
        against the tokens now in play.
        """
        mode = "file" if value == self.MODE_FILE else "console"
        if mode == self._dest_mode:
            return
        if self._busy:
            self.dest_switch.set(self.MODE_FILE if self._dest_mode == "file"
                                 else self.MODE_CONSOLE)
            return
        self._dest_mode = mode
        self._dest_version += 1
        if self._plan is not None:
            self.log.log("Switched where the agents go — check again before "
                         "migrating.")
        self._forget_plan_view()
        self._refresh_gates()

    def _forget_plan_view(self):
        """Drop the plan, the run made from it, and what shows them."""
        self._plan = None
        self._plan_inputs = None
        self._plan_saved_at = ""
        self._report = None
        self._error.pop("check", None)
        self._error.pop("live", None)
        if self._sched_when is not None:
            self._cancel_schedule()
            cli_log("Scheduled agent migration disarmed — the plan it was "
                    "armed with is gone.", "warning")
        self.plan_list.clear(self._plan_placeholder())
        _retext(self._plan_detail, text="")
        self.live_prog.set(0)
        _retext(self._timer_lbl, text="")

    def _plan_placeholder(self):
        if self._dest_mode == "file":
            return ("Nothing checked yet — press “Check what will move” and "
                    "every source site is listed here beside the row of the "
                    "file it pairs with.")
        return ("Nothing checked yet — press “Check what will move” and every "
                "source group is listed here beside the destination group it "
                "lands in.")

    def _load_dest_file(self):
        path = filedialog.askopenfilename(
            title="Open a site-token file — or a saved destination map",
            filetypes=[SITE_TOKEN_FILETYPES[0],
                       ("Destination map (JSON)", "*.json"),
                       ("All files", "*.*")])
        if not path:
            return
        if path.lower().endswith(".json"):
            try:
                with open(path, encoding="utf-8") as f:
                    data = json.load(f)
            except Exception as e:      # noqa: BLE001 - shown to the operator
                messagebox.showerror("Could not read that file", str(e))
                return
            if isinstance(data, dict):
                data = data.get("map")
            if not isinstance(data, list):
                messagebox.showerror(
                    "Not a destination map",
                    "Expected the JSON list of accounts, sites and groups "
                    "that Agent Migration saves.")
                return
            self._set_file_map(data, path)
            return
        try:
            data = read_site_tokens(path)
        except Exception as e:          # noqa: BLE001 - shown to the operator
            messagebox.showerror("Could not read that file", str(e))
            return
        self._set_site_tokens(data)

    def _paste_site_tokens(self):
        """The two columns copied straight out of Excel or an e-mail."""
        try:
            text = self.clipboard_get()
        except tk.TclError:
            text = ""
        if not text.strip():
            messagebox.showinfo(
                "Nothing to paste",
                "Copy the Site Name and Site Token columns — headers "
                "included — from Excel or the e-mail, then try again.")
            return
        try:
            data = parse_site_tokens_text(text)
        except ValueError as e:
            messagebox.showerror("Could not read what was pasted", str(e))
            return
        self._set_site_tokens(data)

    def _set_site_tokens(self, data):
        """Adopt a site-token file: forget the old plan, say what it holds."""
        self._site_tokens = data
        self._file_map, self._file_path = None, data.get("source", "")
        self._dest_version += 1
        self._forget_plan_view()
        entries = data["entries"]
        bad = sum(1 for e in entries if e["problem"])
        for skip in (data.get("skipped") or [])[:20]:
            self.log.log(f"Site tokens: line {skip['row']} skipped — "
                         f"{skip['reason']}.")
        cli_log(f"Loaded {len(entries)} site token(s) from {data['source']}"
                + (f" — {bad} row(s) need fixing" if bad else "") + ".",
                "warning" if bad else "info")
        self.app.log_audit(
            "agent_site_tokens_loaded", sites=len(entries), problems=bad,
            consoles=[console_host(c) for c in data.get("consoles") or []],
            source=data["source"])
        self._refresh_gates()

    def _set_file_map(self, data, path):
        """Adopt a destination map saved earlier, instead of the console."""
        self._file_map, self._file_path = data, path
        self._site_tokens = None
        self._dest_version += 1
        self._set_map(data, path)
        accounts, sites, groups = map_counts(data)
        stranded = sites_without_token(data)
        cli_log(f"Loaded destination map {path}: {accounts} account(s), "
                f"{sites} site(s), {groups} group(s)"
                + (f" — {len(stranded)} site(s) have no registration token"
                   if stranded else "") + ".",
                "warning" if stranded else "info")
        self._refresh_gates()

    def _set_map(self, data, path):
        """Adopt a destination map; a plan built from another one goes."""
        self._map_data = data
        self._map_path = path
        self._forget_plan_view()

    def _view_file_rows(self):
        data = self._site_tokens or {}
        entries = data.get("entries") or []
        if not entries:
            return
        _TableDialog(
            self, "Site tokens — " + (os.path.basename(
                data.get("source") or "") or "pasted text"),
            "Every row of the file: where its agents land, the console its "
            "token belongs to, and whether the row can be used. Tokens are "
            "never shown — only enough of the key to tell two apart.",
            [_col(name, name, weight) for name, weight in
             zip(SITE_TOKEN_COLUMNS, (4, 4, 2, 3))],
            self._site_token_rows(entries),
            colour=lambda r: WARN if r["_bad"] else None)

    @staticmethod
    def _site_token_rows(entries):
        rows = []
        for e in entries:
            rows.append({
                "Site": e["source"]
                + (f"  ({e['account']})" if e["account"] else ""),
                "Lands on": file_destination(e),
                "Token": token_hint(e["token"]) if e["token"] else "—",
                "Check": e["problem"] or ("ok · paired by hand"
                                          if e.get("paired") else "ok"),
                "_bad": bool(e["problem"])})
        return rows

    def _file_kind(self):
        if self._site_tokens:
            return "tokens"
        return "map" if self._file_map is not None else None

    # ── step 2 · check: pair the source with where it goes ────────────

    def _run_match(self, reread=True):
        if self._busy:
            return
        file_mode = self._dest_mode == "file"
        kind = self._file_kind()
        try:
            src = self._source_api()
            if not self._dest_ready():
                raise ValueError(
                    "Load the site-token file in step 1 first." if file_mode
                    else "Connect the DESTINATION console first — or switch "
                         "step 1 to a site-token file.")
            dst = None if file_mode else self._dest_api()
        except ValueError as e:
            messagebox.showwarning("Cannot check yet", str(e))
            return

        account_ids = self.src_accounts.get().strip()
        site_ids = self.src_sites.get().strip()
        src_url = str(getattr(src, "base_url", "") or "")
        dst_url = str(getattr(dst, "base_url", "") or "")
        tokens, file_map = self._site_tokens, self._file_map
        renames = [dict(r) for r in self._renames]
        reuse = None if reread or file_mode else self._map_data
        names = list(self._dest_accounts)
        inputs = self._inputs()

        self._forget_plan_view()
        self._stop.clear()
        _retext(self.check_stop_btn, state="normal")
        self._busy = "check"
        self._check_phase = "Starting…"
        self.plan_filter.set("Everything")
        self.plan_list.set_filter("all")
        self._show_tab("plan")
        self._refresh_gates()
        self.log.log("Checking what will move — nothing is sent…")

        pump = _EventPump(self, self.plan_list.put,
                          on_flush=self._paint_check_progress)
        pump.start()

        def do():
            if file_mode and kind == "tokens":
                return None, names, build_file_plan(
                    src, tokens, account_ids, site_ids, log=self._check_say,
                    should_stop=self._stop.is_set, on_row=pump.push,
                    source_url=src_url)
            if file_mode:
                map_data = file_map
            elif reuse is not None:
                map_data = reuse
            else:
                self._check_say("Reading the destination…")
                map_data, found = read_destination(
                    src, dst, account_ids, site_ids, renames,
                    log=self._check_say, should_stop=self._stop.is_set)
                names[:] = found
            self._check_say("Pairing the source with it…")
            return map_data, names, build_match_plan(
                src, map_data, account_ids, site_ids, log=self._check_say,
                should_stop=self._stop.is_set, on_row=pump.push)

        def done(result):
            pump.drain()
            pump.stop()
            map_data, found, plan = result
            self._busy = None
            if map_data is not None:
                self._map_data = map_data
                self._map_path = self._file_path if file_mode else dst_url
            if not file_mode:
                self._dest_accounts = list(found)
                if reuse is None:
                    self._keep_destination()
                    _accounts, sites, groups = map_counts(map_data)
                    self.app.log_audit("agent_map_build", sites=sites,
                                       groups=groups, path=dst_url)
            self._plan = plan
            self._plan_inputs = inputs
            self._plan_saved_at = ""
            self._report = None
            self._remember_plan(plan, account_ids, site_ids)
            self._paint_plan(plan)
            self._refresh_gates()
            cli_log(f"Checked: {plan['ready']} group(s) ready covering "
                    f"{plan['agents']} agent(s); {plan['blocked']} scope(s) "
                    f"cannot move.",
                    "warning" if plan["blocked"] else "success")

        def fail(e):
            pump.stop()
            self._busy = None
            self._error["check"] = str(e)
            self.plan_list.render()
            self._refresh_gates()
            cli_log(f"Check failed: {e}", "error")

        run_async(self, do, done, fail)

    def _check_say(self, msg):
        """A progress line from the worker: to the log, and onto step 2."""
        def show(m=msg):
            self.log.log(m)
            if self._busy == "check":
                self._check_phase = m.strip()
                self._paint_check_line()
        try:
            self.after(0, show)
        except Exception:
            pass

    def _paint_check_progress(self):
        self.plan_list.render()
        self._paint_check_line()

    def _paint_check_line(self, force=False):
        """Step 2's progress line, a few times a second at most."""
        now = time.monotonic()
        if not force and now - self._last_check_paint < self.COUNTS_EVERY:
            if self._check_job is None:
                self._check_job = self.after(
                    int(self.COUNTS_EVERY * 1000) + 20,
                    lambda: self._paint_check_line(force=True))
            return
        if self._check_job is not None:
            try:
                self.after_cancel(self._check_job)
            except Exception:
                pass
        self._check_job = None
        self._last_check_paint = now
        if self._busy == "check":
            _retext(self._check_line, text=self._check_progress_text(),
                    text_color=TEXT_MUTED)
            self._paint_list_counts()

    def _stop_run(self):
        self._stop.set()
        _retext(self.check_stop_btn, state="disabled")
        _retext(self.live_stop_btn, state="disabled")
        self.log.log("Stopping — finishing the current group…")

    def _paint_plan(self, plan):
        if not plan["rows"]:
            self.plan_list.clear(
                "No source scope produced a row — the accounts are empty, or "
                "the site filter left nothing in.")
        self.plan_list.render()
        # The console keeps the full text version, for pasting into a ticket.
        self.log.log(plan_summary_text(plan))

    def _on_plan_filter(self, value):
        self.plan_list.set_filter(self.PLAN_FILTERS.get(value, "all"))
        self._paint_list_counts()

    def _on_plan_pick(self, row):
        state = row.get("state")
        where = row.get("Status") if state == "blocked" \
            else row.get("Destination")
        count = int((row.get("Agents") if state == "ready"
                     else row.get("Total")) or 0)
        head = f"{scope_path(row)}  ──▶  {where}"
        if count:
            head += f"   ·   {count:,} agent{'' if count == 1 else 's'}"
        tail = row.get("reason") or plan_note(row)
        _retext(self._plan_detail, text=head + (f"\n{tail}" if tail else ""))

    def _export_match(self):
        plan = self._plan or {}
        if not plan.get("rows"):
            messagebox.showinfo("Nothing to export", "Check what will move "
                                                     "first.")
            return
        label, value = self._dest_source()
        ts = datetime.now().strftime("%Y%m%d-%H%M")
        export_agent_report(build_match_report(plan, value, map_label=label),
                            default_name=f"s1-agent-match-plan-{ts}")

    # ── remembering the plan ──────────────────────────────────────────

    def _fingerprint(self, account_ids, site_ids):
        src = self.app.source_api
        return plan_fingerprint(str(getattr(src, "base_url", "") or ""),
                                account_ids, site_ids,
                                f"live:{self._map_path}", self._map_data)

    def _keep_destination(self):
        """Keep the destination read (and the name fixes) for next time."""
        try:
            save_destination(self._map_data, self._map_path, self._renames)
        except OSError as e:
            self.log.log(f"Could not keep a copy of the destination read: {e}")

    def _remember_plan(self, plan, account_ids, site_ids):
        """Cache a complete plan so reopening the app skips the walk.

        Only plans read live from the destination console: a site-token
        file or a saved map is quick to check again.
        """
        if plan.get("stopped") or not plan.get("rows"):
            return
        if plan.get("mode") == "file" or self._dest_mode == "file":
            return
        try:
            save_plan(plan, self._fingerprint(account_ids, site_ids))
        except OSError as e:
            self.log.log(f"Could not cache the plan: {e}")

    def _restore_last_plan(self):
        """Bring back the last plan for these two consoles, so reopening
        the app doesn't mean walking both of them again."""
        if self._plan is not None or self._busy or self._dest_mode != "console":
            return False
        src = self.app.source_api
        dst = getattr(self.app, "dest_api", None)
        if src is None or dst is None:
            return False
        src_url = str(getattr(src, "base_url", "") or "")
        dst_url = str(getattr(dst, "base_url", "") or "")
        scope = cached_plan_scope()
        saved = load_destination()
        if (not scope or not saved or scope.get("source") != src_url
                or saved.get("destination") != dst_url):
            return False
        accounts, sites = scope.get("accounts", ""), scope.get("site", "")
        try:
            plan, saved_at = load_plan(
                plan_fingerprint(src_url, accounts, sites, f"live:{dst_url}",
                                 saved["map"]), saved["map"])
        except Exception:
            return False
        if not plan:
            return False
        self._fill(self.src_accounts, split_ids(accounts))
        self._fill(self.src_sites, split_ids(sites))
        self._map_data, self._map_path = saved["map"], dst_url
        self._renames = list(saved.get("renames") or [])
        self._plan, self._plan_saved_at = plan, saved_at
        self._plan_inputs = self._inputs()
        self.plan_list.set_rows(plan["rows"])
        self._paint_plan(plan)
        self._refresh_gates()
        cli_log(f"Restored the plan from {plan_age(saved_at)} — "
                f"{plan['ready']} group(s), {plan['agents']} agent(s). Check "
                f"again if the consoles have changed.", "info")
        return True

    # ── fixing names ──────────────────────────────────────────────────

    def _fix_names(self):
        plan = self._plan or {}
        if plan.get("mode") == "file":
            self._pair_names()
            return
        if not (plan.get("unmatched") or self._renames):
            messagebox.showinfo("Nothing to fix", "No name mismatches.")
            return
        partial = (self._dest_mode == "console"
                   and bool(split_ids(self.src_sites.get())))
        _NameFixDialog(self, NameComparison(
            plan, self._map_data, self._dest_accounts, self._renames,
            partial_sites=partial), self._after_fix)

    def _after_fix(self, chosen, removed=()):
        renames = [{"kind": item["kind"], "old": dest, "new": item["name"],
                    "within": list(item.get("within") or [])}
                   for item, dest in chosen]
        removed = [dict(r) for r in removed or []]
        if removed:
            self._renames = [r for r in self._renames if r not in removed]
        self._renames.extend(renames)
        self._dest_version += 1
        for r in renames:
            cli_log(f"Name fix: destination {r['kind']} '{r['old']}' is the "
                    f"source's '{r['new']}'.", "info")
        for r in removed:
            cli_log(f"Name fix removed: destination {r['kind']} '{r['old']}' "
                    f"is no longer used for the source's '{r['new']}'.",
                    "info")
        self.app.log_audit("agent_name_fix", fixes=len(renames),
                           removed=len(removed))
        console = self._dest_mode == "console"
        whole = console and (bool(removed) or any(
            r["kind"] == "account" for r in renames))
        if not whole and self._map_data is not None:
            for r in removed:           # a map file: its own name comes back
                rename_in_map(self._map_data, r["kind"], r["new"], r["old"],
                              within=r.get("within"))
            done = apply_renames(self._map_data, renames)
            if console and done < len(renames):
                whole = True            # a typed name the read never saw
            elif console:
                self._keep_destination()
        self._run_match(reread=whole)

    def _pair_names(self):
        """Let the operator choose, for each source site the file doesn't
        name, the row of the file it gets — nothing is guessed."""
        plan = self._plan or {}
        rows = self._pair_rows(plan)
        unlisted = plan.get("unlisted") or []
        if not (rows and unlisted):
            hint = ""
            if unlisted and (plan.get("unused") or plan.get("spare")):
                hint = ("\n\nThe rows no site uses have token problems — "
                        "fix them in the file and load it again.")
            elif unlisted:
                hint = ("\n\nEvery row of the file found its site — add the "
                        "missing sites to the file and load it again.")
            messagebox.showinfo(
                "Nothing to pair",
                "Pairing needs a source site the file doesn't name and a "
                "row of the file no site uses." + hint)
            return
        _SitePairDialog(self, unlisted, rows, self._after_pair)

    @staticmethod
    def _pairable(unused):
        """Rows worth pairing — a broken token stays broken when paired."""
        return [e for e in unused if not e["problem"]]

    @classmethod
    def _pair_rows(cls, plan):
        """The rows of the file a site it doesn't name can take: those no
        source site matched, and those outside the chosen sites."""
        return cls._pairable((plan.get("unused") or [])
                             + (plan.get("spare") or []))

    def _after_pair(self, pairs):
        paired = pair_file_rows(pairs)
        for entry, site in pairs:
            cli_log(f"Paired row {entry['row']} of the file "
                    f"('{entry['dest']}') with source site '{site['site']}' "
                    f"({site['account']}).", "info")
        self.app.log_audit("agent_site_tokens_paired", rows=paired)
        self._dest_version += 1
        self._forget_plan_view()
        self._run_match()

    # ── step 3 · migrate, and watch it happen ─────────────────────────

    def _run_live(self, scheduled: bool = False):
        def refuse(title, text):
            if scheduled:
                cli_log(f"Scheduled migration did not start: {text}", "error")
            else:
                messagebox.showwarning(title, text)

        if not self._plan:
            refuse("Check first", "Run step 2 so you can see what will move.")
            return
        if self._plan_stale():
            refuse("Check again", "Step 1 changed since the plan was made — "
                                  "check again first.")
            return
        if self._report is not None:
            refuse("Check again", "This plan has already been migrated — "
                                  "check again to see what's left.")
            return
        if self._busy:
            refuse("Busy", "Something else is still running — wait for it "
                           "to finish.")
            return
        try:
            api = self._source_api()
        except ValueError as e:
            refuse("Cannot migrate", str(e))
            return
        planned = self._plan.get("agents", 0)
        if not planned:
            refuse("Nothing to move", "The plan has no agents to move.")
            return

        label, where = self._dest_source()
        source_url = str(getattr(api, "base_url", "") or "")
        guide = getattr(self.app, "guide", None)
        if (not scheduled and guide is not None
                and self._plan.get("mode") != "file"
                and not guide.allow("agents")):
            return
        if not scheduled and not messagebox.askyesno(
                "Migrate for real",
                f"Move {planned:,} agent(s) in {self._plan.get('ready', 0):,} "
                f"group(s)\n\nfrom   {self._console_name(api)}\ninto     "
                f"{self._dest_name()}\n\n{label}: {where}\n\nThis cannot be "
                f"undone from here — agents would have to be migrated back. "
                f"Continue?"):
            return

        rows = list(self._plan["rows"])
        verify = bool(self.verify_var.get())
        self._stop.clear()
        _retext(self.live_stop_btn, state="normal")
        self._busy = "live"
        self._report = None
        self._error.pop("live", None)
        self._live_counts = {"moved": 0, "pending": 0, "failed": 0,
                             "sent": 0, "total": planned}
        self._live_where = ""
        self.live_feed.clear("Waiting for the first batch…")
        self.live_feed.set_follow(True)
        self.live_filter.set("All agents")
        self.live_feed.set_filter("all")
        _retext(self._agent_detail, text="")
        self.cmd_log.clear()
        self.live_prog.set(0)
        self._run_started = time.monotonic()
        self._tick_timer()
        self._show_tab("agents")
        self._refresh_gates()
        cli_log("LIVE agent migration starting…", "cmd")

        pump = _EventPump(self, self._on_live_event,
                          on_flush=self._paint_live_view)
        pump.start()

        def do():
            return run_live_migration(
                api, rows, on_event=pump.push,
                should_stop=self._stop.is_set, verify=verify)

        def done(run):
            pump.drain()
            pump.stop()
            self._busy = None
            self._report = run
            self._last_run = {"run": run, "label": label, "value": where,
                              "source": source_url}
            self._paint_live_view(force=True)
            self._finish_live(run, scheduled)

        def fail(e):
            pump.stop()
            self._busy = None
            self._error["live"] = str(e)
            self._paint_live_view(force=True)
            self._refresh_gates()
            cli_log(f"The migration failed: {e}", "error")

        run_async(self, do, done, fail)

    def _on_live_event(self, event):
        """One event from the running migration, on the main thread.
        Records state only — drawing happens once per batch."""
        kind = event["type"]
        if kind == "start":
            self._live_counts["total"] = event["agents"]
        elif kind == "group":
            row = event["row"]
            self._live_where = (f"group {event['index']} of {event['total']}"
                                f" — {where_of(row)}  ──▶  "
                                f"{row['Destination']}")
        elif kind == "request":
            flt = event["filter"]
            self.cmd_log.add(
                f"▸ POST {event['endpoint']}  ids={event['count']}  "
                f"groupIds={flt.get('groupIds')}  token={event['token']}"
                + (f"  batch {event['batch']}/{event['of']}"
                   if event["of"] > 1 else ""))
        elif kind == "response":
            if event["ok"]:
                self.cmd_log.add(f"  ✓ accepted {event['affected']} of "
                                 f"{event['of']}")
            else:
                self.cmd_log.add(f"  ✕ {event['error']}")
        elif kind == "error":
            self.cmd_log.add(f"  ✕ {event['where']}: {event['error']}")
        elif kind == "roster":
            for rec in event["recs"]:
                self.live_feed.put(rec)
        elif kind == "agent":
            rec = event["rec"]
            self.live_feed.put(rec)
            state = rec["Result"]
            if state in ("moved", "pending", "failed"):
                self._live_counts[state] += 1
                self._live_counts["sent"] += 1

    def _paint_live_view(self, force: bool = False):
        """Repaint once per batch of events, not once per agent.

        The feed is a fixed pool and costs almost nothing, so it follows
        every batch. The counters and the progress line each cost tens of
        milliseconds — CustomTkinter redraws far more than the widget you
        touched — so they are held to a few times a second. `force` is
        used at the end so the final numbers are never a frame behind.
        """
        self.cmd_log.flush()
        self.live_feed.render()
        now = time.monotonic()
        if force or now - self._last_counts_paint >= self.COUNTS_EVERY:
            self._last_counts_paint = now
            self._paint_live_counts()

    def _paint_live_counts(self):
        c = self._live_counts
        done = c["moved"] + c["pending"] + c["failed"]
        self.live_prog.set(done / c["total"] if c["total"] else 0)
        if self._busy == "live":
            _retext(self._live_line, text=self._live_progress_text(),
                    text_color=TEXT)
        self._paint_list_counts()

    def _tick_timer(self):
        self._timer_job = None
        secs = int(time.monotonic() - self._run_started)
        _retext(self._timer_lbl, text=f"{secs // 60:02d}:{secs % 60:02d}")
        if self._busy == "live":
            self._timer_job = self.after(1000, self._tick_timer)

    def _finish_live(self, run, scheduled):
        moved, pending = run["moved"], run["pending"]
        failed, errors = run["failed"], len(run["errors"])
        self._tick_timer()
        # The run is over, so stop chasing the tail and read from the top.
        self.live_feed.set_follow(False)
        if failed:
            self.live_filter.set("Failed")
            self.live_feed.set_filter("failed")
        else:
            self.live_feed.render()
        self._refresh_gates()
        level = "warning" if (failed or errors) else "success"
        cli_log(f"Migration finished: {moved} moved, {pending} pending, "
                f"{failed} failed.", level)
        self.log.log(live_summary_text(run))
        self.app.log_audit(
            "agent_mapped_migrate", dry_run=False,
            candidates=len(run["agents"]), moved=moved, pending=pending,
            failed=failed, errors=errors,
            tokens="file" if (self._plan or {}).get("mode") == "file"
            else "map")
        if not scheduled:
            messagebox.showinfo(
                "Migration finished",
                f"{moved} agent(s) confirmed on the destination.\n"
                f"{pending} accepted, waiting to check in.\n"
                f"{failed} failed.\n\nExport results in step 3 keeps the "
                f"record; the Status report confirms the final picture once "
                f"the agents have re-registered.")

    def _on_live_filter(self, value):
        self.live_feed.set_filter(self.LIVE_FILTERS.get(value, "all"))
        self._paint_list_counts()

    def _on_agent_pick(self, rec):
        _retext(self._agent_detail, text=(
            f"{rec.get('Computer', '')}   ·   {rec.get('OS', '')}   ·   "
            f"{rec.get('Account', '')} › {rec.get('Site', '')} › "
            f"{rec.get('Group', '')}  ──▶  {rec.get('Destination', '')}\n"
            f"{rec.get('Result', '')}: {rec.get('Detail') or '—'}   "
            f"(agent id {rec.get('Agent ID') or rec.get('id', '')})"))

    def _export_live(self):
        last = self._last_run or {}
        run = last.get("run") or {}
        if not run.get("agents"):
            messagebox.showinfo("Nothing to export", "Run the migration "
                                                     "first.")
            return
        ts = datetime.now().strftime("%Y%m%d-%H%M")
        export_agent_report(
            build_live_report(run, source=last.get("source", ""),
                              map_path=last.get("value", ""),
                              map_label=last.get("label")
                              or "Destination map file"),
            default_name=f"s1-agent-live-run-{ts}")

    # ── scheduling ────────────────────────────────────────────────────

    def _open_schedule(self):
        if not self._plan or self._plan_stale() or self._report is not None:
            messagebox.showwarning("Check first", "Run step 2 first, so the "
                                   "scheduled run has a checked plan.")
            return
        _ScheduleDialog(self)

    def _arm(self, date_str, time_str, parent=None):
        """Arm the unattended run. True when armed."""
        where = {"parent": parent} if parent is not None else {}
        try:
            when = parse_schedule(date_str, time_str)
        except ValueError as e:
            messagebox.showwarning("Check the date and time", str(e), **where)
            return False
        if not self._plan or self._plan_stale() or self._report is not None:
            messagebox.showwarning(
                "Check first",
                "Run step 2 so the scheduled run has a checked plan — and so "
                "you know how many agents it will move.", **where)
            return False
        if not messagebox.askyesno(
                "Arm the scheduler",
                f"At {when:%Y-%m-%d %H:%M} this will move "
                f"{self._plan.get('agents', 0):,} agent(s) for real.\n\n"
                f"Keep the app open and the machine awake until then. "
                f"Continue?", **where):
            return False
        self._sched_when = when
        cli_log(f"Agent migration scheduled for {when:%Y-%m-%d %H:%M}.",
                "info")
        self._tick()
        return True

    def _disarm(self):
        self._cancel_schedule()
        cli_log("Scheduled agent migration cancelled.", "warning")
        self._refresh_gates()

    def _cancel_schedule(self):
        if self._sched_job is not None:
            try:
                self.after_cancel(self._sched_job)
            except Exception:
                pass
        self._sched_job = None
        self._sched_when = None

    def _tick(self):
        self._sched_job = None
        if self._sched_when is None:
            return
        if (self._sched_when - datetime.now()).total_seconds() <= 0:
            self._cancel_schedule()
            cli_log("Scheduled agent migration firing now.", "cmd")
            self._run_live(scheduled=True)
            self._refresh_gates()
            return
        self._refresh_gates()
        self._sched_job = self.after(1000, self._tick)

    # ── the status report ─────────────────────────────────────────────

    def _on_passphrase_toggle(self):
        if self.pass_var.get():
            self.log.log(
                "Passphrases: one API call per agent (slow on a large scope), "
                "each fetch is written to the source console's activity log, "
                "and the exported report becomes sensitive.")

    @staticmethod
    def _status_colour(rec):
        return {"Failed": ACCENT, "Pending": WARN}.get(
            rec.get("Migration Status"))

    def _run_status(self):
        if self._busy:
            return
        try:
            api = self._source_api()
        except ValueError as e:
            messagebox.showwarning("Console", str(e))
            return
        account_ids = self.src_accounts.get().strip()
        site_id = self.src_sites.get().strip()
        if not account_ids and not site_id and not messagebox.askyesno(
                "Whole console",
                "Step 1 names no account or site, so this reports on EVERY "
                "agent the source token can see. On a large console that is "
                "a lot of paging.\n\nContinue?"):
            return
        fetch = bool(self.pass_var.get())
        if fetch and not messagebox.askyesno(
                "Fetch passphrases?",
                "Including passphrases reads each agent's uninstall passphrase "
                "from the source console.\n\n"
                "•  One API call per agent — this can take a long time on a "
                "large scope.\n"
                "•  SentinelOne records every passphrase fetch in the source "
                "console's ACTIVITY LOG.\n"
                "•  The exported report will contain secrets — treat it as "
                "sensitive.\n\nContinue?"):
            return

        self._stop.clear()
        self._busy = "status"
        self.status_stats.clear()
        self.status_table.clear("Counting…")
        self._show_tab("status")
        self._refresh_gates()

        def do():
            return count_agents_by_status(
                api, account_ids, site_id, fetch, log=self._say,
                should_stop=self._stop.is_set)

        def done(result):
            self._busy = None
            self._status = result
            counts = result.get("counts") or {}
            decom = len(result.get("decommissioned") or [])
            self.status_stats.show([
                (counts.get("Migrated", 0), "migrated", GREEN),
                (counts.get("Pending", 0), "pending", WARN),
                (counts.get("Failed", 0), "failed",
                 ACCENT if counts.get("Failed") else TEXT_MUTED),
                (counts.get("N/A", 0), "not applicable", TEXT_MUTED),
                (decom, "decommissioned", TEXT_MUTED),
            ])
            rows = list(result.get("agents") or []) + \
                list(result.get("decommissioned") or [])
            if rows:
                self.status_table.set_rows(rows)
            else:
                self.status_table.clear("No agents in that scope.")
            self._refresh_gates()
            self.log.log(status_summary_text(result))
            cli_log(f"Status report: {sum(counts.values())} agent(s).",
                    "success")

        def fail(e):
            self._busy = None
            self.status_table.clear(f"The report failed: {e}")
            self._refresh_gates()

        run_async(self, do, done, fail)

    def _export_status(self):
        result = self._status or {}
        if not (result.get("agents") or result.get("decommissioned")):
            messagebox.showinfo("Nothing to export", "Run the report first.")
            return
        scope = (self.src_sites.get().strip()
                 or self.src_accounts.get().strip() or "whole console")
        ts = datetime.now().strftime("%Y%m%d-%H%M")
        export_agent_report(
            build_status_report(
                result, console=str(getattr(self.app.source_api, "base_url",
                                            "") or ""),
                scope=scope, passphrases=bool(self.pass_var.get())),
            default_name=f"s1-agent-status-{ts}")

    def on_show(self):
        setter = getattr(self.app, "set_active_console", None)
        if callable(setter):
            setter("")
        if (self._plan is None and self._busy is None
                and not self.src_accounts.get().strip()
                and not self.src_sites.get().strip()):
            self._restore_last_plan()
        self._refresh_gates()


class _NameTree(_VirtualTable):
    """The source scope as a tree, beside what the destination has.

    Rows come flattened from `NameComparison.visible`. A click on the
    ▸/▾ — or a double-click anywhere on the row — folds a branch.
    """

    TONE = {"same": GREEN, "fixed": BRAND, "unfixed": WARN, "chosen": BRAND,
            "missing": ACCENT, "waiting": TEXT_FAINT}
    LOUD = ("missing", "chosen")

    def __init__(self, master, on_select, on_toggle, columns=None, tone=None,
                 loud=None, empty="Nothing in the source scope.", **kw):
        self.tone = tone or self.TONE
        self.loud = tuple(loud or self.LOUD)
        super().__init__(master, columns or [
            _col("On the source", "_name", 5),
            _col("", "_kind", 0, minsize=58),
            _col("Agents", "_agents", 0, anchor="e", mono=True, minsize=64),
            _col("On the destination", "_dest", 6)],
            empty=empty, on_select=on_select, **kw)
        self._on_toggle = on_toggle

    def _apply_scale(self, scale=None):
        super()._apply_scale(scale)
        self._ruler = None

    def width_of(self, text) -> int:
        """Pixels `text` takes in the tree's font."""
        if getattr(self, "_ruler", None) is None:
            family, size = self._fonts[(False, False)][:2]
            self._ruler = tkfont.Font(root=self, family=family, size=size)
        return self._ruler.measure(text)

    def _make_row(self, slot):
        row = super()._make_row(slot)
        row["cells"][0].bind("<Button-1>", lambda e, s=slot: self._press(s, e))
        for widget in [row["frame"]] + row["cells"]:
            widget.bind("<Double-Button-1>", lambda _e, s=slot: self._fold(s))
        return row

    def _rec(self, slot):
        key = self._pool[slot]["key"]
        return None if key is None else self._records.get(key)

    def _press(self, slot, event):
        rec = self._rec(slot)
        if rec and rec["_branch"] and event.x <= rec["_zone"]:
            self._on_toggle(rec["node"])
            return
        self._click(slot)

    def _fold(self, slot):
        rec = self._rec(slot)
        if rec and rec["_branch"]:
            self._on_toggle(rec["node"])

    def cell(self, rec, key):
        state = rec["state"]
        if key == "_dest":
            return rec["_dest"], self.tone.get(state, TEXT)
        if key in ("_kind", "_agents"):
            return rec[key], TEXT_FAINT if key == "_kind" else TEXT_MUTED
        return rec["_name"], (TEXT if state in self.loud
                              else TEXT_FAINT if state == "waiting"
                              else TEXT_MUTED)

    def refill(self, rows, select=None):
        """New rows, keeping the scroll position and the given selection."""
        offset = self._offset
        self.set_rows(rows)
        self._selected = select if select in self._records else None
        self._offset = offset
        self.render()

    def reveal(self, key):
        """Scroll just enough to show the row `key`."""
        if self._dirty:
            self._rebuild()
        if key not in self._view:
            return
        i = self._view.index(key)
        if i < self._offset:
            self._offset = i
        elif i >= self._offset + self._fit:
            self._offset = i - self._fit + 1
        self.render()


class _NameFixDialog(ctk.CTkToplevel):
    """Compare the source with the destination, and fix what doesn't pair.

    Left: the source scope as a tree — account › site › group — beside
    what the destination has in the same place. Right: the scope you
    select — why it doesn't pair, every destination scope in that place
    and what already pairs with each, where else its name turns up — and
    the pick. Nothing is guessed and nothing is sent: the picks become
    name fixes, applied when the plan is checked again.
    """

    TITLE = "Fix names — source vs destination"
    FILTERS = {"Needs a fix": "fix", "All": "all", "Fixed": "done"}
    MARK = {"same": "✓", "fixed": "✎", "unfixed": "↺", "chosen": "→",
            "missing": "✕", "waiting": "·"}
    TONE = _NameTree.TONE
    LOUD = _NameTree.LOUD
    COLUMNS = None
    EMPTY = "Nothing in the source scope."
    ATTENTION = ("missing",)
    COLLAPSE = ("missing",)
    PICKABLE = ("missing", "chosen")
    OTHER = ("fixed", "unfixed", "waiting")
    UNFIX = ("fixed", "unfixed")
    SEARCH_HINT = "Search source or destination names…"
    CAND_HINT = "Filter — or type the exact destination name"
    CAND_HEADS = ("Destination", "Already pairs with")
    CAND_EMPTY = "Nothing was read in this place."
    USE_IDLE = "Use selected"
    CLEAR_TEXT = "Undo pick"
    FOOTNOTE = "Picks are kept only when you apply them."
    APPLY_IDLE = "Apply & check again"

    def __init__(self, parent, comparison, on_done):
        super().__init__(parent)
        self.title(self.TITLE)
        width = min(1180, max(900, self.winfo_screenwidth() - 120))
        height = min(740, max(560, self.winfo_screenheight() - 160))
        top = parent.winfo_toplevel()
        x = top.winfo_rootx() + (top.winfo_width() - width) // 2
        y = top.winfo_rooty() + (top.winfo_height() - height) // 3
        self.geometry(f"{width}x{height}+{max(x, 0)}+{max(y, 0)}")
        self.minsize(900, 560)
        self.cmp = comparison
        self._on_done = on_done
        self._selected = None
        self._pick = None
        self._cands = []
        self._zones = {}
        self._collapsed = {k for k, n in comparison.nodes.items()
                           if n["children"]
                           and comparison.state(k) in self.COLLAPSE}
        self.grid_columnconfigure(0, weight=1, uniform="pane", minsize=420)
        self.grid_columnconfigure(1, weight=1, uniform="pane", minsize=420)
        self.grid_rowconfigure(1, weight=1)
        self._build_top()
        self.tree = _NameTree(self, self._on_node, self._toggle,
                              columns=self.COLUMNS, tone=self.TONE,
                              loud=self.LOUD, empty=self.EMPTY)
        self.tree.grid(row=1, column=0, sticky="nsew", padx=(16, 8),
                       pady=(0, 10))
        self._build_detail()
        self._build_bottom()
        self._leaf = self._pad_like("▾  ")
        modes = {mode: label for label, mode in self.FILTERS.items()}
        summary = comparison.summary()
        self._filter.set(modes["fix"] if any(summary.get(s)
                                             for s in self.ATTENTION)
                         else modes["all"])
        self._refresh()
        self._select(self._next_missing())
        self.bind("<Escape>", lambda _e: self.destroy())
        self.transient(parent.winfo_toplevel())
        self.after(120, self._grab)

    def _intro(self):
        """The line above the tree, and its colour."""
        return ("Your source scope — account › site › group — beside what "
                "the destination has in the same place. ✕ means the "
                "destination has nothing of that name there: select it to "
                "see what the destination does have and what already pairs "
                "with each, then pick the one it really is. Nothing is sent "
                "— picks are applied when the plan is checked again, and "
                "re-applied every time the destination is read.", TEXT_MUTED)

    def _build_top(self):
        head = ctk.CTkFrame(self, fg_color="transparent")
        head.grid(row=0, column=0, columnspan=2, sticky="ew", padx=16,
                  pady=(14, 8))
        head.grid_columnconfigure(0, weight=1)
        text, tone = self._intro()
        intro = self._intro_lbl = ctk.CTkLabel(
            head, anchor="w", justify="left", font=(UI_FONT, 12),
            text_color=tone, text=text)
        intro.grid(row=0, column=0, sticky="ew")
        _wrap(head, intro)
        bar = ctk.CTkFrame(head, fg_color="transparent")
        bar.grid(row=1, column=0, sticky="ew", pady=(10, 0))
        bar.grid_columnconfigure(2, weight=1)
        self._search = ctk.CTkEntry(
            bar, width=300, height=32, placeholder_text=self.SEARCH_HINT)
        self._search.grid(row=0, column=0, sticky="w")
        self._search.bind("<KeyRelease>", lambda _e: self._refresh())
        self._filter = ctk.CTkSegmentedButton(
            bar, values=list(self.FILTERS), height=30, font=(UI_FONT, 12),
            command=lambda _v: self._refresh())
        self._filter.grid(row=0, column=1, sticky="w", padx=(10, 0))
        self._counts = ctk.CTkLabel(bar, text="", anchor="e",
                                    font=(UI_FONT, 12), text_color=TEXT_MUTED)
        self._counts.grid(row=0, column=2, sticky="e", padx=(10, 0))

    def _build_detail(self):
        card = ctk.CTkFrame(self, fg_color=CARD, corner_radius=RADIUS_LG,
                            border_width=1, border_color=BORDER)
        card.grid(row=1, column=1, sticky="nsew", padx=(8, 16), pady=(0, 10))
        card.grid_columnconfigure(0, weight=1)
        card.grid_rowconfigure(5, weight=1)
        self._d_kind = ctk.CTkLabel(card, text="", anchor="w", height=14,
                                    font=(UI_FONT, 10, "bold"),
                                    text_color=TEXT_FAINT)
        self._d_kind.grid(row=0, column=0, sticky="ew", padx=16, pady=(14, 0))
        self._d_name = ctk.CTkLabel(card, text="", anchor="w", justify="left",
                                    font=(UI_FONT, 15, "bold"))
        self._d_name.grid(row=1, column=0, sticky="ew", padx=16)
        self._d_path = ctk.CTkLabel(card, text="", anchor="w", justify="left",
                                    font=(UI_FONT, 12), text_color=TEXT_MUTED)
        self._d_path.grid(row=2, column=0, sticky="ew", padx=16)
        self._d_state = ctk.CTkLabel(card, text="", anchor="w",
                                     justify="left", font=(UI_FONT, 12))
        self._d_state.grid(row=3, column=0, sticky="ew", padx=16,
                           pady=(10, 0))
        self._d_else = ctk.CTkLabel(card, text="", anchor="w", justify="left",
                                    font=(UI_FONT, 11), text_color=TEXT_MUTED)
        self._d_else.grid(row=4, column=0, sticky="ew", padx=16, pady=(6, 0))
        for label in (self._d_name, self._d_path, self._d_state,
                      self._d_else):
            _wrap(card, label, pad=32)

        box = self._pick_box = ctk.CTkFrame(card, fg_color="transparent")
        box.grid(row=5, column=0, sticky="nsew", padx=16, pady=(12, 0))
        box.grid_columnconfigure(0, weight=1)
        box.grid_rowconfigure(2, weight=1)
        self._c_title = ctk.CTkLabel(box, text="", anchor="w",
                                     font=(UI_FONT, 12, "bold"))
        self._c_title.grid(row=0, column=0, sticky="ew")
        self._c_search = ctk.CTkEntry(box, height=30,
                                      placeholder_text=self.CAND_HINT)
        self._c_search.grid(row=1, column=0, sticky="ew", pady=6)
        self._c_search.bind("<KeyRelease>",
                            lambda _e: self._filter_candidates())
        self._c_search.bind("<Return>", lambda _e: self._use_best())
        self._c_list = _VirtualTable(
            box, [_col(self.CAND_HEADS[0], "shown", 5),
                  _col(self.CAND_HEADS[1], "used", 4)],
            empty=self.CAND_EMPTY,
            on_select=self._on_candidate,
            colour=lambda r: None if r["_free"] else TEXT_FAINT)
        self._c_list.grid(row=2, column=0, sticky="nsew")
        self._c_note = ctk.CTkLabel(box, text="", anchor="w", justify="left",
                                    font=(UI_FONT, 11))
        self._c_note.grid(row=3, column=0, sticky="ew", pady=(6, 0))
        _wrap(box, self._c_note, pad=4)
        acts = ctk.CTkFrame(box, fg_color="transparent")
        acts.grid(row=4, column=0, sticky="w", pady=(8, 14))
        self._use_btn = ctk.CTkButton(
            acts, text=self.USE_IDLE, width=150, height=32,
            corner_radius=RADIUS_MD, font=(UI_FONT, 12, "bold"),
            border_color=BORDER, text_color_disabled=TEXT_FAINT,
            command=self._use, **_GHOSTED)
        self._use_btn.grid(row=0, column=0, padx=(0, 6))
        self._clear_btn = _ghost(acts, self.CLEAR_TEXT, self._clear, width=110)
        self._clear_btn.grid(row=0, column=1)
        self._typed_btn = _ghost(acts, "Use as typed", self._use_typed,
                                 width=150)
        self._typed_btn.grid(row=1, column=0, columnspan=2, sticky="w",
                             pady=(6, 0))

        self._other = ctk.CTkFrame(card, fg_color="transparent")
        self._other.grid(row=6, column=0, sticky="w", padx=16, pady=(12, 14))
        self._unfix_btn = _ghost(self._other, "Remove this fix",
                                 self._toggle_fix, width=150)
        self._unfix_btn.grid(row=0, column=0, padx=(0, 6))
        self._goto_btn = _ghost(self._other, "Show the site",
                                self._goto_parent, width=150)
        self._goto_btn.grid(row=0, column=1)

    def _build_bottom(self):
        bar = ctk.CTkFrame(self, fg_color="transparent")
        bar.grid(row=2, column=0, columnspan=2, sticky="ew", padx=16,
                 pady=(0, 14))
        bar.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(bar, text=self.FOOTNOTE,
                     anchor="w", font=(UI_FONT, 11),
                     text_color=TEXT_FAINT).grid(row=0, column=0, sticky="ew")
        _ghost(bar, "Cancel", self.destroy, width=100, height=34).grid(
            row=0, column=1, padx=(0, 8))
        self._apply_btn = ctk.CTkButton(
            bar, text=self.APPLY_IDLE, width=240, height=34,
            corner_radius=RADIUS_MD, font=(UI_FONT, 13, "bold"),
            border_color=BORDER, text_color_disabled=TEXT_FAINT,
            command=self._apply, **_GHOSTED)
        self._apply_btn.grid(row=0, column=2)

    def _grab(self):
        try:
            self.lift()
            self.focus_force()
            self.grab_set()
        except Exception:
            pass            # closed before it was ever shown

    # ── the tree ──────────────────────────────────────────────────────

    def _mode(self):
        return self.FILTERS.get(self._filter.get(), "all")

    def _pad_like(self, text):
        """Spaces about as wide as `text`, so leaves line up with branches."""
        width, pad = self.tree.width_of(text), ""
        while len(pad) < 12 and self.tree.width_of(pad + " ") <= width:
            pad += " "
        return pad

    def _zone(self, prefix):
        key = (prefix, self.tree._scale)
        if key not in self._zones:
            self._zones[key] = self.tree.width_of(prefix) + 2
        return self._zones[key]

    def _order(self):
        return self.cmp.visible(self._mode(), self._search.get(),
                                self._collapsed)

    def _refresh(self):
        cmp = self.cmp
        searching = bool(self._search.get().strip())
        rows = []
        for key, depth in self._order():
            node = cmp.nodes[key]
            branch = bool(node["children"])
            mark = ((("▸  " if key in self._collapsed and not searching
                      else "▾  ")) if branch else self._leaf)
            prefix = "      " * depth + mark
            rows.append({"key": "\x1f".join(key), "node": key,
                         "_branch": branch, "_zone": self._zone(prefix),
                         "_name": prefix + node["name"], "_kind": node["kind"],
                         "_agents": self._count_text(key),
                         "_dest": cmp.dest_text(key),
                         "state": cmp.state(key)})
        self.tree.refill(rows, "\x1f".join(self._selected)
                         if self._selected else None)
        self._paint_counts()

    def _count_text(self, key):
        agents = self.cmp.agents(key)
        return f"{agents:,}" if agents else "—"

    def _count_bits(self):
        s = self.cmp.summary()
        bits = [f"{s['missing']:,} to fix", f"{s['chosen']:,} picked"]
        if s["fixed"] or s["unfixed"]:
            bits.append(f"{s['fixed'] + s['unfixed']:,} fixed earlier")
        bits.append(f"{s['same']:,} pair by name")
        if s["waiting"]:
            bits.append(f"{s['waiting']:,} waiting for their site")
        return bits

    def _changes(self):
        return self.cmp.summary()["chosen"] + len(self.cmp.removed())

    def _apply_text(self, changes):
        return f"Apply {_plural(changes, 'change')} & check again"

    def _paint_counts(self):
        _retext(self._counts, text="   ·   ".join(self._count_bits()))
        changes = self._changes()
        _retext(self._apply_btn, state="normal" if changes else "disabled",
                text=self._apply_text(changes) if changes else self.APPLY_IDLE,
                **(_GO if changes else _GHOSTED))

    def _toggle(self, key):
        if key in self._collapsed:
            self._collapsed.discard(key)
        else:
            self._collapsed.add(key)
        self._refresh()

    def _on_node(self, rec):
        self._select(rec["node"], scroll=False)

    def _select(self, key, scroll=True):
        self._selected = key if key in self.cmp.nodes else None
        row = "\x1f".join(key) if self._selected else None
        self.tree._selected = row if row in self.tree._records else None
        if scroll and self.tree._selected:
            self.tree.reveal(row)
        else:
            self.tree.render()
        self._show_detail()

    def _next_missing(self, after=None):
        order = [k for k, _d in self._order()]
        if after in order:
            i = order.index(after)
            order = order[i + 1:] + order[:i]
        return next((k for k in order if self.cmp.state(k) in self.ATTENTION),
                    None)

    # ── the selected scope ────────────────────────────────────────────

    def _show_detail(self):
        key, cmp = self._selected, self.cmp
        if key is None:
            for label, text in ((self._d_kind, ""),
                                (self._d_name, "Nothing selected"),
                                (self._d_path, "Select a scope on the left "
                                               "to compare it."),
                                (self._d_state, ""), (self._d_else, "")):
                _retext(label, text=text)
            _grid_shown(self._d_else, False)
            _grid_shown(self._pick_box, False)
            _grid_shown(self._other, False)
            return
        node, state = cmp.nodes[key], cmp.state(key)
        _retext(self._d_kind, text=self._detail_head(node["kind"]))
        _retext(self._d_name, text=node["name"])
        _retext(self._d_path, text="   ·   ".join(self._detail_bits(key)))
        _retext(self._d_state, text=f"{self.MARK[state]}  {cmp.reason(key)}",
                text_color=self.tree.tone.get(state, TEXT))
        text = self._elsewhere_text(key, state)
        _retext(self._d_else, text=text)
        _grid_shown(self._d_else, bool(text))
        pickable = state in self.PICKABLE
        _grid_shown(self._pick_box, pickable)
        _grid_shown(self._other, state in self.OTHER)
        _grid_shown(self._unfix_btn, state in self.UNFIX)
        _retext(self._unfix_btn, text="Keep this fix" if state == "unfixed"
                else "Remove this fix")
        _grid_shown(self._goto_btn, state == "waiting")
        if state == "waiting":
            _retext(self._goto_btn, text="Show the "
                    + cmp.nodes[node["parent"]]["kind"])
        if pickable:
            self._fill_candidates(key)

    def _detail_head(self, kind):
        return f"SOURCE {kind.upper()}"

    def _detail_bits(self, key):
        cmp = self.cmp
        node = cmp.nodes[key]
        kind = node["kind"]
        where = " › ".join(node["path"][:-1])
        bits = [f"in {where}" if where else "a source account"]
        if node["children"]:
            bits.append(_plural(len(node["children"]),
                                "group" if kind == "site" else "site"))
        if cmp.agents(key):
            bits.append(_plural(cmp.agents(key), "agent"))
        return bits

    def _elsewhere_text(self, key, state):
        node = self.cmp.nodes[key]
        kind = node["kind"]
        seen = self.cmp.elsewhere(key) if state in self.PICKABLE else []
        if not seen:
            return ""
        more = f" and {len(seen) - 5:,} more" if len(seen) > 5 else ""
        return (f"Elsewhere on the destination, a {kind} named "
                f"“{node['name']}” is in "
                + ", ".join(f"“{w}”" for w in seen[:5]) + more
                + f". A {kind} only pairs inside its own "
                  f"{'account' if kind == 'site' else 'site'}.")

    def _cand_title(self, kind):
        return f"Pick the destination {kind} it really is"

    def _cand_used(self, c, kind):
        if c["taken_by"]:
            return (f"source {kind} “{c['taken_by']}”"
                    + (" · fixed earlier" if c["fixed"] else ""))
        if c["chosen_for"]:
            return f"your pick for “{c['chosen_for']}”"
        return "nothing — free"

    def _fill_candidates(self, key):
        cmp, kind = self.cmp, self.cmp.nodes[key]["kind"]
        self._cands = cmp.candidates(key)
        _retext(self._c_title, text=self._cand_title(kind))
        self._c_list._heads[0].configure(text=f"DESTINATION {kind.upper()}S")
        rows = []
        for c in self._cands:
            used = self._cand_used(c, kind)
            rows.append({"key": c["name"], "shown": c["shown"], "used": used,
                         "_free": not (c["taken_by"] or c["chosen_for"]),
                         "cand": c})
        self._c_list.set_rows(rows)
        self._c_list.filter_by(None)
        if self._c_search.get():
            self._c_search.delete(0, "end")
        self._pick = None
        current = cmp.choices.get(key)
        if current and cmp.typed(key):
            self._c_search.insert(0, current)
            self._filter_candidates()
            return
        if current:
            self._pick = next((c for c in self._cands
                               if c["name"] == current), None)
            self._c_list._selected = current if self._pick else None
            self._c_list.render()
        self._paint_pick()

    def _filter_candidates(self):
        text = norm_name(self._c_search.get())
        self._c_list.filter_by((lambda r, t=text: t in norm_name(r["shown"]))
                               if text else None)
        self._paint_pick()

    def _on_candidate(self, rec):
        self._pick = rec["cand"]
        self._paint_pick()

    def _paint_pick(self):
        key, cmp = self._selected, self.cmp
        if key is None:
            return
        typed = " ".join(self._c_search.get().split())
        verdict, holder = cmp.check(key, typed) if typed else ("", "")
        pick = self._pick
        free = bool(pick) and not (pick["taken_by"] or pick["chosen_for"])
        _retext(self._use_btn, state="normal" if free else "disabled",
                text=self._use_text(pick) if pick else self.USE_IDLE,
                **(_PRIMARY if free else _GHOSTED))
        unknown = bool(typed) and verdict == "unknown"
        _grid_shown(self._typed_btn, unknown)
        if unknown:
            _retext(self._typed_btn, text=self._typed_text(typed))
        _grid_shown(self._clear_btn, cmp.state(key) == "chosen")
        note, tone = self._pick_note(key, pick, typed, verdict, holder)
        _retext(self._c_note, text=note, text_color=tone)

    def _use_text(self, pick):
        return f"Use “{_short(pick['shown'], 24)}”"

    def _typed_text(self, typed):
        return f"Use “{_short(typed, 20)}” as typed"

    def _pick_note(self, key, pick, typed, verdict, holder):
        """The line under the destination list, and its colour."""
        cmp = self.cmp
        node = cmp.nodes[key]
        kind, place = node["kind"], cmp.where(key)
        unknown = bool(typed) and verdict == "unknown"
        spare = sum(1 for c in self._cands
                    if not (c["taken_by"] or c["chosen_for"]))
        if pick and pick["taken_by"]:
            who = pick["taken_by"]
            note, tone = (f"“{pick['shown']}” already pairs with the source "
                          f"{kind} “{who}”. A name fix renames that "
                          f"destination {kind}, so pointing “{node['name']}” "
                          f"at it as well would take it away from “{who}”. "
                          f"Pick another, or leave this one unmatched.",
                          ACCENT)
        elif pick and pick["chosen_for"]:
            note, tone = (f"“{pick['shown']}” is already your pick for "
                          f"“{pick['chosen_for']}” — one destination {kind} "
                          f"takes only one source {kind} here.", ACCENT)
        elif pick:
            note, tone = (f"“{pick['shown']}” is in {place} and nothing pairs "
                          f"with it yet.", GREEN)
        elif unknown:
            note = f"No {kind} named “{typed}” was read from {place}."
            if kind == "site" and cmp.partial_sites:
                note += (" Only the sites named in step 1 were read, so it may "
                         "still be there.")
            note += (" Use it only if you know the destination has exactly "
                     "this name.")
            tone = WARN
        elif typed and verdict == "taken":
            note, tone = (f"“{typed}” is in {place}, but it already pairs with "
                          f"“{holder}”.", ACCENT)
        elif typed:
            note, tone = (f"“{typed}” is in {place} and free — select it in "
                          f"the list.", GREEN)
        elif not self._cands:
            note = f"Nothing was read from {place} to pick from."
            if kind == "site" and cmp.partial_sites:
                note += (" Only the sites named in step 1 were read — type the "
                         "destination site's exact name.")
            tone = WARN
        else:
            note, tone = (f"{_plural(len(self._cands), kind)} in {place} · "
                          f"{spare:,} free. Closest names first; grey ones "
                          f"already pair with another source {kind}.",
                          TEXT_MUTED)
        return note, tone

    # ── picks ─────────────────────────────────────────────────────────

    def _use(self):
        key, pick = self._selected, self._pick
        if key is None or not pick or not self.cmp.choose(key, pick["shown"]):
            self._paint_pick()
            return
        self._collapsed.discard(key)
        self._changed(advance=True)

    def _use_best(self):
        """Return in the filter: the selected scope, or the only free one."""
        if self._pick and not (self._pick["taken_by"]
                               or self._pick["chosen_for"]):
            self._use()
            return
        self._c_list.counts()
        free = [self._c_list._records[k]["cand"] for k in self._c_list._view
                if self._c_list._records[k]["_free"]]
        if len(free) == 1:
            self._pick = free[0]
            self._use()

    def _use_typed(self):
        key = self._selected
        typed = " ".join(self._c_search.get().split())
        if key is not None and typed and self.cmp.choose(
                key, typed, allow_unknown=True):
            self._changed(advance=True)

    def _clear(self):
        if self._selected is not None:
            self.cmp.unchoose(self._selected)
            self._changed()

    def _toggle_fix(self):
        key = self._selected
        if key is not None:
            self.cmp.unfix(key, keep=self.cmp.state(key) == "unfixed")
            self._changed()

    def _goto_parent(self):
        key = self._selected
        if key is not None and self.cmp.nodes[key]["parent"]:
            self._select(self.cmp.nodes[key]["parent"])

    def _changed(self, advance=False):
        self._refresh()
        nxt = self._next_missing(self._selected) if advance else None
        self._select(nxt if nxt is not None else self._selected)

    def _apply(self):
        chosen, removed = self.cmp.chosen(), self.cmp.removed()
        self.destroy()
        if chosen or removed:
            self._on_done(chosen, removed)


class _PairSiteList(_VirtualTable):
    """The source sites the file doesn't name, and the row each one gets."""

    def __init__(self, master, picks, show_account=False, **kw):
        self._picks = picks
        self._show_account = show_account
        super().__init__(master, [
            _col("Source site", "_site", 5),
            _col("Agents", "_agents", 0, anchor="e", mono=True, minsize=56),
            _col("Gets the token of", "_pick", 4)], **kw)

    def cell(self, rec, key):
        site = rec["site"]
        if key == "_site":
            return (f"{site['site']}  ·  {site['account']}"
                    if self._show_account else site["site"]), TEXT
        if key == "_agents":
            return f"{int(site.get('agents') or 0):,}", TEXT_MUTED
        entry = self._picks.get(rec["key"])
        if entry is None:
            return "not paired yet", TEXT_FAINT
        return f"row {entry['row']}  ·  {entry['source']}", GREEN


class _PairRowList(_VirtualTable):
    """The rows of the file a site can take, closest name first. A
    double-click uses the row."""

    def __init__(self, master, owner, on_use, **kw):
        self._owner = owner
        self._on_use = on_use
        super().__init__(master, [
            _col("Row", "_row", 0, anchor="e", mono=True, minsize=44),
            _col("Site in the file", "_name", 4),
            _col("Lands on", "_dest", 4),
            _col("", "_note", 3)], **kw)

    def _make_row(self, slot):
        row = super()._make_row(slot)
        for widget in [row["frame"]] + row["cells"]:
            widget.bind("<Double-Button-1>",
                        lambda _e, s=slot: self._double(s))
        return row

    def _double(self, slot):
        if self._pool[slot]["key"] is not None:
            self._on_use()

    def cell(self, rec, key):
        entry = rec["entry"]
        taken = self._owner(entry)
        if key == "_row":
            return str(entry["row"]), TEXT_FAINT
        if key == "_name":
            return entry["source"], TEXT_FAINT if taken else TEXT
        if key == "_dest":
            return file_destination(entry), (TEXT_FAINT if taken
                                             else TEXT_MUTED)
        if taken:
            return f"picked for {taken}", TEXT_FAINT
        return ("≈ same name", GREEN) if rec["same"] else ("", TEXT_FAINT)


class _SitePairDialog(ctk.CTkToplevel):
    """Choose, for each source site the file doesn't name, the row of the
    file whose token its agents get.

    Partners often spell a site differently from the source —
    "shira-telecom" in the file, "Shira Telecom" on the console. The rows
    are offered closest name first, and the same name spelt differently is
    selected for a look — but nothing is paired until the operator uses a
    row, because a wrong guess sends one customer's agents into another
    customer's site. A row's token goes to one site only.
    """

    TITLE = "Pair sites — the source beside the file"

    def __init__(self, parent, sites, rows, on_done):
        super().__init__(parent)
        self.title(self.TITLE)
        width = min(1100, max(860, self.winfo_screenwidth() - 160))
        height = min(660, max(500, self.winfo_screenheight() - 200))
        top = parent.winfo_toplevel()
        x = top.winfo_rootx() + (top.winfo_width() - width) // 2
        y = top.winfo_rooty() + (top.winfo_height() - height) // 3
        self.geometry(f"{width}x{height}+{max(x, 0)}+{max(y, 0)}")
        self.minsize(860, 500)
        self._on_done = on_done
        self._sites = list(sites)
        self._rows = list(rows)
        self._picks = {}
        self._current = None
        self._row = None
        self.grid_columnconfigure(0, weight=2, uniform="pane")
        self.grid_columnconfigure(1, weight=3, uniform="pane")
        self.grid_rowconfigure(1, weight=1)

        head = ctk.CTkFrame(self, fg_color="transparent")
        head.grid(row=0, column=0, columnspan=2, sticky="ew", padx=16,
                  pady=(14, 10))
        head.grid_columnconfigure(0, weight=1)
        intro = ctk.CTkLabel(
            head, anchor="w", justify="left", font=(UI_FONT, 12),
            text_color=TEXT_MUTED,
            text="These source sites have agents but no row in the file "
                 "under their name, so they stay put. Select a site, then "
                 "pick the row of the file whose token its agents should get "
                 "— the file may spell it differently, with hyphens, other "
                 "capitals or extra words. Closest names come first, and a "
                 "row goes to one site. Step 2 checks again with your "
                 "choices.")
        intro.grid(row=0, column=0, sticky="ew")
        _wrap(head, intro)

        self.site_list = _PairSiteList(
            self, self._picks,
            show_account=len({s["account_id"] for s in self._sites}) > 1,
            on_select=lambda r: self._select_site(r["key"]),
            empty="Every source site is in the file.")
        self.site_list.grid(row=1, column=0, sticky="nsew", padx=(16, 8),
                            pady=(0, 10))
        self._build_rows()

        bar = ctk.CTkFrame(self, fg_color="transparent")
        bar.grid(row=2, column=0, columnspan=2, sticky="ew", padx=16,
                 pady=(0, 14))
        bar.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(bar, text="Nothing is sent, and the file on disk is "
                               "left as it is.",
                     anchor="w", font=(UI_FONT, 11),
                     text_color=TEXT_FAINT).grid(row=0, column=0, sticky="ew")
        _ghost(bar, "Cancel", self.destroy, width=100, height=34).grid(
            row=0, column=1, padx=(0, 8))
        self._apply_btn = ctk.CTkButton(
            bar, text="Pair & check again", width=230, height=34,
            corner_radius=RADIUS_MD, font=(UI_FONT, 13, "bold"),
            border_color=BORDER, text_color_disabled=TEXT_FAINT,
            command=self._apply, **_GHOSTED)
        self._apply_btn.grid(row=0, column=2)

        self.site_list.set_rows([{"key": i, "site": s}
                                 for i, s in enumerate(self._sites)])
        self._select_site(0 if self._sites else None)
        self.bind("<Escape>", lambda _e: self.destroy())
        self.transient(parent.winfo_toplevel())
        self.after(120, self._grab)

    def _build_rows(self):
        card = ctk.CTkFrame(self, fg_color=CARD, corner_radius=RADIUS_LG,
                            border_width=1, border_color=BORDER)
        card.grid(row=1, column=1, sticky="nsew", padx=(8, 16), pady=(0, 10))
        card.grid_columnconfigure(0, weight=1)
        card.grid_rowconfigure(3, weight=1)
        self._title = ctk.CTkLabel(card, text="", anchor="w", justify="left",
                                   font=(UI_FONT, 13, "bold"))
        self._title.grid(row=0, column=0, sticky="ew", padx=16, pady=(14, 0))
        _wrap(card, self._title, pad=32)
        self._sub = ctk.CTkLabel(card, text="", anchor="w", justify="left",
                                 font=(UI_FONT, 12), text_color=TEXT_MUTED)
        self._sub.grid(row=1, column=0, sticky="ew", padx=16)
        self._search = ctk.CTkEntry(
            card, height=30,
            placeholder_text="Search the file's rows by name or row number…")
        self._search.grid(row=2, column=0, sticky="ew", padx=16, pady=8)
        self._search.bind("<KeyRelease>", lambda _e: self._filter_rows())
        self._search.bind("<Return>", lambda _e: self._use_only())
        self.row_list = _PairRowList(
            card, self._owner, self._use, on_select=self._on_row,
            empty="No row of the file is free — each one already has its "
                  "source site.")
        self.row_list.grid(row=3, column=0, sticky="nsew", padx=16)
        self._note = ctk.CTkLabel(card, text="", anchor="w", justify="left",
                                  font=(UI_FONT, 11))
        self._note.grid(row=4, column=0, sticky="ew", padx=16, pady=(8, 0))
        _wrap(card, self._note, pad=32)
        acts = ctk.CTkFrame(card, fg_color="transparent")
        acts.grid(row=5, column=0, sticky="w", padx=16, pady=(8, 14))
        self._use_btn = ctk.CTkButton(
            acts, text="Use selected row", width=220, height=32,
            corner_radius=RADIUS_MD, font=(UI_FONT, 12, "bold"),
            border_color=BORDER, text_color_disabled=TEXT_FAINT,
            command=self._use, **_GHOSTED)
        self._use_btn.grid(row=0, column=0, padx=(0, 6))
        self._clear_btn = _ghost(acts, "Leave unpaired", self._leave,
                                 width=130)
        self._clear_btn.grid(row=0, column=1)

    def _grab(self):
        try:
            self.lift()
            self.focus_force()
            self.grab_set()
        except Exception:
            pass            # closed before it was ever shown

    # ── choosing ──────────────────────────────────────────────────────

    def _site(self):
        return None if self._current is None else self._sites[self._current]

    def _owner(self, entry) -> str:
        """The other site this row is already picked for, or ''."""
        for index, picked in self._picks.items():
            if picked is entry and index != self._current:
                return self._sites[index]["site"]
        return ""

    def _select_site(self, index):
        self._current = index
        self.site_list._selected = index
        site = self._site()
        if site is None:
            self._row = None
            self.row_list.clear()
            self._paint()
            return
        key = loose_name(site["site"])
        ranked = rank_file_rows(site["site"], self._rows)
        self.row_list.set_rows([{"key": e["row"], "entry": e,
                                 "same": loose_name(e["source"]) == key}
                                for e in ranked])
        self._row = self._picks.get(index)
        if (self._row is None and ranked
                and loose_name(ranked[0]["source"]) == key
                and not self._owner(ranked[0])):
            self._row = ranked[0]       # selected for a look, not picked
        self.row_list._selected = self._row["row"] if self._row else None
        self._filter_rows()

    def _filter_rows(self):
        text = self._search.get().strip()
        query = loose_name(text)

        def match(rec):
            entry = rec["entry"]
            return (query in loose_name(entry["source"])
                    or query in loose_name(entry["dest"])
                    or text == str(entry["row"]))
        self.row_list.filter_by(match if query else None)
        self._paint()

    def _on_row(self, rec):
        self._row = rec["entry"]
        self._paint()

    def _use(self):
        site, entry = self._site(), self._row
        if site is None or entry is None or self._owner(entry):
            return
        self._picks[self._current] = entry
        nxt = next((i for i in range(len(self._sites))
                    if i not in self._picks), None)
        if nxt is not None:
            self._select_site(nxt)
        else:
            self._paint()

    def _use_only(self):
        """Return in the search box: use the one free row it left."""
        self.row_list.counts()
        free = [self.row_list._records[k]["entry"] for k in self.row_list._view
                if not self._owner(self.row_list._records[k]["entry"])]
        if len(free) == 1:
            self._row = free[0]
            self.row_list._selected = free[0]["row"]
            self._use()

    def _leave(self):
        if self._current is None:
            return
        self._picks.pop(self._current, None)
        self._row = None
        self.row_list._selected = None
        self._paint()

    # ── painting ──────────────────────────────────────────────────────

    def _row_note(self, site, entry, picked):
        """(text, colour, can use) under the rows."""
        if site is None:
            return "", TEXT_MUTED, False
        name = site["site"]
        agents = _plural(int(site.get("agents") or 0), "agent")
        if entry is None:
            if picked is not None:
                return (f"“{name}” gets the token of row {picked['row']}, "
                        f"“{picked['source']}”. Select another row to change "
                        f"it.", GREEN, False)
            if not self._rows:
                return "", TEXT_MUTED, False
            return (f"Select the row whose token the {agents} of “{name}” "
                    f"should get — closest names first.", TEXT_MUTED, False)
        taken = self._owner(entry)
        if taken:
            return (f"Row {entry['row']} is already your pick for “{taken}” "
                    f"— a row's token goes to one site.", ACCENT, False)
        if picked is entry:
            return (f"“{name}” gets the token of row {entry['row']} and lands "
                    f"in {file_destination(entry)}.", GREEN, False)
        lead = (f"Row {entry['row']}, “{entry['source']}”, is the same name "
                f"spelt differently. "
                if loose_name(entry["source"]) == loose_name(name) else "")
        return (lead + f"Its {agents} will move with this row's token into "
                f"{file_destination(entry)}.", TEXT, True)

    def _paint(self):
        site, entry = self._site(), self._row
        picked = self._picks.get(self._current) if site else None
        _retext(self._title, text=(f"Rows of the file for “{site['site']}”"
                                   if site else "Nothing to pair"))
        _retext(self._sub, text=(
            f"{site['account']}  ·  "
            f"{_plural(int(site.get('agents') or 0), 'agent')} on the source"
            if site else ""))
        text, tone, can_use = self._row_note(site, entry, picked)
        _retext(self._note, text=text, text_color=tone)
        if can_use:
            _retext(self._use_btn, state="normal",
                    text=f"Use row {entry['row']} for "
                         f"“{_short(site['site'], 24)}”", **_PRIMARY)
        else:
            _retext(self._use_btn, state="disabled", text="Use selected row",
                    **_GHOSTED)
        _grid_shown(self._clear_btn, picked is not None)
        count = len(self._picks)
        if count:
            _retext(self._apply_btn, state="normal",
                    text=f"Pair {_plural(count, 'site')} & check again", **_GO)
        else:
            _retext(self._apply_btn, state="disabled",
                    text="Pair & check again", **_GHOSTED)
        self.site_list.render()
        self.row_list.render()

    def _apply(self):
        pairs = [(entry, self._sites[index])
                 for index, entry in sorted(self._picks.items())]
        self.destroy()
        if pairs:
            self._on_done(pairs)
